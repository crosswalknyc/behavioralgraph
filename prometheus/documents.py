"""Documents in the Prometheus chat: upload, read, edit.

Jenna 2026-10-06: "update prometheus so that you can upload a document
to it. like a powerpoint to be edited, or pdf or other data."

A user drops a file on the chat (or taps the paperclip). It lands in
S3 under the asker's thread, its text is pulled out once, and the
thread carries it from then on:

  * a question about it ("what does slide 4 say", "summarize this",
    "what is the churn number in the pdf") is answered from the
    document's own text, with slide or page references;
  * an instruction about a deck or Word file ("change the title on
    slide 3 to X", "replace 22.8M with 411,324", "delete slide 7") is
    turned into a small edit plan, applied in place with the file's
    formatting kept, and handed back as a download;
  * a sheet (csv / xlsx) is summarized by column and can be asked about.

Nothing here touches the catalog as a published figure: a document a
user brought is theirs, not ours. Every path is fail-safe and the lane
returns None whenever the ask is not about the attached files.
"""
import io
import json
import os
import re
import time
import traceback
import uuid
from datetime import datetime, timezone

from .host import host

UPLOAD_PREFIX = 'system/prometheus_uploads/'
OUTPUT_PREFIX = 'generated_decks/edited/'
MAX_BYTES = 25 * 1024 * 1024
MAX_TEXT = 60_000          # characters kept per document
CONTEXT_CHARS = 18_000     # characters of document text handed to the model
KINDS = {
    '.pptx': 'deck', '.pdf': 'pdf', '.docx': 'doc', '.xlsx': 'sheet', '.xlsm': 'sheet',
    '.csv': 'sheet', '.tsv': 'sheet', '.txt': 'text', '.md': 'text', '.json': 'text',
    '.png': 'image', '.jpg': 'image', '.jpeg': 'image', '.webp': 'image',
}
KIND_NOUN = {'deck': 'deck', 'pdf': 'PDF', 'doc': 'document', 'sheet': 'sheet',
             'text': 'text file', 'image': 'image'}

_DOC_REF_RX = re.compile(
    r"\b(this|the|that|my|attached|uploaded)\s+(deck|slides?|presentation|powerpoint|pptx|ppt|pdf|"
    r"doc|document|word file|file|sheet|spreadsheet|csv|xlsx|excel|upload|attachment|image|screenshot)\b"
    r"|\bslide\s*\d+\b|\bpage\s*\d+\b|\bin the (deck|pdf|file|document|sheet|upload)\b"
    r"|\b(summari[sz]e|tl;?dr|read|skim|go through|walk me through)\s+(it|this|that)\b"
    r"|\bwhat does (it|this|that) say\b|\bfrom the (deck|pdf|file|document|sheet)\b", re.I)
_EDIT_RX = re.compile(
    r"\b(change|replace|swap|update|rename|retitle|rewrite|reword|edit|fix|correct|delete|remove|"
    r"drop|add|insert|append|move|reorder|make (it|the \w+) say|should say|should read)\b", re.I)


def _now():
    return datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')


def kind_of(name):
    ext = os.path.splitext(str(name or '').lower())[1]
    return KINDS.get(ext)


def safe_name(name):
    base = os.path.basename(str(name or 'file'))
    base = re.sub(r'[^A-Za-z0-9._ ()-]+', '_', base).strip(' ._') or 'file'
    return base[:120]


# -------------------------------------------------------------- extract

def _pptx_inventory(data):
    from pptx import Presentation  # type: ignore
    prs = Presentation(io.BytesIO(data))
    slides = []
    for n, slide in enumerate(prs.slides, start=1):
        title = ''
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:
            title = ''
        texts = []
        for shape in _iter_shapes(slide.shapes):
            try:
                if getattr(shape, 'has_text_frame', False) and shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t and t != title:
                        texts.append(t)
                if getattr(shape, 'has_table', False) and shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            texts.append(' | '.join(cells))
            except Exception:
                continue
        notes = ''
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ''
        slides.append({'n': n, 'title': title, 'text': '\n'.join(texts), 'notes': notes})
    return slides


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        try:
            if shape.shape_type == 6 and hasattr(shape, 'shapes'):  # group
                for s in _iter_shapes(shape.shapes):
                    yield s
        except Exception:
            continue


def _pdf_text(data):
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages[:80], start=1):
            try:
                pages.append((i, (page.extract_text() or '').strip()))
            except Exception:
                pages.append((i, ''))
        return pages, len(reader.pages)
    except Exception:
        pass
    try:
        import fitz  # type: ignore
        doc = fitz.open(stream=data, filetype='pdf')
        pages = [(i + 1, (doc[i].get_text() or '').strip()) for i in range(min(len(doc), 80))]
        return pages, len(doc)
    except Exception:
        return [], 0


def _docx_text(data):
    import docx  # type: ignore
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


def _sheet_summary(name, data):
    import pandas as pd  # type: ignore
    ext = os.path.splitext(name.lower())[1]
    frames = {}
    if ext in ('.csv', '.tsv'):
        sep = '\t' if ext == '.tsv' else ','
        frames['Sheet1'] = pd.read_csv(io.BytesIO(data), sep=sep, low_memory=False)
    else:
        frames = pd.read_excel(io.BytesIO(data), sheet_name=None)
    sheets, lines = [], []
    for sname, df in list(frames.items())[:12]:
        cols = [str(c) for c in df.columns][:60]
        sheets.append({'name': str(sname), 'rows': int(len(df)), 'columns': cols})
        lines.append(f"Sheet {sname}: {len(df):,} rows, {len(df.columns)} columns")
        lines.append('Columns: ' + ', '.join(cols))
        try:
            lines.append(df.head(12).to_string(max_cols=20, max_colwidth=40))
        except Exception:
            pass
        try:
            num = df.select_dtypes('number')
            if len(num.columns):
                lines.append('Numeric summary:\n' + num.describe().round(2).to_string(max_cols=12))
        except Exception:
            pass
    return sheets, '\n'.join(lines)


def extract(name, data):
    """{'kind', 'text', 'slides', 'pages', 'rows', 'sheets', 'summary'}.
    Never raises; an unreadable file still lands with an honest note."""
    kind = kind_of(name) or 'text'
    out = {'kind': kind, 'text': '', 'slides': [], 'pages': 0, 'rows': 0, 'sheets': [], 'summary': ''}
    try:
        if kind == 'deck':
            slides = _pptx_inventory(data)
            out['slides'] = slides
            out['pages'] = len(slides)
            out['text'] = '\n\n'.join(
                f"Slide {s['n']}: {s['title'] or '(no title)'}\n{s['text']}"
                + (f"\nNotes: {s['notes']}" if s['notes'] else '') for s in slides)
            out['summary'] = f"{len(slides)} slide{'s' if len(slides) != 1 else ''}"
        elif kind == 'pdf':
            pages, n = _pdf_text(data)
            out['pages'] = n
            out['text'] = '\n\n'.join(f"Page {i}:\n{t}" for i, t in pages if t)
            out['summary'] = f"{n} page{'s' if n != 1 else ''}" + ('' if out['text'] else ', no readable text (a scan)')
        elif kind == 'doc':
            out['text'] = _docx_text(data)
            out['summary'] = f"{len(out['text'].split()):,} words"
        elif kind == 'sheet':
            sheets, text = _sheet_summary(name, data)
            out['sheets'], out['text'] = sheets, text
            out['rows'] = sum(s['rows'] for s in sheets)
            out['summary'] = (f"{out['rows']:,} rows across {len(sheets)} sheet{'s' if len(sheets) != 1 else ''}"
                              if len(sheets) != 1 else f"{out['rows']:,} rows, {len(sheets[0]['columns'])} columns")
        elif kind == 'image':
            out['summary'] = 'image'
        else:
            out['text'] = data.decode('utf-8', errors='replace')
            out['summary'] = f"{len(out['text'].split()):,} words"
    except Exception as e:
        traceback.print_exc()
        out['summary'] = out['summary'] or 'could not be read'
        out['error'] = str(e)[:200]
    out['text'] = (out['text'] or '')[:MAX_TEXT]
    return out


# ---------------------------------------------------------------- store

def _s3():
    return host.s3_client, host.bucket


def _index_key(uname, tid):
    return f"{UPLOAD_PREFIX}{_safe(uname)}/{_safe(tid or 'active')}/index.json"


def _safe(s):
    return re.sub(r'[^A-Za-z0-9_.@-]+', '_', str(s or ''))[:80] or 'x'


def load_index(uname, tid, s3=None, bucket=None):
    if s3 is None:
        s3, bucket = _s3()
    try:
        return json.loads(s3.get_object(Bucket=bucket, Key=_index_key(uname, tid))['Body'].read().decode('utf-8'))
    except Exception:
        return {'attachments': []}


def _save_index(uname, tid, idx, s3, bucket):
    s3.put_object(Bucket=bucket, Key=_index_key(uname, tid),
                  Body=json.dumps(idx, ensure_ascii=False).encode('utf-8'), ContentType='application/json')


def store_upload(uname, tid, name, data, *, s3=None, bucket=None):
    """Write the file and its text, append to the thread's attachment
    index, return the record."""
    if s3 is None:
        s3, bucket = _s3()
    name = safe_name(name)
    kind = kind_of(name)
    if not kind:
        raise ValueError('unsupported file type')
    if len(data) > MAX_BYTES:
        raise ValueError('file is larger than 25 MB')
    uid = uuid.uuid4().hex[:12]
    base = f"{UPLOAD_PREFIX}{_safe(uname)}/{_safe(tid or 'active')}/{uid}_"
    file_key, text_key = base + name, base + 'text.txt'
    s3.put_object(Bucket=bucket, Key=file_key, Body=data, ContentType='application/octet-stream')
    ex = extract(name, data)
    if ex.get('text'):
        s3.put_object(Bucket=bucket, Key=text_key, Body=ex['text'].encode('utf-8'), ContentType='text/plain')
    rec = {'upload_id': uid, 'name': name, 'kind': kind, 'bytes': len(data), 'file_key': file_key,
           'text_key': text_key if ex.get('text') else '', 'pages': ex.get('pages') or 0,
           'rows': ex.get('rows') or 0, 'sheets': ex.get('sheets') or [],
           'slides': [{'n': s['n'], 'title': s['title']} for s in (ex.get('slides') or [])],
           'summary': ex.get('summary') or '', 'excerpt': (ex.get('text') or '')[:600],
           'uploaded_at': _now(), 'error': ex.get('error', '')}
    idx = load_index(uname, tid, s3, bucket)
    idx.setdefault('attachments', []).append(rec)
    idx['attachments'] = idx['attachments'][-20:]
    _save_index(uname, tid, idx, s3, bucket)
    return rec


def attachments(uname, tid, s3=None, bucket=None):
    return list((load_index(uname, tid, s3, bucket) or {}).get('attachments') or [])


def document_text(rec, s3=None, bucket=None):
    if not rec or not rec.get('text_key'):
        return ''
    if s3 is None:
        s3, bucket = _s3()
    try:
        return s3.get_object(Bucket=bucket, Key=rec['text_key'])['Body'].read().decode('utf-8', 'replace')
    except Exception:
        return ''


def document_bytes(rec, s3=None, bucket=None):
    if s3 is None:
        s3, bucket = _s3()
    return s3.get_object(Bucket=bucket, Key=rec['file_key'])['Body'].read()


# ------------------------------------------------------------- wording

def ack(rec):
    """The reply after an upload lands, with the next moves as chips."""
    noun = KIND_NOUN.get(rec.get('kind'), 'file')
    head = f"Got it: {rec['name']}" + (f" ({rec['summary']})" if rec.get('summary') else '') + '.'
    if rec.get('kind') in ('deck', 'doc'):
        body = (f" Ask me to summarize the {noun}, pull a number or a line from it, or tell me what to "
                f"change (\"change the title on slide 3 to ...\", \"replace 22.8M with 411,324\") and I will "
                f"hand back an edited copy with the formatting kept.")
        chips = ['Summarize it', 'What are the key numbers in it?', 'Change the title on slide 1']
    elif rec.get('kind') == 'sheet':
        body = " Ask me what is in it, for a number or a column, or how it compares with a profile you have."
        chips = ['Summarize it', 'What are the columns and what do they hold?']
    elif rec.get('kind') == 'pdf':
        body = (" Ask me to summarize it or pull a number or a passage from it. I read PDFs; I do not edit "
                "them in place, but I can turn what you want changed into a new deck or document.")
        chips = ['Summarize it', 'What are the key numbers in it?']
    elif rec.get('kind') == 'image':
        body = " I have it on the thread. Tell me what you want done with it."
        chips = []
    else:
        body = " Ask me anything about it."
        chips = ['Summarize it']
    if rec.get('error') or (rec.get('kind') in ('deck', 'pdf', 'doc', 'sheet') and not rec.get('excerpt')):
        body = " I have the file, but I could not read text out of it. If it is a scan or an image export, send the original and I will take it from there."
        chips = []
    return head + body, chips


def references_document(text, recs, history=None):
    """Is this ask about the attached files? Named reference words, a
    slide or page number, the file's own name, or the turn right after
    an upload."""
    t = str(text or '')
    if not recs:
        return False
    if _DOC_REF_RX.search(t):
        return True
    low = t.lower()
    for r in recs:
        stem = os.path.splitext(str(r.get('name') or ''))[0].lower()
        if stem and len(stem) >= 4 and stem in low:
            return True
    # The turn right after the upload: "summarize it", "key numbers?"
    turns = [h for h in (history or []) if isinstance(h, dict)]
    for h in reversed(turns[-3:]):
        if str(h.get('role') or '') in ('agent', 'assistant'):
            if (h.get('meta') or {}).get('kind') == 'upload_ack':
                return True
            break
    return False


def is_edit_ask(text):
    return bool(_EDIT_RX.search(str(text or '')))


# ---------------------------------------------------------------- edits

EDIT_SYSTEM = """You turn one instruction about a presentation or document into
a small JSON edit plan. You are given the slide (or section) inventory:
numbers, titles and text. Return STRICT JSON only:

{"ops": [
   {"op": "replace", "find": str, "replace": str, "slide": int|null},
   {"op": "set_title", "slide": int, "text": str},
   {"op": "delete_slide", "slide": int},
   {"op": "add_slide", "title": str, "bullets": [str, ...], "after": int|null},
   {"op": "set_notes", "slide": int, "text": str}
 ],
 "summary": str}

Rules: only what the instruction asks for, nothing extra. Every
"find" string must appear verbatim in the inventory text (copy it
exactly, including punctuation and numbers). A number change is a
replace on the exact old number string. When the instruction names a
slide, keep the op on that slide. If the instruction cannot be mapped
to these ops, return {"ops": [], "summary": "<one plain sentence saying
what you could not do>"}."""


def plan_edits(claude_data, rec, text, doc_text):
    inv = doc_text[:CONTEXT_CHARS]
    user = json.dumps({'file': rec.get('name'), 'kind': rec.get('kind'),
                       'instruction': str(text or '').strip(), 'inventory': inv})
    data = claude_data(EDIT_SYSTEM, user, max_tokens=2500, temperature=0.1, surface='document_edit')
    if not isinstance(data, dict):
        return [], ''
    ops = [o for o in (data.get('ops') or []) if isinstance(o, dict) and o.get('op')]
    return ops, str(data.get('summary') or '').strip()


def _replace_in_text_frame(tf, find, repl):
    """Replace inside runs when the match sits in one run (formatting
    kept exactly); when it spans runs, rewrite the paragraph into its
    first run (paragraph style kept)."""
    n = 0
    for p in tf.paragraphs:
        done = False
        for r in p.runs:
            if find in r.text:
                r.text = r.text.replace(find, repl)
                n += 1
                done = True
        if done:
            continue
        joined = ''.join(r.text for r in p.runs)
        if find in joined and p.runs:
            p.runs[0].text = joined.replace(find, repl)
            for r in p.runs[1:]:
                r.text = ''
            n += 1
    return n


def _slide_text_frames(slide):
    for shape in _iter_shapes(slide.shapes):
        try:
            if getattr(shape, 'has_text_frame', False) and shape.has_text_frame:
                yield shape.text_frame
            if getattr(shape, 'has_table', False) and shape.has_table:
                for row in shape.table.rows:
                    for c in row.cells:
                        yield c.text_frame
        except Exception:
            continue


def apply_edits_pptx(data, ops):
    from pptx import Presentation  # type: ignore
    from pptx.util import Pt  # type: ignore
    prs = Presentation(io.BytesIO(data))
    applied, skipped = [], []
    slides = list(prs.slides)

    def _slide(n):
        try:
            n = int(n)
        except (TypeError, ValueError):
            return None
        return slides[n - 1] if 1 <= n <= len(slides) else None

    to_delete = []
    for op in ops:
        kind = str(op.get('op') or '')
        try:
            if kind == 'replace':
                find, repl = str(op.get('find') or ''), str(op.get('replace') or '')
                if not find:
                    skipped.append('replace with an empty find'); continue
                targets = [_slide(op.get('slide'))] if op.get('slide') else slides
                hits = 0
                for s in targets:
                    if s is None:
                        continue
                    for tf in _slide_text_frames(s):
                        hits += _replace_in_text_frame(tf, find, repl)
                    try:
                        if s.has_notes_slide:
                            hits += _replace_in_text_frame(s.notes_slide.notes_text_frame, find, repl)
                    except Exception:
                        pass
                (applied if hits else skipped).append(
                    f'replaced "{find}" with "{repl}"' + (f' ({hits} place{"s" if hits != 1 else ""})' if hits else ' (not found)'))
            elif kind == 'set_title':
                s = _slide(op.get('slide'))
                if s is None or s.shapes.title is None:
                    skipped.append(f"set title on slide {op.get('slide')} (no title box)"); continue
                tf = s.shapes.title.text_frame
                if tf.paragraphs and tf.paragraphs[0].runs:
                    tf.paragraphs[0].runs[0].text = str(op.get('text') or '')
                    for r in tf.paragraphs[0].runs[1:]:
                        r.text = ''
                    for p in tf.paragraphs[1:]:
                        for r in p.runs:
                            r.text = ''
                else:
                    tf.text = str(op.get('text') or '')
                applied.append(f"slide {op.get('slide')} title set to \"{op.get('text')}\"")
            elif kind == 'set_notes':
                s = _slide(op.get('slide'))
                if s is None:
                    skipped.append(f"notes on slide {op.get('slide')}"); continue
                s.notes_slide.notes_text_frame.text = str(op.get('text') or '')
                applied.append(f"slide {op.get('slide')} notes updated")
            elif kind == 'delete_slide':
                s = _slide(op.get('slide'))
                if s is None:
                    skipped.append(f"delete slide {op.get('slide')} (no such slide)"); continue
                to_delete.append(int(op.get('slide')))
            elif kind == 'add_slide':
                layout = None
                for lay in prs.slide_layouts:
                    if 'title and content' in str(lay.name or '').lower():
                        layout = lay; break
                if layout is None:
                    layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
                new = prs.slides.add_slide(layout)
                if new.shapes.title is not None:
                    new.shapes.title.text = str(op.get('title') or '')
                body = next((sh for sh in new.placeholders if sh.placeholder_format.idx != 0), None)
                bullets = [str(b) for b in (op.get('bullets') or []) if str(b).strip()]
                if body is not None and bullets:
                    tf = body.text_frame
                    tf.text = bullets[0]
                    for b in bullets[1:]:
                        p = tf.add_paragraph(); p.text = b
                    for p in tf.paragraphs:
                        for r in p.runs:
                            r.font.size = r.font.size or Pt(18)
                after = op.get('after')
                try:
                    if after is not None:
                        xml_slides = prs.slides._sldIdLst  # noqa: SLF001
                        el = xml_slides[-1]
                        xml_slides.remove(el)
                        xml_slides.insert(int(after), el)
                except Exception:
                    pass
                applied.append(f"added slide \"{op.get('title')}\"" + (f" after slide {after}" if after else ' at the end'))
            else:
                skipped.append(f'unknown op {kind}')
        except Exception as e:
            traceback.print_exc()
            skipped.append(f'{kind}: {str(e)[:80]}')
    if to_delete:
        xml_slides = prs.slides._sldIdLst  # noqa: SLF001
        ids = list(xml_slides)
        for n in sorted(set(to_delete), reverse=True):
            try:
                el = ids[n - 1]
                rid = el.rId
                prs.part.drop_rel(rid)
                xml_slides.remove(el)
                applied.append(f'deleted slide {n}')
            except Exception as e:
                skipped.append(f'delete slide {n}: {str(e)[:60]}')
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), applied, skipped


def apply_edits_docx(data, ops):
    import docx  # type: ignore
    d = docx.Document(io.BytesIO(data))
    applied, skipped = [], []

    def _paras():
        for p in d.paragraphs:
            yield p
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        yield p
    for op in ops:
        kind = str(op.get('op') or '')
        if kind == 'replace':
            find, repl = str(op.get('find') or ''), str(op.get('replace') or '')
            hits = 0
            for p in _paras():
                done = False
                for r in p.runs:
                    if find in r.text:
                        r.text = r.text.replace(find, repl); hits += 1; done = True
                if not done:
                    joined = ''.join(r.text for r in p.runs)
                    if find in joined and p.runs:
                        p.runs[0].text = joined.replace(find, repl)
                        for r in p.runs[1:]:
                            r.text = ''
                        hits += 1
            (applied if hits else skipped).append(f'replaced "{find}" with "{repl}"' + ('' if hits else ' (not found)'))
        elif kind == 'add_slide':
            d.add_paragraph(str(op.get('title') or ''), style='Heading 1')
            for b in (op.get('bullets') or []):
                d.add_paragraph(str(b), style='List Bullet')
            applied.append(f"added section \"{op.get('title')}\"")
        else:
            skipped.append(f'{kind} does not apply to a document')
    out = io.BytesIO()
    d.save(out)
    return out.getvalue(), applied, skipped


def publish_output(rec, data, s3=None, bucket=None):
    """Write the edited copy and return a 7-day link."""
    if s3 is None:
        s3, bucket = _s3()
    stem, ext = os.path.splitext(rec['name'])
    out_name = f"{stem} (edited){ext}"
    key = f"{OUTPUT_PREFIX}{uuid.uuid4().hex[:12]}/{safe_name(out_name)}"
    ctype = ('application/vnd.openxmlformats-officedocument.presentationml.presentation' if ext.lower() == '.pptx'
             else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    s3.put_object(Bucket=bucket, Key=key, Body=data, ContentType=ctype,
                  ContentDisposition=f'attachment; filename="{safe_name(out_name)}"')
    url = s3.generate_presigned_url('get_object', Params={'Bucket': bucket, 'Key': key}, ExpiresIn=7 * 24 * 3600)
    return key, url, safe_name(out_name)


# ---------------------------------------------------------------- read

READ_SYSTEM = """You answer a question about a document the user attached to the
chat, using ONLY the document text you are given. Plain English, short
sentences, numbers exactly as the document states them, and point to
the slide or page ("slide 4", "page 12") each fact comes from. If the
document does not contain the answer, say so in one sentence and name
the closest thing it does say. Return STRICT JSON:
{"answer": str, "followups": [str, str]}"""


def answer_question(claude_data, recs, text, doc_texts):
    docs = []
    budget = CONTEXT_CHARS
    for rec, dt in zip(recs, doc_texts):
        piece = (dt or '')[:max(2000, budget // max(1, len(recs)))]
        docs.append({'file': rec.get('name'), 'kind': rec.get('kind'), 'summary': rec.get('summary'), 'text': piece})
    user = json.dumps({'question': str(text or '').strip(), 'documents': docs})
    data = claude_data(READ_SYSTEM, user, max_tokens=1800, temperature=0.2, surface='document_read')
    if not isinstance(data, dict) or not str(data.get('answer') or '').strip():
        return '', []
    fus = [str(f) for f in (data.get('followups') or []) if str(f).strip()][:3]
    return str(data['answer']).strip(), fus


# ---------------------------------------------------------------- lane

def _claude():
    try:
        if host.has('claude_data'):
            return host.claude_data
    except Exception:
        pass
    return None


def _raw(reply, followups=None, **extra):
    out = {'success': True, 'action': 'answer', 'reply': reply, 'followups': list(followups or []),
           'offer_deck': False, 'deck_angle': None}
    out.update(extra)
    return out


def answer(text, uname, tid, history=None, *, s3=None, bucket=None, claude_data=None):
    """None when the thread has no attachments or the ask is not about
    them. Otherwise the raw analyze payload: an edited copy with a
    download link, or an answer read out of the document."""
    if s3 is None:
        try:
            s3, bucket = _s3()
        except Exception:
            return None
    recs = attachments(uname, tid, s3, bucket)
    if not recs or not references_document(text, recs, history):
        return None
    claude_data = claude_data or _claude()
    # The most recent file the ask names, else the most recent file.
    low = str(text or '').lower()
    target = None
    for r in reversed(recs):
        stem = os.path.splitext(str(r.get('name') or ''))[0].lower()
        if stem and len(stem) >= 4 and stem in low:
            target = r; break
    target = target or recs[-1]
    editable = target.get('kind') in ('deck', 'doc')
    if editable and is_edit_ask(text) and claude_data is not None:
        doc_text = document_text(target, s3, bucket)
        ops, summary = plan_edits(claude_data, target, text, doc_text)
        if not ops:
            return _raw(summary or "I could not map that to a change in the file. Tell me the exact words or number to change and what to change them to.",
                        ['Summarize it'], document='edit_unmapped')
        data = document_bytes(target, s3, bucket)
        if target.get('kind') == 'deck':
            out, applied, skipped = apply_edits_pptx(data, ops)
        else:
            out, applied, skipped = apply_edits_docx(data, ops)
        if not applied:
            return _raw("I found the file but none of those changes matched its text: " + '; '.join(skipped[:4])
                        + ". Tell me the exact words as they appear and I will try again.", ['Summarize it'],
                        document='edit_nohit')
        key, url, out_name = publish_output(target, out, s3, bucket)
        lines = '\n'.join(f"- {a}" for a in applied[:12])
        tail = ('\n\nNot applied: ' + '; '.join(skipped[:4]) + '.') if skipped else ''
        reply = (f"Done. Here is {out_name} with the formatting kept:\n{lines}{tail}\n\n"
                 f"The link is good for 7 days. Tell me the next change and I will work from this copy.")
        # The edited copy becomes the thread's current version.
        try:
            new_rec = store_upload(uname, tid, out_name, out, s3=s3, bucket=bucket)
            new_rec['derived_from'] = target.get('upload_id')
            idx = load_index(uname, tid, s3, bucket)
            for r in idx.get('attachments') or []:
                if r.get('upload_id') == new_rec['upload_id']:
                    r['derived_from'] = target.get('upload_id')
            _save_index(uname, tid, idx, s3, bucket)
        except Exception:
            traceback.print_exc()
        return _raw(reply, ['Summarize it', 'Make another change'],
                    file_link={'url': url, 'label': f'Download {out_name}'}, document='edited')
    if claude_data is None:
        return _raw(f"I have {target['name']} on the thread ({target.get('summary') or 'read'}). "
                    f"The reading step is not available right now; try again in a minute.", [], document='no_model')
    # Read from the file the ask names (default: the most recent
    # version); older copies ride along only when the ask names them.
    use = [target] + [r for r in recs if r is not target
                      and os.path.splitext(str(r.get('name') or ''))[0].lower() in low
                      and len(os.path.splitext(str(r.get('name') or ''))[0]) >= 4][-2:]
    texts = [document_text(r, s3, bucket) for r in use]
    reply, fus = answer_question(claude_data, use, text, texts)
    if not reply:
        return _raw(f"I have {target['name']} on the thread but could not read an answer to that out of it. "
                    f"Ask for a specific slide, page, number or section.", ['Summarize it'], document='read_empty')
    return _raw(reply, fus, document='read')
