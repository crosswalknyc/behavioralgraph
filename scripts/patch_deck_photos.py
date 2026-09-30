#!/usr/bin/env python3
"""Deck photography (Jenna 2026-09-30: "when building decks can
prometheus find cool images to put into the decks to mirror how they
look in the designed decks we do manually").

Mirrors the manual deck grammar (the Paige Bueckers reference): full-
bleed photography under a graphite scrim on the cover, the first big
statement page, and the close. Images come from the standing image
resolver (image_backfill: IMDb, Wikipedia, web image search - the
same stack that populates dashboard profile art). Everything is
fail-safe: no image, no network, or a kill switch means the deck
renders text-only exactly as today."""
import ast
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / 'deck_builder.py'
PMA = ROOT / 'prometheus_analysis.py'
APP = ROOT / 'app.py'


def splice(src, old, new, desc):
    n = src.count(old)
    if n != 1:
        raise RuntimeError(f'[{desc}] anchor count {n}')
    return src.replace(old, new)


# ========================= deck_builder.py =========================
db = DB.read_text(encoding='utf-8')

PHOTO_LAYER = '''

# ---------------- Deck photography (2026-09-30, Jenna) -------------
# "when building decks can prometheus find cool images to put into
# the decks to mirror how they look in the designed decks we do
# manually". Full-bleed photos under a graphite scrim on the cover,
# the first big statement page, and the close - the manual deck
# grammar. Sourced through the standing image resolver
# (image_backfill: IMDb, Wikipedia, web image search). Fail-safe by
# design: any miss renders the slide text-only exactly as before.

_PHOTO_CACHE = {}
_PHOTO_MIN_W = 640
_PHOTO_MIN_H = 420


def _photos_enabled():
    return os.environ.get('PROMETHEUS_DECK_PHOTOS', '1').strip().lower() \\
        not in ('0', 'false', 'no')


def _blob_px(blob):
    try:
        from pptx.parts.image import Image as _PImg
        return _PImg.from_blob(blob).size
    except Exception:
        return (0, 0)


_PHOTO_EXTS = ('jpg', 'jpeg', 'png', 'gif', 'bmp', 'tiff')


def _blob_embeddable(blob):
    """True when python-pptx can embed the image (WEBP and friends
    cannot ship in a PPTX)."""
    try:
        from pptx.parts.image import Image as _PImg
        img = _PImg.from_blob(blob)
        return (str(img.ext or '').lower() in _PHOTO_EXTS
                and img.size[0] >= _PHOTO_MIN_W
                and img.size[1] >= _PHOTO_MIN_H)
    except Exception:
        return False


def _fetch_photos(subject, kind, max_n):
    """Candidate photo bytes for the deck subject, best sources
    first. Never raises."""
    import re as _re
    import urllib.parse as _up
    try:
        import image_backfill as ib
    except Exception:
        return []
    master = {'person': 'TALENT', 'title': 'CONTENT'}.get(kind, 'BRAND')
    urls = []
    try:
        u, _src = ib.resolve_image_url(subject, master)
        if u:
            urls.append(u)
    except Exception:
        pass
    try:
        u = ib.wiki_lookup(subject)
        if u:
            urls.append(u)
    except Exception:
        pass
    try:
        data, _ct = ib._http_get(
            ib.BING_IMAGES_URL.format(q=_up.quote(f"{subject} photo")),
            ib.HEADERS_JSON, timeout=8.0)
        html = data.decode('utf-8', 'ignore') if data else ''
        for raw in _re.findall(
                r'&quot;murl&quot;:&quot;([^&"]+)&quot;', html):
            u = raw.replace('\\\\/', '/').strip()
            low = u.lower().split('?', 1)[0]
            if low.endswith(('.jpg', '.jpeg', '.png', '.webp')):
                urls.append(u)
            if len(urls) >= max_n + 5:
                break
    except Exception:
        pass
    seen, blobs = set(), []
    for u in urls:
        if not u or u in seen:
            continue
        seen.add(u)
        try:
            got = ib.download_image(u)
        except Exception:
            got = None
        if not got:
            continue
        blob = got[0] if isinstance(got, tuple) else got
        if not blob:
            continue
        if _blob_embeddable(blob):
            blobs.append(bytes(blob))
        if len(blobs) >= max_n:
            break
    return blobs


def _deck_photos(subject, kind='person', max_n=3):
    """Up to max_n photo blobs for the subject, cached per process.
    Empty when photos are disabled, unresolvable, or too small."""
    subject = str(subject or '').strip()
    if not subject:
        return []
    key = f"{subject.lower()}|{str(kind or 'person').lower()}"
    if key in _PHOTO_CACHE:
        return list(_PHOTO_CACHE[key])
    blobs = []
    if _photos_enabled():
        try:
            blobs = _fetch_photos(subject, kind, max_n)
        except Exception:
            blobs = []
    _PHOTO_CACHE[key] = list(blobs)
    return list(blobs)


def _add_photo_fill(s, blob, x, y, w, h):
    """Place a photo to exactly fill the target box: stretch to the
    box, then crop the source so the aspect holds (cover-fill, never
    distorted)."""
    import io as _io
    pic = s.shapes.add_picture(_io.BytesIO(blob), x, y,
                               width=w, height=h)
    try:
        iw, ih = pic.image.size
        if iw and ih:
            tgt = float(w) / float(h)
            src = float(iw) / float(ih)
            if src > tgt:
                frac = 1.0 - (tgt / src)
                pic.crop_left = frac / 2
                pic.crop_right = frac / 2
            elif src < tgt:
                # Portrait into a landscape box: bias the window
                # toward the top of the image so faces stay in frame.
                frac = 1.0 - (src / tgt)
                pic.crop_top = frac * 0.22
                pic.crop_bottom = frac * 0.78
    except Exception:
        pass
    return pic


def _photo_moment(s, blob, alpha_pct):
    """Full-bleed photo plus scrim; a bad blob renders the slide
    text-only instead of failing the deck."""
    try:
        _add_photo_fill(s, blob, 0, 0, SW, SH)
    except Exception:
        return False
    _scrim(s, alpha_pct)
    return True


def _scrim(s, alpha_pct=64, color=GRAPHITE):
    """Graphite veil over a full-bleed photo so type stays readable
    (brand rule: scrim type over photographs)."""
    sh = rect(s, 0, 0, SW, SH, color)
    try:
        from pptx.oxml.ns import qn
        srgb = sh.fill.fore_color._xFill.find(qn('a:srgbClr'))
        if srgb is not None:
            srgb.append(srgb.makeelement(
                qn('a:alpha'), {'val': str(int(alpha_pct) * 1000)}))
    except Exception:
        pass
    return sh


class _Deck:'''

db = splice(db, '''class _Deck:''', PHOTO_LAYER.replace('\\\\', '\\'),
            'photo layer')

# ---- cover: photo + scrim under the type ----
db = splice(db, '''def _sl_cover(d, sl):
    s = d.new_slide(dark=True)
    dot(s, M, Inches(0.540))''', '''def _sl_cover(d, sl):
    s = d.new_slide(dark=True)
    if sl.get("_photo"):
        _photo_moment(s, sl["_photo"], 64)
    dot(s, M, Inches(0.540))''', 'cover photo')

# ---- hero (dark ground only): photo + scrim ----
db = splice(db, '''    s = d.new_slide(dark=not orchid, fill=ORCHID if orchid else None)
    d.chrome(s, sl.get("eyebrow") or "Signal", sl.get("title") or "",''',
            '''    s = d.new_slide(dark=not orchid, fill=ORCHID if orchid else None)
    if not orchid and sl.get("_photo"):
        _photo_moment(s, sl["_photo"], 68)
    d.chrome(s, sl.get("eyebrow") or "Signal", sl.get("title") or "",''',
            'hero photo')

# ---- hero_proof: photo + scrim ----
db = splice(db, '''def _sl_hero_proof(d, sl):
    s = d.new_slide(dark=True)
    d.chrome(s, sl.get("eyebrow") or "Proof", sl.get("title") or "",''',
            '''def _sl_hero_proof(d, sl):
    s = d.new_slide(dark=True)
    if sl.get("_photo"):
        _photo_moment(s, sl["_photo"], 70)
    d.chrome(s, sl.get("eyebrow") or "Proof", sl.get("title") or "",''',
            'hero_proof photo')

# ---- close: photo + scrim under the cards ----
db = splice(db, '''def _sl_close(d, sl):
    s = d.new_slide(dark=True)
    d.chrome(s, sl.get("eyebrow") or "Close", sl.get("title") or "",''',
            '''def _sl_close(d, sl):
    s = d.new_slide(dark=True)
    if sl.get("_photo"):
        _photo_moment(s, sl["_photo"], 70)
    d.chrome(s, sl.get("eyebrow") or "Close", sl.get("title") or "",''',
            'close photo')

# ---- entry point: resolve photos, assign the moments ----
db = splice(db, '''def render_insights_deck(plan, out_path, static_dir=None):
    """Render a slide-plan dict to a finished PPTX at out_path.
    Returns the number of slides rendered."""
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    lw, lk = _resolve_logos(static_dir)
    d = _Deck(prs, lw, lk)
    rendered = 0''',
            '''def render_insights_deck(plan, out_path, static_dir=None,
                        photo_subject='', photo_kind=''):
    """Render a slide-plan dict to a finished PPTX at out_path.
    Returns the number of slides rendered.

    Photography (2026-09-30): the plan's image_subject (or the
    caller's photo_subject fallback) resolves to real photos placed
    under a graphite scrim on the cover, the first hero or
    hero_proof, and the close - the manual deck grammar. A slide
    carrying photo: false opts out; no photos means text-only."""
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    lw, lk = _resolve_logos(static_dir)
    d = _Deck(prs, lw, lk)
    # Render from copies: the caller's plan is never mutated, so a
    # repeat render of the same plan is deterministic.
    _slides = [dict(sl) for sl in (plan.get("slides") or [])
               if isinstance(sl, dict)]
    try:
        _p_subj = str(plan.get("image_subject")
                      or photo_subject or '').strip()
        _p_kind = str(plan.get("image_kind")
                      or photo_kind or 'person').strip().lower()
        photos = _deck_photos(_p_subj, _p_kind) if _p_subj else []
        if photos:
            moments = []
            for want in (('cover',), ('hero', 'hero_proof'),
                         ('close',)):
                hit = next(
                    (sl for sl in _slides
                     if str(sl.get("type") or '').strip().lower()
                     in want and sl.get("photo") is not False), None)
                if hit is not None:
                    moments.append(hit)
            for i, sl in enumerate(moments):
                sl["_photo"] = photos[i % len(photos)]
    except Exception:
        pass
    rendered = 0
    for sl in _slides:
        fn = _RENDERERS.get(str(sl.get("type") or "").strip().lower())''', 'render entry photos')

db = splice(db, '''    rendered = 0
    for sl in _slides:
        fn = _RENDERERS.get(str(sl.get("type") or "").strip().lower())
    for sl in (plan.get("slides") or []):
        if not isinstance(sl, dict):
            continue
        fn = _RENDERERS.get(str(sl.get("type") or "").strip().lower())''',
            '''    rendered = 0
    for sl in _slides:
        fn = _RENDERERS.get(str(sl.get("type") or "").strip().lower())''',
            'render loop copies')

ast.parse(db)
DB.write_text(db, encoding='utf-8')
print('deck_builder.py: photography layer applied, ast clean')

# ===================== prometheus_analysis.py ======================
pma = PMA.read_text(encoding='utf-8')

pma = splice(
    pma,
    'Omit slides the data cannot carry (no live events means no '
    'second-screen slide; no avid cut means no avid tier tile). '
    'Never pad: a 14-slide deck that is all signal beats a 20-slide '
    'deck with filler.',
    'Omit slides the data cannot carry (no live events means no '
    'second-screen slide; no avid cut means no avid tier tile). '
    'Never pad: a 14-slide deck that is all signal beats a 20-slide '
    'deck with filler.\n\n'
    'ART DIRECTION. Set a top-level "image_subject": the one person, '
    'brand, or title the deck is about, spelled exactly as publicly '
    'known, and "image_kind": one of "person", "title", "brand". The '
    'renderer places real photography of that subject under a dark '
    'scrim on the cover, the first big statement page, and the '
    'close. Set "photo": false on any of those slides that must stay '
    'type-only.',
    'prompt art direction')

pma = splice(pma, '''    out = {
        'title': _messy_int_in_text(
            subject, scrub_user_text(str(plan.get('title') or ''))),
        'filename_stem': re.sub(
            r'[^A-Za-z0-9_]+', '_',
            str(plan.get('filename_stem') or '')).strip('_')[:60],
    }''', '''    out = {
        'title': _messy_int_in_text(
            subject, scrub_user_text(str(plan.get('title') or ''))),
        'filename_stem': re.sub(
            r'[^A-Za-z0-9_]+', '_',
            str(plan.get('filename_stem') or '')).strip('_')[:60],
    }
    # Art direction rides through (2026-09-30 deck photography).
    out['image_subject'] = scrub_user_text(
        str(plan.get('image_subject') or '')).strip()[:80]
    _ik = str(plan.get('image_kind') or '').strip().lower()
    if _ik not in ('person', 'title', 'brand'):
        _ik = ''
    out['image_kind'] = _ik''', 'enforce art fields')

pma = splice(pma, '''        cleaned = _clean_deck_value(subject, sl)
        cleaned['type'] = stype
        slides.append(cleaned)''', '''        cleaned = _clean_deck_value(subject, sl)
        cleaned['type'] = stype
        if sl.get('photo') is False:
            cleaned['photo'] = False
        slides.append(cleaned)''', 'photo opt-out survives clean')

ast.parse(pma)
PMA.write_text(pma, encoding='utf-8')
print('prometheus_analysis.py: art direction wired, ast clean')

# ============================ app.py ===============================
app = APP.read_text(encoding='utf-8')

app = splice(app, '''        deck_builder.render_insights_deck(plan, local,
                                          static_dir=static_dir)''',
             '''        deck_builder.render_insights_deck(
            plan, local, static_dir=static_dir,
            photo_subject=str(subject or p_meta.get('name') or ''))''',
             'app photo subject')

ast.parse(app)
APP.write_text(app, encoding='utf-8')
print('app.py: photo subject passed, ast clean')
