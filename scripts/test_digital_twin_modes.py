#!/usr/bin/env python3
"""Pilot Testing phase 2: the screening modes, the diagnostics, the
comparisons, the recruit block, the critics room, the hooks, the report,
the pricing and the partner surface (2026-10-28). Hermetic: FakeS3,
REGRESSION_TEST_MODE=1, no ffmpeg, no model.
"""
from __future__ import annotations

import json
import os
import re
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

os.environ.setdefault('REGRESSION_TEST_MODE', '1')
os.environ.setdefault('CORPUS_CATALOG', '0')

from migration import digital_twin_pilot as PI  # noqa: E402
from migration import digital_twin_modes as M  # noqa: E402
from migration import digital_twin_diagnostics as D  # noqa: E402
from migration import digital_twins as DT  # noqa: E402

sys.path.insert(0, HERE)
from test_digital_twin_pilot import _beals, _young, FakeS3  # noqa: E402

LOGLINE = 'A chef returns to her hometown diner after a scandal and has one summer to win the town back.'
SCRIPT = """INT. DINER - MORNING

MAE wipes the counter. The bell rings.

          MAE
     We're closed.

          RAY
     Not for me, you're not.

Ray sits. Mae does not look up.

EXT. MAIN STREET - DAY

Mae crosses the street. Everyone looks. A dog barks. The mayor pretends not to see her.

          MAE
     Say it to my face.

          MAYOR
     Welcome home, Mae.

INT. DINER - NIGHT

The room is full. Ray raises a glass.

          RAY
     To second helpings.
"""
TRANSCRIPT = [{'t': 10, 'end': 14, 'text': 'We are closed', 'speaker': 'Mae'}, {'t': 30, 'end': 34, 'text': 'Not for me', 'speaker': 'Ray'},
              {'t': 300, 'end': 304, 'text': 'Say it to my face', 'speaker': 'Mae'}, {'t': 500, 'end': 504, 'text': 'To second helpings', 'speaker': 'Ray'},
              {'t': 540, 'end': 544, 'text': 'Welcome home', 'speaker': 'Mayor'}]
BANNED = ('synth', 'synthetic', 'modeled', 'estimated', 'pipeline', 'claude', 'hostmap', 's3://', 'worker', 'persona', 'agent', '\u2014', '\u2013')


def _pilot(twin=None, **kw):
    twin = twin or _beals()
    sess = PI.open_session([twin], title=kw.pop('title', 'Pilot'), in_room=kw.pop('in_room', 12), **{k: v for k, v in kw.items() if k in ('mode', 'logline', 'text', 'options', 'rough_cut', 'critics_room')})
    return twin, sess


def _no_banned(obj, where=''):
    blob = json.dumps(obj, default=str).lower()
    hits = [b for b in BANNED if b in blob]
    return hits


class Modes(unittest.TestCase):
    def test_each_mode_opens_with_its_battery_and_norms_cell(self):
        twin = _beals()
        for mode in ('pilot', 'feature', 'episode', 'binge', 'trailer', 'spot', 'title', 'key_art', 'script', 'concept'):
            kw = {}
            if mode in ('title', 'key_art'):
                kw['options'] = ['The Diner', 'Hometown Heat', 'Second Helping']
            if mode in ('script', 'concept'):
                kw['text'] = SCRIPT if mode == 'script' else LOGLINE
            sess = PI.open_session([twin], title=f"{mode} test", in_room=8, mode=mode, logline=LOGLINE, **kw)
            self.assertEqual(sess['mode'], mode)
            self.assertEqual(PI.charge_due(sess), DT.MODE_PRICE_KEY[mode])
            PI.prepare_watch(sess, {'jb': twin}, duration_s=(95 if mode in ('trailer', 'spot') else 600 if mode in ('pilot', 'feature', 'episode', 'binge') else 0), title=f"{mode} cut")
            self.assertEqual(sess['status'], 'ready')
            fmt = sess['format']
            self.assertIn(fmt, ('pilot', 'feature', 'episode', 'trailer', 'title', 'key_art', 'script', 'concept'))
            bat = sess['scorecard']['battery']
            self.assertTrue(bat['questions'], mode)
            self.assertEqual(bat['format'], fmt)
            norms = sess['scorecard']['norms']
            self.assertEqual(norms['format'], fmt)
            self.assertTrue(norms['all'].get('total_positive') or norms['all'].get('tune_in'), mode)
            self.assertIn('mode_read', sess)
            self.assertEqual(sess['mode_read'].get('mode'), mode)

    def test_trailer_lift_computed_and_in_range(self):
        twin, sess = _pilot(mode='trailer', logline=LOGLINE, in_room=10)
        PI.prepare_watch(sess, {'jb': twin}, duration_s=95, title='Hurry Up Tomorrow trailer')
        r = sess['mode_read']
        for k in ('interest_before_pct', 'interest_after_pct', 'definite_interest_lift_pp', 'would_watch_pct', 'who_for', 'understood', 'expect', 'sold_it', 'quadrants', 'read'):
            self.assertIn(k, r)
        self.assertGreaterEqual(r['interest_before_pct'], 0.0)
        self.assertLessEqual(r['interest_before_pct'], 100.0)
        self.assertAlmostEqual(r['definite_interest_lift_pp'], round(r['interest_after_pct'] - r['interest_before_pct'], 1), places=1)
        self.assertGreaterEqual(r['definite_interest_lift_pp'], -100.0)
        self.assertLessEqual(r['definite_interest_lift_pp'], 100.0)
        self.assertEqual(len(r['understood']), 10)
        self.assertTrue(r['who_for']['line'].startswith('Reads as a title for'))
        self.assertEqual(len(r['quadrants']), 4)
        self.assertEqual(sess['format'], 'trailer')
        self.assertEqual(PI.play_quantity(sess), 1)
        sess['assets'] = [{'id': 'a1', 's3_key': 'x'}, {'id': 'a2', 's3_key': 'y'}, {'id': 'a3', 's3_key': 'z'}, {'id': 'a4', 's3_key': 'w'}]
        self.assertEqual(PI.play_quantity(sess), 3)
        self.assertFalse(_no_banned(r))

    def test_trailer_second_asset_same_room_and_compare(self):
        twin, sess = _pilot(mode='trailer', logline=LOGLINE, in_room=8)
        PI.prepare_watch(sess, {'jb': twin}, duration_s=60, title='Trailer A')
        rec = M.prepare_extra_asset(sess, {'id': 'b', 'slot': 1, 'label': 'Trailer B', 'title': 'Trailer B', 'runtime_s': 45}, {'jb': twin})
        self.assertIn('trailer', rec)
        self.assertEqual(len(rec['dial']), len(rec['t']))
        self.assertEqual(set(rec['seat_dials'].keys()), set(sess['seat_dials'].keys()))
        M.finish_mode_read(sess, {'jb': twin})
        cmp_ = sess['mode_read'].get('compare')
        self.assertTrue(cmp_)
        self.assertEqual(len(cmp_['ranked']), 2)
        self.assertIn(cmp_['winner'], ('Trailer A', 'Trailer B'))

    def test_title_test_ranks_six_options_with_reasons(self):
        opts = ['The Diner', 'Hometown Heat', 'Second Helping', 'What Happened to Mae', 'Service', 'Burned']
        twin, sess = _pilot(mode='title', logline=LOGLINE, options=opts, in_room=12)
        PI.prepare_watch(sess, {'jb': twin}, title='Six titles')
        r = sess['mode_read']
        self.assertEqual(r['mode'], 'title')
        self.assertEqual(len(r['options']), 6)
        self.assertEqual([o['rank'] for o in r['options']], [1, 2, 3, 4, 5, 6])
        tot = sum(o['preference_pct'] for o in r['options'])
        self.assertAlmostEqual(tot, 100.0, delta=0.6)
        self.assertTrue(all(o['preference_pct'] > 0 for o in r['options']))
        self.assertTrue(r['options'][0]['reasons'], 'the winner carries reasons')
        self.assertTrue(all('fit_pct' in o and 'unique_pct' in o for o in r['options']))
        self.assertEqual(len(r['quadrants']), 4)
        self.assertEqual(len(sess['scenes']), 6)
        self.assertIn(r['winner_label'], opts)
        self.assertTrue(r['read'])

    def test_key_art_test_has_attention_and_lift(self):
        twin, sess = _pilot(mode='key_art', logline=LOGLINE, options=[{'label': 'Art A', 'caption': 'A woman in a diner doorway at night, the title in red'},
                                                                        {'label': 'Art B', 'caption': 'A bright sunny street, the whole ensemble smiling'}], in_room=10)
        PI.prepare_watch(sess, {'jb': twin}, title='Two pieces of art')
        r = sess['mode_read']
        self.assertEqual(r['mode'], 'key_art')
        for o in r['options']:
            self.assertIn('attention_pct', o)
            self.assertIn('interest_lift_pp', o)
        self.assertTrue(M.accept_image_name('poster.png'))
        self.assertFalse(M.accept_image_name('poster.mp4'))

    def test_script_scene_split_on_sluglines(self):
        sc = M.script_scenes(SCRIPT)
        self.assertEqual(sc['scene_count'], 3)
        self.assertEqual(sc['scenes'][0]['slug'], 'INT. DINER - MORNING')
        self.assertEqual(sc['scenes'][1]['slug'], 'EXT. MAIN STREET - DAY')
        self.assertIn('Mae', sc['cast'])
        self.assertIn('Ray', sc['cast'])
        self.assertTrue(sc['transcript'])
        self.assertGreater(sc['runtime_s'], 15)
        twin, sess = _pilot(mode='script', text=SCRIPT, in_room=8)
        PI.prepare_watch(sess, {'jb': twin}, title='Second Helping script')
        self.assertEqual(sess['format'], 'script')
        self.assertEqual(len(sess['scenes']), 3)
        r = sess['mode_read']
        self.assertEqual(len(r['scenes']), 3)
        for row in r['scenes']:
            self.assertIn('story', row['notes'])
            self.assertIn('structure', row['notes'])
            self.assertIn('character', row['notes'])
        self.assertIn('script_story', r['norm_metrics'])
        # no sluglines: paragraphs are the scenes
        self.assertGreaterEqual(M.script_scenes('A short pitch.\n\nAnother paragraph.')['scene_count'], 2)
        self.assertTrue(M.accept_text_name('script.pdf'))
        self.assertEqual(M.text_from_upload(b'hello world', 'a.txt'), 'hello world')

    def test_concept_read_has_battery(self):
        twin, sess = _pilot(mode='concept', text=LOGLINE + ' Her ex runs the bank that holds the note.', in_room=10)
        PI.prepare_watch(sess, {'jb': twin}, title='Second Helping')
        r = sess['mode_read']
        for k in ('definitely_interested_pct', 'unique_pct', 'who_for', 'would_watch_pct', 'strongest', 'doubt', 'quadrants', 'read'):
            self.assertIn(k, r)
        self.assertEqual(r['strongest'], sorted(r['strongest'], key=lambda x: -x['pct']))
        self.assertEqual(r['doubt'], sorted(r['doubt'], key=lambda x: -x['pct']))
        self.assertEqual(sess['format'], 'concept')
        self.assertFalse(_no_banned(r))

    def test_episode_knows_the_prior_cut_and_binge_builds(self):
        twin, a = _pilot(title='Pilot', in_room=8)
        PI.prepare_watch(a, {'jb': twin}, duration_s=300, title='Episode 1')
        b = PI.open_session([twin], title='Episode 2', in_room=8, mode='episode')
        b['prior_read'] = {'id': a['id'], 'title': a['title'], 'scorecard': a['scorecard'], 'seat_dials': a['seat_dials'], 'frame': 'the room watched episode 1'}
        PI.prepare_watch(b, {'jb': twin}, duration_s=300, title='Episode 2')
        r = b['mode_read']
        self.assertEqual(r['mode'], 'episode')
        self.assertIsNotNone(r['commitment_pct'])
        self.assertEqual(r['prior']['session_id'], a['id'])
        self.assertIsNotNone(r['tune_in_delta_pp'])
        reads = [{'label': 'Ep 1', 'scorecard': a['scorecard'], 'dial': a['dial']['value']}, {'label': 'Ep 2', 'scorecard': b['scorecard'], 'dial': b['dial']['value']}]
        bg = M.binge_read(b, reads)
        self.assertEqual(len(bg['episodes']), 2)
        self.assertEqual(len(bg['cumulative_tune_in']), 2)
        self.assertTrue(bg['read'])

    def test_rough_cut_skips_marketing_locked_runs_it(self):
        twin, rough = _pilot(title='Rough', in_room=8, rough_cut=True)
        PI.prepare_watch(rough, {'jb': twin}, duration_s=300, title='Rough cut')
        self.assertNotIn('marketing_read', rough)
        twin, locked = _pilot(title='Locked', in_room=8)
        PI.prepare_watch(locked, {'jb': twin}, duration_s=300, title='Locked cut')
        mk = locked['marketing_read']
        self.assertIn('who_for', mk)
        self.assertIn('trailer_moments', mk)
        self.assertTrue(mk['read'])

    def test_set_stimulus_and_options_patch(self):
        twin, sess = _pilot(mode='title', in_room=6)
        ch = PI.set_options(sess, {'options': ['A', 'B', 'C', 'D', 'E', 'F', 'G'], 'logline': LOGLINE, 'cast': ['Mae', 'Ray'], 'critics_room': True})
        self.assertEqual(ch['options'], 6)
        self.assertTrue(ch['logline'])
        self.assertEqual(ch['cast'], 2)
        self.assertTrue(sess['critics_room'])
        self.assertEqual([o['id'] for o in sess['options']], list('abcdef'))

    def test_graded_shares_are_not_unanimous(self):
        rows = M.ordered_probs(60.0, M.CUTS['interest4'], list('abcd'))
        self.assertAlmostEqual(sum(rows.values()), 1.0, places=6)
        self.assertGreater(rows['a'], 0.4)
        self.assertGreater(rows['b'], 0.1)
        young = _young()
        sess = PI.open_session([young], title='Young', in_room=12, mode='trailer', logline=LOGLINE)
        PI.prepare_watch(sess, {'ym': young}, duration_s=80, title='An action trailer')
        r = sess['mode_read']
        self.assertGreater(r['interest_after_pct'], 1.0)
        self.assertLess(r['interest_after_pct'], 99.0)


class Diagnostics(unittest.TestCase):
    def _pilot_with_cast(self):
        twin = _beals()
        sess = PI.open_session([twin], title='Pilot', in_room=12)
        sess['cast'] = ['Mae', 'Ray', 'Lorelai']
        PI.prepare_watch(sess, {'jb': twin}, duration_s=600, title='A comedy pilot', transcript=TRANSCRIPT)
        return twin, sess

    def test_character_table_from_transcript_with_norms(self):
        twin, sess = self._pilot_with_cast()
        d = sess['diagnostics']['characters']
        names = [c['name'] for c in d['characters']]
        self.assertIn('Mae', names)
        self.assertIn('Ray', names)
        self.assertIn('Mayor', names)
        lead = d['characters'][0]
        self.assertEqual(lead['role'], 'lead')
        self.assertIsNotNone(lead['excellent_pct'])
        self.assertEqual(lead['norm'], 50.0)
        self.assertEqual(d['characters'][1]['norm'], 45.0)
        self.assertTrue(d['relationship'])
        self.assertTrue(d['follow'])
        self.assertTrue(d['read'])
        for c in d['characters']:
            self.assertGreater(c['excellent_pct'], 0.0)
            self.assertLess(c['excellent_pct'], 100.0)
        self.assertFalse(_no_banned(d))

    def test_pace_confusion_ending_present(self):
        twin, sess = self._pilot_with_cast()
        p = sess['diagnostics']['pace']
        self.assertEqual([a['act'] for a in p['acts']], ['beginning', 'middle', 'turn', 'ending'])
        for a in p['acts']:
            self.assertAlmostEqual((a['too_slow_pct'] or 0) + (a['just_right_pct'] or 0) + (a['too_fast_pct'] or 0), 100.0, delta=0.6)
        self.assertIn('confusion_pct', p)
        self.assertIn('confusions', p)
        self.assertIsNotNone(p['beginning']['excellent_pct'])
        self.assertIsNotNone(p['ending']['excellent_pct'])
        self.assertTrue(p['running_time_read'])
        self.assertTrue(p['read'])

    def test_scenes_liked_best_and_least_with_counts_and_quotes(self):
        twin, sess = self._pilot_with_cast()
        s = sess['diagnostics']['scenes']
        self.assertTrue(s['liked_best'])
        self.assertTrue(s['liked_least'])
        top = s['liked_best'][0]
        self.assertGreaterEqual(top['mentions'], 1)
        self.assertTrue(top['quotes'])
        self.assertTrue(top['quotes'][0]['line'])
        self.assertEqual(s['liked_best'], sorted(s['liked_best'], key=lambda x: -x['mentions']))
        self.assertIn('not_recommending', s)
        self.assertIn('similar_titles', s)
        self.assertTrue(s['appeal']['line'].startswith('Plays as a title for'))

    def test_components_with_norms(self):
        twin, sess = self._pilot_with_cast()
        c = sess['diagnostics']['components']
        keys = [r['key'] for r in c['rows']]
        self.assertEqual(keys, ['humor', 'story', 'beginning', 'ending', 'music', 'settings', 'action', 'suspense'])
        humor = c['rows'][0]
        self.assertEqual(humor['norm'], 40.0)
        self.assertIsNotNone(humor['excellent_pct'])
        nm = PI.norm_metrics(sess)
        self.assertIn('component_humor', nm)
        self.assertIn('lead_character', nm)
        self.assertIn('component_pace_just_right', nm)

    def test_diagnostics_lines_clean(self):
        twin, sess = self._pilot_with_cast()
        lines = D.diagnostics_lines(sess)
        self.assertGreaterEqual(len(lines), 3)
        for l in lines:
            self.assertFalse(_no_banned(l), l)
            self.assertNotIn('\u2014', l)


class Comparison(unittest.TestCase):
    def _pair(self):
        from migration import digital_twin_compare as C
        twin = _beals()
        a = PI.open_session([twin], title='Pilot', in_room=10)
        PI.prepare_watch(a, {'jb': twin}, duration_s=600, title='A comedy pilot', transcript=TRANSCRIPT)
        b = C.open_recut(a, [twin], created_by='jenna')
        PI.prepare_watch(b, {'jb': twin}, duration_s=540, title='A comedy pilot', transcript=[t for t in TRANSCRIPT if t['t'] != 300])
        return twin, a, b

    def test_recut_keeps_seats_and_personas_and_aligns_scenes(self):
        from migration import digital_twin_compare as C
        twin, a, b = self._pair()
        self.assertEqual(b['lineage']['parent'], a['id'])
        self.assertEqual(b['cut_label'], 'Cut B')
        self.assertEqual([r['name'] for r in a['group']['roster']], [r['name'] for r in b['group']['roster']])
        self.assertEqual([r.get('seat') for r in a['group']['roster']], [r.get('seat') for r in b['group']['roster']])
        self.assertEqual(PI.charge_due(b), 'digital_twin_recut')
        self.assertTrue(b['prior_read']['frame'].startswith('The room watched'))
        self.assertEqual(a['recuts'][0]['id'], b['id'])
        cmp_ = C.compare_recut(a, b)
        self.assertEqual(len(cmp_['dial']['a']), len(cmp_['dial']['b']))
        self.assertEqual(len(cmp_['dial']['t']), len(cmp_['dial']['b']))
        self.assertTrue(cmp_['scorecard_deltas'])
        self.assertTrue(cmp_['scene_deltas'])
        self.assertTrue(cmp_['act_deltas'])
        self.assertGreaterEqual(cmp_['aligned_by_words'], 1)
        self.assertTrue(cmp_['same_room'])
        self.assertTrue(cmp_['read'])
        self.assertFalse(_no_banned(cmp_))
        al = C.align_scenes([{'id': 0, 't_start': 0, 't_end': 10, 'dialogue': 'we are closed not for me you are not'}, {'id': 1, 't_start': 10, 't_end': 20, 'dialogue': 'say it to my face'}],
                            [{'id': 0, 't_start': 0, 't_end': 9, 'dialogue': 'we are closed, not for me you are not'}], 20, 9)
        self.assertEqual(al['pairs'][0]['a_id'], 0)
        self.assertEqual(al['pairs'][0]['method'], 'words')
        self.assertEqual(al['trimmed'], [1])

    def test_two_audiences_overlay_returns_two_lines(self):
        twin, young = _beals(), _young()
        sess = PI.open_session([twin, young], title='Two', in_room=12)
        PI.prepare_watch(sess, {'jb': twin, 'ym': young}, duration_s=300, title='An action pilot')
        ta = sess.get('two_audiences')
        self.assertTrue(ta)
        self.assertEqual(len(ta['lines']), 2)
        for line in ta['lines']:
            self.assertEqual(len(line['dial']), len(ta['t']))
            self.assertIsNotNone(line['total_positive_pct'])
        self.assertIn('splits', ta)
        self.assertTrue(ta['read'])
        self.assertIn('two_audiences', PI.public_session(sess))

    def test_quota_block_reseats_to_the_quotas(self):
        from migration import digital_twin_compare as C
        twin = _beals()
        sess = PI.open_session([twin], title='Quota', in_room=12)
        rec = C.apply_quotas(sess, [twin], {'under_35_pct': 65, 'women_pct': 50, 'fans_of_lead_pct': 40, 'must_have_seen': 2, 'lead': 'Jennifer Beals', 'comps': ['Flashdance', 'The L Word', 'Taken']})
        roster = sess['group']['roster']
        self.assertEqual(len(roster), 12)
        women = sum(1 for r in roster if r['gender'] == 'F')
        under = sum(1 for r in roster if int(r['age']) < 35)
        self.assertEqual(women, 6)
        self.assertIn(under, (7, 8))
        self.assertEqual(sum(1 for r in roster if r.get('fan_of_lead')), 5)
        self.assertTrue(all(len(r.get('seen_titles') or []) >= 2 for r in roster))
        self.assertIn('65% under 35', rec['concept'])
        self.assertIn('must have seen 2 of', rec['concept'])
        self.assertGreater(rec['invited'], 12)
        self.assertNotEqual(rec['invited'] % 10, 0)
        self.assertIn('invited for every seat', rec['invite_ratio'])
        self.assertEqual(len(rec['quotas_met']), 3)
        PI.prepare_watch(sess, {'jb': twin}, duration_s=400, title='A pilot')
        self.assertIn('walkouts', sess['recruit'])
        self.assertIn('stats_line', sess['recruit'])
        self.assertIn('seated', sess['recruit']['stats_line'])
        # a room that has played is left alone
        PI.mark_play_paid(sess)
        before = [r['name'] for r in sess['group']['roster']]
        C.apply_quotas(sess, [twin], {'women_pct': 100})
        self.assertEqual(before, [r['name'] for r in sess['group']['roster']])

    def test_comps_attach_from_a_fake_catalog(self):
        from migration import digital_twin_compare as C
        twin = _beals()
        sess = PI.open_session([twin], title='Comps', in_room=8)
        sess['comps'] = [{'title': 'Flashdance', 'total_positive_pct': 78.4, 'definitely_recommend_pct': 61.2}, 'The L Word']
        PI.prepare_watch(sess, {'jb': twin}, duration_s=300, title='A drama pilot')
        payload = {'meta': twin['meta'], 'runs': [{'id': 'r1', 'kind': 'greenlight', 'inputs': {'title': 'Second Helping'}, 'outputs': {'title': 'Second Helping', 'odds_pct': 64.2, 'expected': 120345}}],
                   'pilot_sessions': [sess, {'id': 'older', 'title': 'Last season', 'scorecard': {'battery': {'headline': {'total_positive_pct': 70.3, 'definitely_recommend_pct': 48.9, 'excellent_pct': 31.1}}}}]}
        other = {'meta': {'name': 'Other', 'key': 'k2'}, 'pilot_sessions': [{'id': 'o1', 'title': 'Other pilot', 'scorecard': {'battery': {'headline': {'total_positive_pct': 60.1, 'definitely_recommend_pct': 40.2, 'excellent_pct': 20.3}}}}]}
        catalog = [{'key': 'k2', 'name': 'Other', 'has_pilot': True}]
        comps = C.comp_set(payload, sess, catalog_runs=catalog, load_twin=lambda k: other)
        titles = [r['title'] for r in comps['rows']]
        self.assertIn('Flashdance', titles)
        self.assertIn('The L Word', titles)
        self.assertNotIn('Second Helping', titles)      # the retired Greenlight run never feeds the comps (2026-10-28)
        self.assertFalse(any('Greenlight' in str(r.get('source')) for r in comps['rows']))
        self.assertIn('Last season', titles)
        self.assertIn('Other pilot', titles)
        self.assertIn('definitely_recommend_pct', comps['versus'])
        self.assertTrue(comps['read'].startswith(('Lands between', 'Above every comp', 'Below every comp')), comps['read'])
        self.assertEqual(comps['versus']['definitely_recommend_pct']['of'], 4)
        self.assertTrue(all('versus' in r for r in comps['rows']))

    def test_critics_room_seats_reviewer_personas(self):
        from migration import digital_twin_compare as C
        twin = _beals()
        sess = PI.open_session([twin], title='Critics', in_room=8, critics_room=True)
        C.seat_critics(sess)
        for r in sess['group']['roster']:
            self.assertIn(r['occupation'], ('Film and TV critic', 'Entertainment journalist'))
            self.assertIn(r['reviewer_lens'], C.REVIEWER_LENS)
            self.assertTrue(r.get('outlet'))
        PI.prepare_watch(sess, {'jb': twin}, duration_s=400, title='A drama pilot')
        cr = sess['critics']
        self.assertGreater(cr['critic_score'], 0)
        self.assertLess(cr['critic_score'], 100)
        self.assertEqual(len(cr['range']), 2)
        self.assertTrue(cr['awards_read'])
        self.assertIn('critic_score', PI.norm_metrics(sess))
        self.assertEqual(PI.charge_due(sess), 'digital_twin_pilot')
        self.assertEqual(C.find_critics_twin([{'name': 'Gilmore Girls', 'key': 'a'}, {'name': 'TV Critics and Journalists', 'key': 'b'}])['key'], 'b')
        self.assertIsNone(C.find_critics_twin([{'name': 'Gilmore Girls', 'key': 'a'}]))


class Hooks(unittest.TestCase):
    def test_concept_test_seats_two_twins_in_one_room(self):
        """Two twins on one idea is a Concept test with both in the room (the
        compare-on-an-idea Greenlight run is gone, 2026-10-28)."""
        from migration import digital_twin_compare as C
        twin, young = _beals(), _young()
        sess = PI.open_session([twin, young], title='Second Helping', in_room=12, mode='concept', text=LOGLINE)
        PI.prepare_watch(sess, {'jb': twin, 'ym': young}, title='Second Helping')
        self.assertEqual(sess['mode'], 'concept')
        self.assertIsNotNone(sess['mode_read']['definitely_interested_pct'])
        ta = sess.get('two_audiences')
        self.assertTrue(ta and len(ta['lines']) == 2)
        self.assertEqual(PI.charge_due(sess), 'digital_twin_concept')
        self.assertFalse(hasattr(C, 'concept_inputs_from_run'))
        self.assertFalse(hasattr(C, 'attach_concept_to_run'))

    def test_talent_attach_by_quadrant_with_watch_with(self):
        from migration import digital_twin_compare as C
        twin = _beals()
        twin['meta']['size_us'] = 9339205
        sess = PI.open_session([twin], title='Pilot', in_room=10)
        PI.prepare_watch(sess, {'jb': twin}, duration_s=300, title='A drama pilot')
        rec = C.talent_attach(sess, twin, 'Jennifer Beals', 'lead')
        self.assertEqual(rec['talent'], 'Jennifer Beals')
        self.assertEqual(len(rec['by_quadrant']), 4)
        self.assertTrue(any(q.get('reach_lift_pct') for q in rec['by_quadrant']))
        self.assertTrue(rec['watch_with'])
        self.assertAlmostEqual(sum(r['pct'] for r in rec['watch_with']), 100.0, delta=0.6)
        self.assertIn('watch it with', rec['read'])
        self.assertEqual(sess['talent_attach'][-1]['talent'], 'Jennifer Beals')

    def test_activation_brief_has_the_spend_engine_handoff_shape(self):
        from migration import digital_twin_compare as C
        from migration import digital_twin_spend as SP
        twin = _beals()
        twin['identity']['platforms'] = [{'label': 'Netflix', 'reach_pct': 61.2}, {'label': 'Hulu', 'reach_pct': 44.3}, {'label': 'YouTube', 'reach_pct': 71.8}]
        sess = PI.open_session([twin], title='Pilot', in_room=12)
        PI.prepare_watch(sess, {'jb': twin}, duration_s=300, title='A drama pilot')
        brief = C.activation_brief(sess, twin)
        hand = SP.handoff_package(twin, {'channels': brief['channels'], 'goal': brief['goal'], 'window': {'days': 42}})
        for k in ('campaign_name', 'text', 'channels', 'goal'):
            self.assertIn(k, brief)
            self.assertIn(k, hand)
        self.assertEqual([c['channel'] for c in brief['channels']], ['YouTube', 'Netflix', 'Hulu'])
        for c in brief['channels']:
            self.assertNotEqual(int(c['spend']) % 10, 0)
        self.assertTrue(brief['segments'])
        self.assertTrue(brief['launch_prompt'].startswith('Run Launch Lab on the'))
        self.assertIn('Daily tracking on', brief['text'])
        self.assertFalse(_no_banned(brief))


class Report(unittest.TestCase):
    def _sess(self):
        from migration import digital_twin_compare as C
        twin = _beals()
        sess = PI.open_session([twin], title='Second Helping pilot', in_room=12)
        sess['cast'] = ['Mae', 'Ray']
        C.apply_quotas(sess, [twin], {'under_35_pct': 60, 'women_pct': 55, 'fans_of_lead_pct': 40, 'must_have_seen': 1, 'lead': 'Mae'})
        PI.prepare_watch(sess, {'jb': twin}, duration_s=600, title='A comedy pilot', transcript=TRANSCRIPT)
        return twin, sess

    def test_report_pages_tldr_first_and_clean(self):
        from migration import digital_twin_report as R
        twin, sess = self._sess()
        pages = R.report_pages({'meta': twin['meta']}, sess)
        self.assertEqual(pages[0]['kind'], 'cover')
        self.assertEqual(pages[1]['kind'], 'bullets')
        self.assertTrue(pages[1]['header'].startswith('What we found'))
        kinds = [p['kind'] for p in pages]
        self.assertIn('image', kinds)
        self.assertGreaterEqual(kinds.count('table'), 5)
        blob = json.dumps([{k: v for k, v in p.items() if k != 'image'} for p in pages], default=str).lower()
        for tok in ('\u2014', '\u2013', 'door', 'synth', 'persona', 'claude', 'screen engine', 'nrg', 'marketcast', 'fishdog', 'pipeline', 'hostmap'):
            self.assertNotIn(tok, blob, tok)
        for p in pages:
            self.assertLessEqual(len(p['header'].split()), 16, p['header'])
        png = R.dial_chart_png(sess)
        self.assertTrue(png and png[:8] == b'\x89PNG\r\n\x1a\n')

    def test_report_deck_passes_brand_lint(self):
        from migration import digital_twin_report as R
        from migration import design_skill as ds
        import tempfile
        twin, sess = self._sess()
        deck = R.build_report_deck({'meta': twin['meta']}, sess)
        self.assertGreater(len(deck), 20000)
        pdf = R.build_report_pdf({'meta': twin['meta']}, sess)
        self.assertTrue(pdf[:4] == b'%PDF')
        d = tempfile.mkdtemp()
        p = os.path.join(d, 'report.pptx')
        with open(p, 'wb') as fh:
            fh.write(deck)
        try:
            from pptx import Presentation
        except Exception:
            self.skipTest('python-pptx not installed')
        prs = Presentation(p)
        texts = []
        per_slide = []
        for sl in prs.slides:
            words = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            per_slide.append(' '.join(words).lower())
            texts.extend(words)
        blob = ' '.join(texts).lower()
        self.assertIn('what we found', blob)
        self.assertIn('what we found', per_slide[1], 'the TLDR is the page after the cover')
        for tok in ('\u2014', '\u2013', ' door', 'synth', 'persona', 'claude', 'screen engine'):
            self.assertNotIn(tok, blob, tok)
        if ds.scripts_dir() is not None:
            rep = ds.brand_lint(p)
            step = next((s for s in rep.steps if s[0] == 'brand_lint'), None)
            self.assertIsNotNone(step)
            self.assertTrue(step[1], step[2][-800:])


class PartnerSurface(unittest.TestCase):
    def test_scrub_and_prompt_parsing(self):
        import digital_twins_partner_api as P
        self.assertEqual(P.scrub_text('Definitely recommend 54.5%, comedy pilots usually land near 50%.'), 'Definitely recommend 54.5%, comedy pilots usually land near 50%.')
        for bad in ('the synth engine read it', 'Claude scored the room', 'the worker restarted', 'hostmap gap', 'the panel of twelve', 'the cohort skews', 'the pipeline ran', 'estimated at 40%'):
            self.assertEqual(P.scrub_text(bad), '', bad)
        self.assertEqual(P.mode_from_prompt('Test this trailer for Hurry Up Tomorrow on the Gilmore Girls twin'), 'trailer')
        self.assertEqual(P.mode_from_prompt('Which of these six titles works for the Beals audience?'), 'title')
        self.assertEqual(P.mode_from_prompt('Read this script with the room'), 'script')
        self.assertEqual(P.mode_from_prompt('Screen the pilot'), 'pilot')
        self.assertEqual(P.mode_from_prompt('anything', 'key art'), 'key_art')
        cat = [{'key': 'a', 'name': 'Gilmore Girls Hulu viewers', 'subject': 'Gilmore Girls'}, {'key': 'b', 'name': 'Jennifer Beals', 'subject': 'Jennifer Beals'}]
        self.assertEqual(P.twin_from_prompt('Test this trailer on the Jennifer Beals twin', cat)['key'], 'b')
        self.assertEqual(P.twin_from_prompt('a trailer for gilmore girls', cat)['key'], 'a')
        self.assertIsNone(P.twin_from_prompt('nothing here', cat))
        self.assertEqual(P.twin_from_prompt('x', cat, explicit='Jennifer Beals')['key'], 'b')

    def test_result_allowlist_scrubs_internal_terms(self):
        import digital_twins_partner_api as P
        twin = _beals()
        sess = PI.open_session([twin], title='Pilot', in_room=10)
        sess['cast'] = ['Mae', 'Ray']
        PI.prepare_watch(sess, {'jb': twin}, duration_s=600, title='A comedy pilot', transcript=TRANSCRIPT)
        sess['screening_id'] = 'scr_test'
        PI.mark_play_paid(sess)
        PI.stop_session(sess)
        sess['scorecard']['norm_lines'].append('the synth engine says hello')
        sess['events'][0]['line'] = 'Claude thought this was funny'
        out = P._result(sess, 'Jennifer Beals')
        self.assertEqual(out['status'], 'complete')
        blob = json.dumps(out, default=str).lower()
        for tok in P._BANNED:
            self.assertNotIn(tok, blob, tok)
        for k in ('scorecard', 'norms', 'quadrants', 'dial', 'reactions', 'characters', 'pace', 'scenes', 'components', 'exports'):
            self.assertIn(k, out)
        for forbidden in ('watch', 'seat_dials', 'prepare_job_id', 'media', 'group', 'created_by', 'twin_keys'):
            self.assertNotIn(forbidden, out)
        self.assertNotIn('the synth engine says hello', json.dumps(out))

    def test_blueprint_routes_through_flask(self):
        try:
            import flask  # noqa: F401
        except Exception:
            self.skipTest('flask not installed in this interpreter')
        import types
        import digital_twins_api as A
        import digital_twins_partner_api as P
        from migration import digital_twin_synthesis as SYN
        from prometheus.host import host
        s3 = FakeS3()
        s3.list_objects_v2 = lambda **kw: {'Contents': [{'Key': k, 'LastModified': '2026'} for k in list(s3.objs) if k.startswith(kw.get('Prefix', ''))]}
        twin = _beals()
        twin['meta']['created_by'] = 'jenna'
        key = 'digital-twins/jenna/Jennifer_Beals_20261028_120000_abc123def0.json.gz'
        SYN.save_twin(s3, key, twin)
        user = {'username': 'acme', 'role': 'user', 'has_digital_twins_iq_access': True, '_auth_via': 'api_key', 'api_key_id': 'k1', 'api_key_scopes': ['ask', 'build']}
        charges = []
        fake_app = types.SimpleNamespace(s3_client=s3, get_current_user=lambda: user, _normalize_role=lambda r: r,
                                         consume_credit=lambda *a, **k: charges.append(k) or True)
        sys.modules['app'] = fake_app
        host.bind(gate=lambda allow_api_key=False: (user, None))
        app = flask.Flask('t')
        app.secret_key = 'x'
        A.register(app)
        self.assertIn('digital_twins_partner_api', app.blueprints)
        c = app.test_client()
        r = c.post('/api/v1/screenings', json={'prompt': 'Run a concept test on the Jennifer Beals twin: ' + LOGLINE, 'mode': 'concept', 'text': LOGLINE})
        self.assertIn(r.status_code, (200, 202), r.get_data(as_text=True)[:300])
        body = r.get_json()
        self.assertTrue(body['success'])
        self.assertEqual(body['mode'], 'concept')
        self.assertTrue(charges, 'the concept price was taken')
        self.assertEqual(charges[-1]['pull_type'], 'Digital Twins IQ Concept Test')
        sid = body['screening_id']
        r2 = c.get(f'/api/v1/screenings/{sid}')
        self.assertEqual(r2.status_code, 200)
        out = r2.get_json()
        self.assertEqual(out['status'], 'complete')
        self.assertIn('scorecard', out)
        self.assertIn('mode_read', out)
        blob = json.dumps(out).lower()
        for tok in P._BANNED:
            self.assertNotIn(tok, blob, tok)
        r3 = c.get(f'/api/v1/screenings/{sid}/dial.png')
        self.assertEqual(r3.status_code, 200)
        r4 = c.get('/api/v1/screenings')
        self.assertEqual(r4.get_json()['screenings'][0]['screening_id'], sid)
        r5 = c.post('/api/v1/screenings', json={'prompt': 'Test the trailer on the Nobody twin'})
        self.assertEqual(r5.status_code, 404)
        # the retired Greenlight run: refused with the plain line, nothing charged; a compare on that scenario the same
        n_charges = len(charges)
        for body_ in ({'kind': 'greenlight', 'inputs': {'title': 'Second Helping', 'platforms': ['Hulu']}}, {'inputs': {}},
                      {'kind': 'compare', 'other_key': key, 'inputs': {'scenario_kind': 'greenlight'}}):
            r6 = c.post(f'/api/digital-twins/{key}/run', json=body_)
            self.assertEqual(r6.status_code, 400, r6.get_data(as_text=True)[:200])
            out6 = r6.get_json()
            self.assertEqual(out6['error'], DT.GREENLIGHT_RETIRED_LINE)
            self.assertTrue(out6.get('retired'))
            self.assertEqual(out6.get('tab'), 'greenlight')
        self.assertEqual(len(charges), n_charges)
        for tok in ('synth', 'pipeline', 'claude', '\u2014', 'odds'):
            self.assertNotIn(tok, DT.GREENLIGHT_RETIRED_LINE.lower())
        # a stored Greenlight run stays on the twin and never reaches the page
        stored = SYN.load_twin(s3, key)
        stored.setdefault('runs', []).insert(0, {'id': 'old', 'kind': 'greenlight', 'label': 'Greenlight Lab', 'question': 'odds?', 'outputs': {'expected': 1, 'low': 1, 'high': 2, 'odds_pct': 50.0}})
        SYN.save_twin(s3, key, stored)
        r7 = c.get(f'/api/digital-twins/{key}')
        self.assertEqual(r7.status_code, 200)
        self.assertNotIn('old', [x['id'] for x in r7.get_json()['twin']['runs']])
        self.assertIn('old', [x['id'] for x in SYN.load_twin(s3, key)['runs']])
        self.assertEqual(c.get(f'/api/digital-twins/{key}/status').get_json()['runs'], len(r7.get_json()['twin']['runs']))
        self.assertEqual(c.post(f'/api/digital-twins/{key}/greenlight/old/ask-the-room', json={}).status_code, 404)
        sys.modules.pop('app', None)


class GreenlightLane(unittest.TestCase):
    """Prometheus routes Greenlight to its two options (Jenna 2026-10-28):
    a logline / synopsis / idea is a Concept test, a script / treatment /
    attached screenplay is a Script read, a bare 'run Greenlight' asks
    which, and two twins on one idea is a Concept test with both seated."""

    @classmethod
    def setUpClass(cls):
        from prometheus import digital_twin_lane as L
        cls.L = L

    def test_option_from_the_words_and_the_file(self):
        L = self.L
        self.assertEqual(L.greenlight_option('Test this concept in Greenlight on the Gilmore Girls twin: ' + LOGLINE), 'concept')
        self.assertEqual(L.greenlight_option('What are the odds this logline clears 2M viewers on Peacock: ' + LOGLINE), 'concept')
        self.assertEqual(L.greenlight_option('Run Greenlight on the Mark Rober twin on this idea: a build box'), 'concept')
        self.assertEqual(L.greenlight_option('Read this script in Greenlight on the Mark Rober twin'), 'script')
        self.assertEqual(L.greenlight_option('Run Greenlight on the Beals twin with the attached treatment'), 'script')
        self.assertEqual(L.greenlight_option('Run Greenlight on the Beals twin with the attached file', SCRIPT), 'script')
        self.assertEqual(L.greenlight_option('Run Greenlight on the Gilmore Girls twin'), '')
        self.assertEqual(L.greenlight_option('Run Greenlight Lab on the Gilmore Girls twin', 'a one page memo with no sluglines'), '')
        self.assertEqual(L.greenlight_text('Test this concept in Greenlight on the Gilmore Girls twin: ' + LOGLINE, 'concept'), LOGLINE)
        self.assertEqual(L.greenlight_text('What are the odds this logline clears 2M viewers on Peacock: ' + LOGLINE, 'concept'), LOGLINE)
        self.assertEqual(L.greenlight_text('Put this to the room: "' + LOGLINE + '"', 'concept'), LOGLINE)
        self.assertEqual(L.greenlight_text('Test a concept in Greenlight on the Beals twin', 'concept', 'The one-sheet text.'), 'The one-sheet text.')
        self.assertEqual(L.greenlight_text('Test a concept in Greenlight on the Beals twin', 'concept'), '')
        self.assertEqual(L.greenlight_text('Read this script in Greenlight on the Beals twin', 'script', SCRIPT), SCRIPT.strip())
        self.assertEqual(L.greenlight_text('Read this script in Greenlight on the Beals twin', 'script'), '')
        self.assertEqual(L.greenlight_text('Read this script in Greenlight on the Beals twin. Script: ' + SCRIPT, 'script'), SCRIPT.strip())

    def test_intents_route_to_greenlight(self):
        L = self.L
        for ask in ('Test this concept in Greenlight on the Gilmore Girls twin: ' + LOGLINE, 'Read this script in Greenlight on the Mark Rober twin',
                    'Run Greenlight on the Gilmore Girls twin', 'Run Greenlight Lab on the Gilmore Girls twin: a new season', 'What are the odds this logline clears 2M viewers on the Beals twin: x'):
            self.assertTrue(L.twin_intent(ask), ask)
            self.assertEqual(L.lab_kind(ask), 'greenlight', ask)
        self.assertEqual(L.lab_kind('Compare the Gilmore Girls and Mark Rober twins on this idea: ' + LOGLINE), 'compare')
        self.assertEqual(L.question_kind('would this logline work for them: ' + LOGLINE), 'greenlight')
        self.assertEqual(L.question_kind('where does the second act drag in this script?'), 'greenlight')
        self.assertEqual(L.question_kind('what does $12.99 a month do to take rate?'), 'pricing')
        self.assertEqual(L.question_kind('how does the next season land on Netflix?'), 'content')
        self.assertFalse(hasattr(L, 'greenlight_inputs'))
        self.assertIn('run_greenlight', L.__all__)

    def test_run_greenlight_opens_the_option_or_asks(self):
        L = self.L
        from migration import digital_twin_synthesis as SYN
        from prometheus.host import host
        s3 = FakeS3()
        twin, young = _beals(), _young()
        twin['meta']['created_by'] = young['meta']['created_by'] = 'jenna'
        k1 = 'digital-twins/jenna/Jennifer_Beals_20261028_120000_abc123def0.json.gz'
        k2 = 'digital-twins/jenna/Young_Men_20261028_120000_abc123def1.json.gz'
        SYN.save_twin(s3, k1, twin)
        SYN.save_twin(s3, k2, young)
        host.bind(s3_client=s3, bucket='b')
        user = {'username': 'jenna', 'role': 'super_admin', 'has_digital_twins_iq_access': True}
        name = twin['meta']['name']
        # bare: one line, two chips
        ask = L.run_greenlight(f'Run Greenlight on the {name} twin', {}, [], user, {}, '')
        self.assertIn('two ways', ask['reply'])
        self.assertIn('Script read', ask['reply'])
        self.assertIn('Concept test', ask['reply'])
        self.assertEqual(len(ask['followups']), 2)
        self.assertEqual(ask['action_hint']['tab'], 'greenlight')
        self.assertFalse(any(s.get('mode') in ('script', 'concept') for s in SYN.load_twin(s3, k1).get('pilot_sessions') or []))
        # a concept without the text: asks for it, opens nothing
        ask2 = L.run_greenlight(f'Test a concept in Greenlight on the {name} twin', {}, [], user, {}, '')
        self.assertIn('Paste the logline', ask2['reply'])
        self.assertEqual(ask2['action_hint']['tab'], 'concept')
        # the concept: opens the test, free until Play, lands on Greenlight under Concept
        out = L.run_greenlight(f'Test this concept in Greenlight on the {name} twin: ' + LOGLINE, {}, [], user, {}, '')
        self.assertIn('Concept test', out['reply'])
        self.assertIn('charged on Play', out['reply'])
        self.assertEqual(out['action_hint']['tab'], 'concept')
        stored = SYN.load_twin(s3, k1)
        sess = next(s for s in stored['pilot_sessions'] if s['id'] == out['action_hint']['session_id'])
        self.assertEqual(sess['mode'], 'concept')
        self.assertEqual(sess['text'], LOGLINE)
        self.assertFalse(sess.get('play_charged'))
        self.assertEqual(PI.charge_due(sess), 'digital_twin_concept')
        self.assertEqual(sess['status'], 'ready')
        self.assertIsNotNone((sess.get('mode_read') or {}).get('definitely_interested_pct'))
        # the script, pasted: a Script read
        out3 = L.run_greenlight(f'Read this script in Greenlight on the {name} twin. Script: ' + SCRIPT, {}, [], user, {}, '')
        self.assertIn('Script read', out3['reply'])
        self.assertEqual(out3['action_hint']['tab'], 'script')
        s3s = next(s for s in SYN.load_twin(s3, k1)['pilot_sessions'] if s['id'] == out3['action_hint']['session_id'])
        self.assertEqual(s3s['mode'], 'script')
        self.assertEqual(PI.charge_due(s3s), 'digital_twin_script')
        # two twins on one idea through the compare lane: a Concept test with both seated, never a run
        out4 = L.run_lab('compare', f"Compare the {name} and {young['meta']['name']} twins on this idea: " + LOGLINE, {}, [], user, {}, '')
        self.assertIn('Concept test', out4['reply'])
        self.assertNotIn('job_id', out4)
        self.assertIn(out4['twin_key'], (k1, k2))
        s4 = next(s for s in SYN.load_twin(s3, out4['twin_key'])['pilot_sessions'] if s['id'] == out4['action_hint']['session_id'])
        self.assertEqual(sorted(s4['twin_keys']), sorted([k1, k2]))
        self.assertTrue(s4.get('two_audiences'))
        self.assertIn('seated too', out4['reply'])
        for k in (k1, k2):
            self.assertFalse(any(r.get('kind') in ('greenlight', 'compare') for r in SYN.load_twin(s3, k).get('runs') or []))
        for payload in (ask, ask2, out, out3, out4):
            blob = json.dumps(payload, default=str).lower()
            for tok in ('synth', 'claude', 'pipeline', 'hostmap', '\u2014', 'odds'):
                self.assertNotIn(tok, blob, tok)


class Page(unittest.TestCase):
    """The dashboard page carries every new tab and its console."""

    @classmethod
    def setUpClass(cls):
        cls.html = open(os.path.join(ROOT, 'templates', 'index.html'), encoding='utf-8').read()
        i = cls.html.index("var TABS = ['twins'")
        start = cls.html.rfind('<script', 0, i)
        cls.block = cls.html[start:cls.html.index('</script>', i)]

    def test_every_new_tab_is_on_the_page(self):
        for tab, pane, host in (('trailer', 'dtwTabTrailer', 'dtwTrConsole'), ('titles', 'dtwTabTitles', 'dtwTiConsole'), ('norms', 'dtwTabNorms', 'dtwNoConsole')):
            self.assertIn(f'data-tab="{tab}"', self.html, tab)
            self.assertIn(f'id="{pane}"', self.html, pane)
            self.assertIn(f'id="{host}"', self.html, host)
        self.assertIn("var TABS = ['twins', 'greenlight', 'launch', 'spend', 'panel', 'pilot', 'trailer', 'titles', 'norms'];", self.html)
        for label in ('>Trailers<', '>Titles and Key Art<', '>Norms<', '>Greenlight<'):
            self.assertIn(label, self.html, label)
        self.assertIn('flex-wrap: wrap', self.html[self.html.index('.dtw-segs {'):self.html.index('.dtw-segs {') + 200])

    def test_greenlight_is_two_options_and_the_standalone_tabs_are_gone(self):
        """Jenna 2026-10-28: Greenlight opens on Script read and Concept; the
        Script Reads and Concepts tabs fold into it; the odds run is gone."""
        bar = self.html[self.html.index('<div class="dtw-segs" role="tablist">'):self.html.index('<div class="dtw-fg-acts" id="dtwFgActs"')]
        labels = re.findall(r'role="tab" onclick="dtwTab\(\'[a-z]+\'\)">([^<]+)</button>', bar)
        self.assertEqual(labels, ['Twins', 'Greenlight', 'Launch', 'Spend Engine', 'Focus Group IQ', 'Pilot Testing', 'Trailers', 'Titles and Key Art', 'Norms'])
        for gone in ('data-tab="script"', 'data-tab="concept"', 'id="dtwTabScript"', 'id="dtwTabConcept"', '>Script Reads<', '>Concepts<'):
            self.assertNotIn(gone, self.html, gone)
        pane = self.html[self.html.index('<section id="dtwTabGreenlight"'):self.html.index('<section id="dtwTabLaunch"')]
        for need in ('data-gl="script"', 'data-gl="concept"', "dtwGreenlightPick('script')", "dtwGreenlightPick('concept')", '>Script read<', '>Concept<',
                     'id="dtwScConsole"', 'id="dtwCoConsole"', 'id="dtwGlScriptPrice"', 'id="dtwGlConceptPrice"'):
            self.assertIn(need, pane, need)
        for gone in ('dtwGlResults', 'Run Greenlight', 'Odds on a title', 'dtwGlText', 'dtwGlOther', 'data-dtw-run'):
            self.assertNotIn(gone, pane, gone)
        # deep links and chips land on Greenlight with the option picked
        self.assertIn("if (GL_OPTIONS.indexOf(tab) >= 0) { S.glOption = tab; tab = 'greenlight'; }", self.block)
        self.assertIn("var GL_OPTIONS = ['script', 'concept'];", self.block)
        self.assertIn("PI_HOST = { pilot: 'dtwPiConsole', trailer: 'dtwTrConsole', titles: 'dtwTiConsole', script: 'dtwScConsole', concept: 'dtwCoConsole' }", self.block)
        self.assertIn("script: 'Script read', concept: 'Concept'", self.block)
        self.assertIn("greenlight: 'Greenlight'", self.block)
        self.assertIn('window.dtwGreenlightPick = function(opt)', self.block)
        self.assertIn('function syncGreenlightOptions()', self.block)
        self.assertIn("find('greenlight script')", self.block)
        self.assertIn("find('greenlight concept')", self.block)
        # the old Greenlight run is off the page
        for gone in ("renderEmpty('dtwGlResults', 'greenlight')", 'function renderGreenlight(', 'window.dtwGlAskRoom', 'Back to the Greenlight run', 'Run Greenlight Lab on the',
                     "dtwRunLab('greenlight')", 'ask-the-room', 'Ask the room again', 'concept_read', "'compare': 'greenlight'", 'How do I run Greenlight Lab'):
            self.assertNotIn(gone, self.html, gone)
        for need in ('function renderRun(', 'id="dtwRunResults"', 'function isRetiredRun(r)', 'function productRuns(t)', "productRuns(t).forEach(", 'S.twin.runs = productRuns(S.twin);'):
            self.assertIn(need, self.html, need)
        self.assertIn("Test a concept in Greenlight on the ' + (tw || 'open') + ' twin'", self.html)
        self.assertIn('How do I use Greenlight on a twin?', self.html)
        self.assertNotIn('three runs', self.block)

    def test_block_has_the_phase_2_renderers_and_no_dashes(self):
        for fn in ('function renderModeRead(', 'function renderDiagnostics(', 'function renderPilotExtras(', 'async function renderNormsTab(', 'window.dtwPilotRecut', 'window.dtwPilotCompare',
                   'window.dtwPilotComps', 'window.dtwPilotCritics', 'window.dtwPilotTalent', 'window.dtwPilotBrief', 'Ask the room', 'report.pptx', 'Every screening adds to these'):
            self.assertIn(fn, self.block, fn)
        self.assertNotIn('\u2014', self.block)
        self.assertNotIn('\u2013', self.block)
        phase2 = self.block[self.block.index('Pilot Testing phase 2: the mode reads'):self.block.index('window._dtwRender.renderPilotConsole = renderPilotConsole;')]
        for tok in ('synth', 'Claude', 'hostmap', 'Screen Engine', 'NRG', 'MarketCast', 'FishDog', 'persona'):
            self.assertNotIn(tok, phase2, tok)


class Pricing(unittest.TestCase):
    def test_charge_due_per_mode(self):
        twin = _beals()
        want = {'pilot': 'digital_twin_pilot', 'feature': 'digital_twin_pilot', 'episode': 'digital_twin_pilot', 'binge': 'digital_twin_pilot', 'trailer': 'digital_twin_trailer',
                'spot': 'digital_twin_trailer', 'title': 'digital_twin_title', 'key_art': 'digital_twin_key_art', 'concept': 'digital_twin_concept', 'script': 'digital_twin_script'}
        for mode, key in want.items():
            sess = PI.open_session([twin], title='t', in_room=12, mode=mode)
            self.assertEqual(PI.charge_due(sess), key, mode)
            self.assertEqual(DT.PRICING[key], {'digital_twin_pilot': 5500.0, 'digital_twin_trailer': 2500.0, 'digital_twin_title': 1500.0, 'digital_twin_key_art': 1500.0,
                                               'digital_twin_concept': 2500.0, 'digital_twin_script': 3500.0}[key])
            PI.mark_play_paid(sess)
            self.assertIsNone(PI.charge_due(sess))
            sess['questions'] = [{'q': str(i)} for i in range(12)]
            self.assertEqual(PI.charge_due(sess), 'digital_twin_pilot_more')
        recut = PI.open_session([twin], title='Cut B', in_room=12)
        recut['lineage'] = {'parent': 'abc', 'label': 'Cut B'}
        self.assertEqual(PI.charge_due(recut), 'digital_twin_recut')
        self.assertEqual(DT.PRICING['digital_twin_recut'], 2750.0)
        for k in ('digital_twin_trailer', 'digital_twin_title', 'digital_twin_key_art', 'digital_twin_concept', 'digital_twin_script', 'digital_twin_recut'):
            self.assertIn(k, DT.PULL_TYPES)

    def test_how_to_items_for_every_tab(self):
        ids = [it['id'] for it in DT.how_to()['items']]
        for tab in ('pilot', 'trailer', 'titles', 'greenlight', 'norms'):
            self.assertIn(tab, ids)
        for folded in ('script', 'concept'):
            self.assertNotIn(folded, ids)          # the two options live inside the Greenlight item (2026-10-28)
        gl = next(it for it in DT.how_to()['items'] if it['id'] == 'greenlight')
        self.assertEqual([o['id'] for o in gl['options']], ['script', 'concept'])
        text = json.dumps(DT.how_to(), default=str).lower()
        for tok in ('screen engine', 'nrg', 'marketcast', 'fishdog', 'synth', 'claude', 'pipeline', '\u2014', 'greenlight lab'):
            self.assertNotIn(tok, text)
        self.assertIn('Trailers', DT.how_to_text('how do I run a trailer test'))
        self.assertIn('Norms', DT.how_to_text('what is the norms tab'))
        self.assertIn('Greenlight.', DT.how_to_text('how do I use greenlight'))
        self.assertIn('Greenlight.', DT.how_to_text('how do I run a script read'))
        self.assertIn('Greenlight.', DT.how_to_text('how do I run a concept test'))

    def test_rate_card_carries_the_modes(self):
        from prometheus import rate_card as RC
        text = RC.render()
        for frag in ('Trailer or TV spot test', 'Title or logline test', 'Key art test', 'Greenlight Concept test', 'Greenlight Script read', 'Recut re-test', 'with its first two runs'):
            self.assertIn(frag, text)
        for gone in ('Greenlight, Launch', 'first three runs'):
            self.assertNotIn(gone, text)
        self.assertEqual(RC.default('digital_twin_trailer'), 2500.0)
        self.assertEqual(RC.default('digital_twin_concept'), 2500.0)
        self.assertEqual(RC.default('digital_twin_script'), 3500.0)
        self.assertNotIn('\u2014', text)


if __name__ == '__main__':
    unittest.main(verbosity=1)
