#!/usr/bin/env python3
"""Deck photography (Jenna 2026-09-30): real photos under a graphite
scrim on the cover, the first statement page, and the close -
mirroring the manually designed decks. Offline by construction: the
photo cache is injected with a locally generated image, and the
kill-switch path proves a text-only render when nothing resolves."""
import os
import struct
import sys
import zlib

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import deck_builder as db  # noqa: E402
import prometheus_analysis as pma  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

FAIL = 0


def check(name, ok, detail=''):
    global FAIL
    print(('PASS' if ok else 'FAIL'), name, detail)
    if not ok:
        FAIL += 1


def png(w, h, rgb=(40, 90, 120)):
    def chunk(t, data):
        c = t + data
        return (struct.pack('>I', len(data)) + c
                + struct.pack('>I', zlib.crc32(c) & 0xffffffff))
    ihdr = struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)
    row = b'\x00' + bytes(rgb) * w
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', ihdr)
            + chunk(b'IDAT', zlib.compress(row * h))
            + chunk(b'IEND', b''))


def pics(slide):
    return [sh for sh in slide.shapes
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE
            and sh.width > 9000000]  # full-bleed only, not the logo


def text_boxes(slide):
    out = []
    for sh in slide.shapes:
        if sh.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
            continue
        if not (sh.text_frame.text or '').strip():
            continue
        out.append((sh.left / 914400.0, sh.top / 914400.0,
                    sh.width / 914400.0, sh.height / 914400.0,
                    sh.text_frame.text.strip()[:40]))
    return out


def overlaps(a, b):
    ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return ox > 0.10 and oy > 0.05


# px gate and blob sizing work on the generated image
blob = png(1216, 812)
check('blob pixel read works', db._blob_px(blob) == (1216, 812),
      repr(db._blob_px(blob)))

# inject the cache so the render never touches the network
db._PHOTO_CACHE['paige bueckers|person'] = [blob]

PLAN = {'title': 'x', 'filename_stem': 'photo_test',
        'image_subject': 'Paige Bueckers', 'image_kind': 'person',
        'slides': [
            {'type': 'cover', 'eyebrow': 'PREPARED FOR NIKE . WNBA',
             'title': 'Her audience out-shops the league.',
             'intro': 'What this deck contains and why it matters.',
             'stats': [{'big': '3.1M', 'label': 'Projected US '
                        'audience'}]},
            {'type': 'hero', 'ground': 'dark', 'title': 'The '
             'sneaker wallet is already open.',
             'big': '3.2x the US average',
             'line': 'Sneaker channel shopping runs ahead of every '
             'peer audience.'},
            {'type': 'bars', 'title': 'Where the money goes.',
             'rows': [{'label': 'Nike', 'value': 61.2,
                       'accent': True},
                      {'label': 'Adidas', 'value': 34.8}]},
            {'type': 'close', 'title': 'Three moves this quarter.',
             'cards': [{'head': 'Sign the sneaker moment',
                        'body': 'The audience already shops the '
                        'channel at three times the US average.'}]},
        ]}

out = '/tmp/deck_photo_test.pptx'
n = db.render_insights_deck(PLAN, out)
check('all slides rendered', n == 4, f'({n})')

prs = Presentation(out)
check('cover carries a full-bleed photo', len(pics(prs.slides[0])) == 1)
check('hero carries a full-bleed photo', len(pics(prs.slides[1])) == 1)
check('data slide stays photo-free', len(pics(prs.slides[2])) == 0)
check('close carries a full-bleed photo', len(pics(prs.slides[3])) == 1)

pic = pics(prs.slides[0])[0]
# 1216x812 (1.4975) into 13.333x7.5 (1.7777): crop top+bottom, with
# the window biased toward the top so faces stay in frame
frac = 1.0 - (1216.0 / 812.0) / (13.333 / 7.5)
check('cover photo is cover-cropped, never distorted',
      abs((pic.crop_top + pic.crop_bottom) - frac) < 0.01
      and pic.crop_top < pic.crop_bottom and pic.crop_left == 0,
      f'top={pic.crop_top:.4f} bottom={pic.crop_bottom:.4f} '
      f'frac~{frac:.4f}')

xml1 = prs.slides[0].shapes._spTree.xml
check('scrim veils the cover photo (64% graphite)',
      'val="64000"' in xml1)

# photo slides stay overlap-clean
clean = True
for si, slide in enumerate(prs.slides, 1):
    tb = text_boxes(slide)
    for i in range(len(tb)):
        for j in range(i + 1, len(tb)):
            if overlaps(tb[i], tb[j]):
                clean = False
                check(f'slide {si}: text overlap', False,
                      f'{tb[i][4]!r} x {tb[j][4]!r}')
check('photo slides stay overlap-clean', clean)

# photo: false opts a moment out
PLAN2 = {k: v for k, v in PLAN.items()}
PLAN2['slides'] = [dict(s) for s in PLAN['slides']]
PLAN2['slides'][0]['photo'] = False
n = db.render_insights_deck(PLAN2, '/tmp/deck_photo_test2.pptx')
prs2 = Presentation('/tmp/deck_photo_test2.pptx')
check('photo:false keeps the cover type-only',
      len(pics(prs2.slides[0])) == 0)

# kill switch: uncached subject with photos disabled never fetches
os.environ['PROMETHEUS_DECK_PHOTOS'] = '0'
try:
    got = db._deck_photos('Nobody Offline Test', 'person')
    check('kill switch returns no photos without a fetch', got == [])
    PLAN3 = dict(PLAN)
    PLAN3 = {**PLAN, 'image_subject': 'Nobody Offline Test'}
    PLAN3['slides'] = [dict(s) for s in PLAN['slides']]
    n = db.render_insights_deck(PLAN3, '/tmp/deck_photo_test3.pptx')
    prs3 = Presentation('/tmp/deck_photo_test3.pptx')
    check('no photos means text-only render, never a failure',
          n == 4 and all(len(pics(s)) == 0 for s in prs3.slides))
finally:
    os.environ.pop('PROMETHEUS_DECK_PHOTOS', None)

# a corrupt blob in the cache renders text-only, never a failure
db._PHOTO_CACHE['garbage subject|person'] = [b'not an image at all']
PLANG = {**PLAN, 'image_subject': 'Garbage Subject'}
PLANG['slides'] = [dict(s) for s in PLAN['slides']]
n = db.render_insights_deck(PLANG, '/tmp/deck_photo_test4.pptx')
prs4 = Presentation('/tmp/deck_photo_test4.pptx')
check('corrupt photo bytes never fail the deck',
      n == 4 and all(len(pics(s)) == 0 for s in prs4.slides))

# plan enforcement carries the art fields
cleaned = pma.enforce_insights_plan(
    {'title': 'x', 'filename_stem': 'y',
     'image_subject': 'Paige Bueckers', 'image_kind': 'person',
     'slides': [{'type': 'cover', 'title': 't', 'photo': False}]},
    'Paige Bueckers')
check('enforce carries image_subject',
      cleaned.get('image_subject') == 'Paige Bueckers')
check('enforce whitelists image_kind',
      cleaned.get('image_kind') == 'person')
check('enforce drops an unknown image_kind',
      pma.enforce_insights_plan(
          {'title': 'x', 'image_kind': 'meme', 'slides': []},
          'x').get('image_kind') == '')
check('photo opt-out survives the plan clean',
      cleaned['slides'][0].get('photo') is False)

# prompt and callsite wiring
check('plan prompt carries the art direction section',
      'ART DIRECTION' in pma.INSIGHTS_DECK_SYSTEM_PROMPT
      and 'image_subject' in pma.INSIGHTS_DECK_SYSTEM_PROMPT)
app_src = open(os.path.join(ROOT, 'app.py'), encoding='utf-8').read()
check('deck job passes the subject for photography',
      'photo_subject=str(subject' in app_src)

print()
if FAIL:
    print(f'{FAIL} CHECK(S) FAILED')
    sys.exit(1)
print('ALL CHECKS PASSED')
