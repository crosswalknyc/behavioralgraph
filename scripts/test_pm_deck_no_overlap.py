#!/usr/bin/env python3
"""No overlapping text in Prometheus decks (Jenna 2026-09-30).

Renders a deliberately pathological deck (the PAW Patrol cover from
the defect screenshot plus worst-case long strings on every slide
type) and asserts, shape by shape, that no two rendered text boxes
overlap on any slide.
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import deck_builder as db  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

FAIL = 0


def check(name, ok, detail=''):
    global FAIL
    print(('PASS' if ok else 'FAIL'), name, detail)
    if not ok:
        FAIL += 1


PAW_TITLE = ('756,363 streams sit on top of 6,632,879 parents, and '
             'Arts and crafts is the lane nobody is selling them.')
PAW_INTRO = ("What this deck contains: the parent pool behind the "
             "film's first week in the Paramount+ top 10, what that "
             "pool already buys, how it answers a PAW Patrol unit, "
             "and the toy demand lane the current line does not "
             "cover. Window Sep 25 2025 to Sep 24 2026, with the "
             "stream week Sep 18 to Sep 24 2026.")
LONG_TITLE = ('The parent pool already buys the adjacent lanes, and '
              'the arts and crafts shelf is where the demand has '
              'nowhere to land.')
LONG_SUB = ('Every figure on this page is the parent pool measured '
            'across the full window, ranked against the US average, '
            'with the strongest lane called out in the accent color '
            'so the shelf gap is impossible to miss.')

PLAN = {'title': 'x', 'filename_stem': 'overlap_test', 'slides': [
    {'type': 'cover', 'eyebrow': 'PAW PATROL . PREPARED FOR SPIN '
     'MASTER . Q4 2026', 'title': PAW_TITLE, 'intro': PAW_INTRO,
     'stats': [{'big': '6.6M', 'label': 'Projected US PAW Patrol '
                'parent pool'},
               {'big': '40.3%', 'label': 'Viewers reaching a PAW '
                'Patrol product page inside 30 days'},
               {'big': '2.6%', 'label': 'Click rate on PAW '
                'Patrol-tagged units'}], 'accent_index': 1},
    {'type': 'argument', 'ground': 'dark', 'title': LONG_TITLE,
     'sub': LONG_SUB, 'cards': [
         {'head': 'The pool is bigger than the franchise shelf '
          'assumes it is', 'body': 'Projected parents in the window '
          'run past the current line plan, and the strongest lane '
          'is one the line does not carry at all, which is the '
          'whole argument of this deck in one card.'}] * 4},
    {'type': 'bars', 'title': LONG_TITLE, 'sub': LONG_SUB,
     'show_index': True, 'read': 'The accent lane clears every '
     'other lane by a third.', 'rows': [
         {'label': 'Arts and crafts activity kits for '
          'preschoolers', 'value': 61.4, 'index': 213,
          'accent': True}] + [
         {'label': 'Building sets', 'value': 44.2, 'index': 151}] * 8},
    {'type': 'split_stats_bars', 'title': LONG_TITLE, 'sub': LONG_SUB,
     'stat_cards': [{'kicker': 'POOL', 'big': '6.6M',
                     'label': 'Projected US parent pool'},
                    {'kicker': 'REACH', 'big': '40.3%',
                     'label': 'Product page reach inside 30 days',
                     'accent': True}],
     'bars_title': 'LANES RANKED', 'read': 'Crafts leads every lane.',
     'rows': [{'label': 'Arts and crafts kits', 'value': 61.4,
               'accent': True}] + [
         {'label': 'Plush', 'value': 38.1}] * 7},
    {'type': 'tiles_row', 'title': LONG_TITLE, 'sub': LONG_SUB,
     'read': 'All three tiles point at the same shelf.', 'tiles': [
         {'big': '61.4%', 'label': 'Parents already buying arts and '
          'crafts kits in the window',
          'body': 'The lane runs ahead of every toy lane the line '
          'carries today, and the gap holds across every age band '
          'in the pool, which makes it the cleanest first unit to '
          'sell against this audience.'}] * 3},
    {'type': 'hero', 'ground': 'dark', 'title': LONG_TITLE,
     'sub': LONG_SUB, 'big': '6,632,879 parents in the pool',
     'line': 'That is the projected US parent pool behind the first '
     'week, measured across the full window.',
     'support': 'The pool is the sale: it is already watching, '
     'already buying adjacent lanes, and not yet offered the lane '
     'it over-indexes on.'},
    {'type': 'table', 'title': LONG_TITLE, 'sub': LONG_SUB,
     'columns': ['Lane', 'Pool share', 'US average', 'Index',
                 'Current line'],
     'rows': [['Arts and crafts activity kits', '61.4%', '28.8%',
               '213', 'Not carried'],
              ['Building sets', '44.2%', '29.3%', '151', 'Carried'],
              ['Plush', '38.1%', '31.2%', '122', 'Carried'],
              ['Vehicles', '35.4%', '30.9%', '115', 'Carried'],
              ['Role play', '31.2%', '27.4%', '114', 'Carried']],
     'accent_row': 0, 'accent_col': 3,
     'read': 'The one lane the line does not carry is the one the '
     'pool over-indexes on hardest.',
     'read2': 'Building sets prove the pool converts when the shelf '
     'exists.'},
    {'type': 'hero_proof', 'title': LONG_TITLE, 'sub': LONG_SUB,
     'big': '756,363 first-week streams',
     'line': 'First week inside the Paramount+ top 10, measured at '
     'the account level across the window.',
     'proofs': [{'fig': '6.6M', 'label': 'Projected US parent pool'},
                {'fig': '40.3%', 'label': 'Product page reach'},
                {'fig': '2.6%', 'label': 'Click rate on tagged '
                 'units'}]},
    {'type': 'paths', 'title': LONG_TITLE, 'sub': LONG_SUB,
     'rows': [{'kind': 'search', 'url': 'google.com/search?q=paw+'
               'patrol+crafts', 'lit': False},
              {'kind': 'cart', 'url': 'target.com/cart/paw-patrol-'
               'craft-kit', 'lit': True}] * 5},
    {'type': 'close', 'title': LONG_TITLE, 'cards': [
        {'head': 'Sell the crafts lane first because the pool '
         'already proved it', 'body': 'The pool over-indexes 213 on '
         'arts and crafts and the line does not carry it, which '
         'makes it the cleanest first unit and the fastest proof '
         'of the partnership.'}] * 4},
    {'type': 'tiles_facts', 'title': LONG_TITLE, 'sub': LONG_SUB,
     'tiles': [{'big': '6.6M', 'label': 'Projected US parent pool '
                'behind the first week'}] * 4, 'accent_index': 0,
     'read': 'The pool is young, urban, and already shopping the '
     'adjacent lanes.',
     'facts': [{'label': 'Parents of kids aged two to seven in the '
                'pool', 'fig': '71.2%', 'note': 'The core preschool '
                'band the franchise sells to, concentrated well '
                'past the US average for streaming families'}] * 6},
]}


def text_boxes(slide):
    out = []
    for sh in slide.shapes:
        if sh.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
            continue
        t = (sh.text_frame.text or '').strip()
        if not t:
            continue
        out.append((sh.left / 914400.0, sh.top / 914400.0,
                    sh.width / 914400.0, sh.height / 914400.0,
                    t[:44]))
    return out


def overlaps(a, b, tol=0.05):
    ax, ay, aw, ah, _ = a
    bx, by, bw, bh, _ = b
    ox = min(ax + aw, bx + bw) - max(ax, bx)
    oy = min(ay + ah, by + bh) - max(ay, by)
    return ox > 0.10 and oy > tol


# sanity: the measurement layer sees the defect title as 3+ lines
check('measurement: PAW title wraps past two lines at 34pt',
      db._line_count(PAW_TITLE, 34, 11.493, bold=True) >= 3)
check('fit: chrome titles step down to two lines or fewer',
      db._fit_size(LONG_TITLE, (34, 31, 28, 26), 11.493, 2) < 34)

out_path = '/tmp/deck_overlap_test.pptx'
n = db.render_insights_deck(PLAN, out_path)
check('all slides rendered', n == len(PLAN['slides']), f'({n})')

prs = Presentation(out_path)
clean = True
for si, slide in enumerate(prs.slides, 1):
    boxes = text_boxes(slide)
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            if overlaps(boxes[i], boxes[j]):
                clean = False
                check(f'slide {si}: no text overlap', False,
                      f'{boxes[i][4]!r} x {boxes[j][4]!r}')
check('no two text boxes overlap on any slide', clean)

# the defect signature specifically: cover intro sits below the title
cover = prs.slides[0]
tb = text_boxes(cover)
title = next(b for b in tb if b[4].startswith('756,363'))
intro = next(b for b in tb if b[4].startswith('What this deck'))
check('cover intro starts below the measured title',
      intro[1] >= title[1] + title[3] - 0.01)

print()
if FAIL:
    print(f'{FAIL} CHECK(S) FAILED')
    sys.exit(1)
print('ALL CHECKS PASSED')
