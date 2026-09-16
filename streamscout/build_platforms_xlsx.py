#!/usr/bin/env python3
"""Build the flashy StreamScout-Platforms.xlsx (20 platforms).

Writes StreamScout-Platforms.xlsx next to this script by default; pass a path
argument to write elsewhere (e.g. python build_platforms_xlsx.py ~/Desktop)."""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

_dest = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.expanduser(_dest), "StreamScout-Platforms.xlsx") \
    if os.path.isdir(os.path.expanduser(_dest)) else os.path.expanduser(_dest)

# ── palette ───────────────────────────────────────────────────────────────────
NAVY   = "0B1F3A"
NAVY2  = "13294B"
ACCENT = "1F6FEB"
WHITE  = "FFFFFF"
INK    = "1B2733"
ZEBRA  = "F3F7FD"
GRID   = "D8E2F0"
# access pills (fill, text)
NOLOGIN = ("DAF2E0", "1A7F37")   # green
LOGIN   = ("FFE3B3", "8A5300")   # amber
APIKEY  = ("D9E6FF", "0A3D91")   # blue
SITU    = ("FFD9C2", "9A3B12")   # orange
CHECK   = ("E7F7EC", "1A7F37")   # green check cell

thin = Side(style="thin", color=GRID)
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# ── data ──────────────────────────────────────────────────────────────────────
# (platform, movies, series, access_key, what_you_get, notes)
VIDEO = [
    ("Peacock",       1, 1, "no",    "watch / show ids",              "Optional company-DB fallback for odd titles"),
    ("Hulu",          1, 1, "no",    "episode ids",                   ""),
    ("Netflix",       1, 1, "login", "title / episode ids",           "Self-driving Firefox window"),
    ("Apple TV+",     1, 1, "no",    "one id per series",             "Series shell — SEASON left blank"),
    ("Paramount Plus",1, 1, "no",    "episode watch URLs",            "Walks every season (all-seasons safe)"),
    ("HBO MAX",       1, 1, "login", "episode ids",                   "Self-driving Firefox window"),
    ("Disney+",       1, 1, "situ",  "episode ids",                   "No-login first; Firefox only for tricky titles"),
    ("Starz",         1, 1, "no",    "episode ids",                   ""),
    ("Hallmark Plus", 1, 1, "no",    "episode ids",                   ""),
    ("Amazon",        1, 1, "no",    "id fragments (ASIN + GTI)",     "Every edition + Amazon Channels (Lionsgate+, Starz…)"),
    ("MGM Plus",      1, 1, "no",    "watch paths",                   "One season shell per season"),
    ("BritBox",       1, 1, "no",    "show / movie shell path",       "One shell per title"),
    ("YouTube",       1, 1, "no",    "every episode watch?v= URL",    "Channel Videos tab; Shorts excluded"),
]
POD = [
    ("Spotify",         1, 1, "api",  "every episode open.spotify URL", "Free developer API key (no browser)"),
    ("Apple Podcasts",  1, 1, "no",   "every episode URL",              "Public feed caps ~200 episodes"),
    ("iHeart",          1, 1, "no",   "every episode URL",              ""),
    ("Pandora",         1, 1, "no",   "every episode URL",              ""),
    ("Amazon Podcasts", 1, 1, "no",   "every episode URL",              "Self-driving Chromium window"),
    ("SiriusXM",        1, 1, "situ", "every episode URL",              "Paste link = no login; title search = 1-time login"),
]
AUDIO = [
    ("Audible",         1, 1, "no",   "direct Audible + via-Amazon",    "TWO rows / title: audible.com/pd + amazon.com/dp (different ASINs)"),
]
ACCESS = {
    "no":    ("No login",      NOLOGIN),
    "login": ("Browser login", LOGIN),
    "api":   ("Free API key",  APIKEY),
    "situ":  ("Situational",   SITU),
}

TOTAL = len(VIDEO) + len(POD) + len(AUDIO)

# ── workbook ──────────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Platforms"
ws.sheet_view.showGridLines = False

HEADERS = ["", "Platform", "Movies", "Series", "Access", "What you get", "Notes"]
WIDTHS  = [4.5, 20, 9.5, 9.5, 16, 30, 50]
for i, w in enumerate(WIDTHS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
NCOL = len(HEADERS)
last = get_column_letter(NCOL)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def band(row, text, bg, fg, size=11, bold=True, h=None, align="left"):
    ws.merge_cells(f"A{row}:{last}{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.fill = fill(bg)
    c.font = Font(name="Aptos", size=size, bold=bold, color=fg)
    c.alignment = Alignment(horizontal=align, vertical="center", indent=1)
    if h:
        ws.row_dimensions[row].height = h


r = 1
# title banner
band(r, "  🛰  StreamScout", NAVY, WHITE, size=22, h=40)
r += 1
band(r, f"  One tool → every watch / play id  ·  {TOTAL} platforms  ·  movies, series, podcasts & audiobooks",
     NAVY2, "BFD3EE", size=11, bold=False, h=22)
r += 1
band(r, "", WHITE, WHITE, h=6)   # spacer
r += 1


def header_row(row):
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = fill(ACCENT)
        c.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="center" if i in (3, 4, 5) else "left",
                                vertical="center", indent=0 if i in (3, 4, 5) else 1)
        c.border = border
    ws.row_dimensions[row].height = 22


def data_rows(rows, start, zebra_offset=0):
    for j, (name, mv, sr, acc, get, notes) in enumerate(rows):
        row = start + j
        z = (j + zebra_offset) % 2
        base = ZEBRA if z else WHITE
        ws.row_dimensions[row].height = 21
        # icon / bullet
        c = ws.cell(row=row, column=1, value="▸")
        c.fill = fill(base); c.font = Font(name="Aptos", size=10, color=ACCENT)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
        # platform name
        c = ws.cell(row=row, column=2, value=name)
        c.fill = fill(base); c.font = Font(name="Aptos", size=11, bold=True, color=INK)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c.border = border
        # movies / series checks
        for col, val in ((3, mv), (4, sr)):
            c = ws.cell(row=row, column=col, value="✓" if val else "—")
            cf, ct = CHECK if val else (base, "9AA7B4")
            c.fill = fill(cf)
            c.font = Font(name="Aptos", size=11, bold=True, color=ct)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
        # access pill
        label, (pf, pt) = ACCESS[acc]
        c = ws.cell(row=row, column=5, value=label)
        c.fill = fill(pf)
        c.font = Font(name="Aptos", size=9.5, bold=True, color=pt)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
        # what you get
        c = ws.cell(row=row, column=6, value=get)
        c.fill = fill(base); c.font = Font(name="Aptos", size=10, color=INK)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c.border = border
        # notes
        c = ws.cell(row=row, column=7, value=notes)
        c.fill = fill(base); c.font = Font(name="Aptos", size=10, color="4A5A6A", italic=bool(notes))
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c.border = border
    return start + len(rows)


header_row(r)
header_start = r
r += 1
# section: streaming video
band(r, f"📺  Streaming Video  ·  {len(VIDEO)}", "E8F0FC", NAVY2, size=10, h=20)
r += 1
r = data_rows(VIDEO, r)
# section: podcasts
band(r, f"🎧  Podcasts  ·  {len(POD)}", "E8F0FC", NAVY2, size=10, h=20)
r += 1
r = data_rows(POD, r, zebra_offset=0)
# section: audiobooks
band(r, f"📚  Audiobooks  ·  {len(AUDIO)}", "E8F0FC", NAVY2, size=10, h=20)
r += 1
r = data_rows(AUDIO, r, zebra_offset=0)

# freeze under header
ws.freeze_panes = f"A{header_start + 1}"

# ── legend / good-to-know ─────────────────────────────────────────────────────
r += 1
band(r, "  Good to know", NAVY, WHITE, size=12, h=26)
r += 1
NOTES = [
    f"15 of {TOTAL} platforms need NO login. Only 5 need anything: Netflix & HBO Max (browser login), "
    "Spotify (free API key), and Disney+ & SiriusXM (situational).",
    "Access legend:  No login = anonymous  ·  Browser login = a Firefox/Chromium window drives itself  ·  "
    "Free API key = Spotify Client ID+Secret  ·  Situational = usually none, occasional login.",
    "Every run writes one CSV to your Desktop:  SHOW · URL · PRODUCTION · PLATFORM · SEASON.  "
    "PRODUCTION (studio) fills in automatically; SEASON is blank for movies, podcasts, audiobooks, Apple TV+, BritBox & YouTube.",
    "Podcast platforms return one row per episode. Streaming platforms handle both movies and series, "
    "with flexible season picks (all · 1-3 · 1,4,6).",
    "SiriusXM title search is Netflix-style: one device code the first time on a new computer, then automatic. "
    "Pasting a SiriusXM show link never needs a login.",
    "Audible returns TWO rows per audiobook — the direct audible.com/pd listen link AND the parallel "
    "amazon.com/dp 'Audible Audio Edition' (a different ASIN) — each labeled Audible / Amazon in the PLATFORM column.",
    "Amazon captures every clickstream form — all offer ASINs, the GTI in both encodings, shells + episodes, "
    "every film edition, and Amazon Channels (Lionsgate+, Starz…).",
]
for n in NOTES:
    ws.merge_cells(f"A{r}:{last}{r}")
    c = ws.cell(row=r, column=1, value="•  " + n)
    c.fill = fill(ZEBRA if (r % 2) else WHITE)
    c.font = Font(name="Aptos", size=9.5, color=INK)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[r].height = 30
    r += 1

wb.save(OUT)
print("wrote", OUT, "· platforms:", TOTAL)
