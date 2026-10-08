"""Attribution IQ: one-click Excel export of everything a campaign holds.

Builds a single ``.xlsx`` workbook for a campaign (Be Good Influence is
the first customer) so the whole Attribution IQ read can travel outside
the dashboard. Every dataset the Attribution IQ tabs draw from lands
on its own sheet: campaign facts, phases, every asset, content types,
paid vs organic, the audiences and congressional districts, cohorts,
demographics, and the multi-touch attribution read (touchpoints,
journeys, funnel, time to action, co-exposure, per-audience views).

Design notes
------------
* Pure read. Nothing here writes to S3 or ClickHouse.
* Fail-soft per sheet. A dataset that cannot be loaded produces a
  one-line sheet saying so instead of killing the download.
* Reader-facing vocabulary only. Internal provenance fields (ingest
  sources, job ids, calibration notes) are dropped on purpose.
* Entry point: :func:`build_workbook` returns the ``.xlsx`` bytes plus a
  suggested filename.
"""

from __future__ import annotations

import datetime as _dt
import io
import json
import re
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Crosswalk palette (deck system): Graphite Teal ground, Off-White type,
# Signal Green accent. Used sparingly: header rows and section labels.
_GRAPHITE = "0C1618"
_SLATE = "15252A"
_OFFWHITE = "E9E8E1"
_SIGNAL = "C7F23E"
_MUTED = "9AA09B"

_FONT = "Inter"

_HEADER_FONT = Font(name=_FONT, bold=True, color=_OFFWHITE, size=10)
_HEADER_FILL = PatternFill("solid", fgColor=_GRAPHITE)
_SECTION_FONT = Font(name=_FONT, bold=True, color=_GRAPHITE, size=11)
_BODY_FONT = Font(name=_FONT, size=10)
_MUTED_FONT = Font(name=_FONT, size=9, color="5C6560")
_TITLE_FONT = Font(name=_FONT, bold=True, size=14, color=_GRAPHITE)

_MAX_SHEET_NAME = 31
_MAX_COL_WIDTH = 60


# ----------------------------------------------------------------------
# Small helpers
# ----------------------------------------------------------------------

def _sheet_title(name: str, taken: set) -> str:
    base = re.sub(r"[\[\]\:\*\?\/\\]", " ", str(name or "Sheet")).strip() or "Sheet"
    base = base[:_MAX_SHEET_NAME]
    title = base
    n = 2
    while title.lower() in taken:
        suffix = f" ({n})"
        title = base[: _MAX_SHEET_NAME - len(suffix)] + suffix
        n += 1
    taken.add(title.lower())
    return title


def _num(v: Any) -> Any:
    """Coerce numeric-looking values; leave everything else alone."""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v if x not in (None, ""))
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def _pct_from_fraction(v: Any) -> Optional[float]:
    try:
        if v is None:
            return None
        return round(float(v) * 100.0, 4)
    except Exception:
        return None


def _safe_float(v: Any) -> Optional[float]:
    try:
        if v is None or v == "":
            return None
        return float(v)
    except Exception:
        return None


def _autosize(ws, min_width: int = 8) -> None:
    widths: Dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value)
            longest = max((len(part) for part in text.split("\n")), default=0)
            widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(_MAX_COL_WIDTH, w + 2))


def _write_table(ws, headers: Sequence[str], rows: Iterable[Sequence[Any]],
                 start_row: int = 1, number_formats: Optional[Dict[int, str]] = None,
                 freeze: bool = True) -> int:
    """Write a header + rows block. Returns the next free row index."""
    r = start_row
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    r += 1
    count = 0
    for row in rows:
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=_num(v))
            cell.font = _BODY_FONT
            if number_formats and c in number_formats and isinstance(cell.value, (int, float)):
                cell.number_format = number_formats[c]
        r += 1
        count += 1
    if count == 0:
        ws.cell(row=r, column=1, value="No rows for this campaign.").font = _MUTED_FONT
        r += 1
    if freeze and start_row == 1:
        ws.freeze_panes = ws.cell(row=2, column=1)
    return r


def _section(ws, row: int, label: str) -> int:
    cell = ws.cell(row=row, column=1, value=label)
    cell.font = _SECTION_FONT
    return row + 1


def _kv_block(ws, row: int, pairs: Sequence[Tuple[str, Any]]) -> int:
    for k, v in pairs:
        ws.cell(row=row, column=1, value=k).font = Font(name=_FONT, bold=True, size=10)
        c = ws.cell(row=row, column=2, value=_num(v))
        c.font = _BODY_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    return row


def _fmt_date(s: Any) -> str:
    try:
        d = _dt.date.fromisoformat(str(s)[:10])
        return d.strftime("%B %-d, %Y")
    except Exception:
        return str(s or "")


def _title_case_key(k: str) -> str:
    return re.sub(r"[_\s]+", " ", str(k)).strip().capitalize()


# ----------------------------------------------------------------------
# Sheet builders. Each takes the workbook + a context dict and never
# raises (the caller wraps them).
# ----------------------------------------------------------------------

def _sheet_campaign(wb, ws, ctx: dict) -> None:
    ov = ctx.get("overview") or {}
    snap = ctx.get("snapshot") or {}
    title = snap.get("title") or {}
    bc = ov.get("brand_config") or title.get("brand_config") or {}
    term = ov.get("terminology") or title.get("terminology") or {}
    is_brand = str(ov.get("title_type") or title.get("title_type") or "").lower() == "brand"
    launch_label = term.get("launch_date_label") or ("Campaign launch" if is_brand else "Opening date")

    ws.cell(row=1, column=1, value=ov.get("display_name") or title.get("display_name") or ctx.get("slug")).font = _TITLE_FONT
    ws.cell(row=2, column=1, value="Attribution IQ export").font = _MUTED_FONT
    row = 4
    pairs: List[Tuple[str, Any]] = [
        ("Campaign", ov.get("display_name") or title.get("display_name")),
        ("Partner", ov.get("distributor") or title.get("distributor")),
        ("Category", ov.get("genre") or title.get("genre") or bc.get("brand_category")),
        (launch_label, _fmt_date(ov.get("opening_date") or title.get("opening_date"))),
        ("Campaign end", _fmt_date(ov.get("end_date") or title.get("end_date"))),
        ("Assets tracked", ov.get("asset_count") or len(snap.get("assets") or [])),
        ("Response window (days)", bc.get("attribution_window_days")),
        ("US adults reached", bc.get("exposed_reach_us")),
        ("Exported", ctx.get("exported_label")),
    ]
    if ctx.get("as_of"):
        pairs.append(("As of", _fmt_date(ctx.get("as_of"))))
    if title.get("notes"):
        pairs.append(("About", title.get("notes")))
    row = _kv_block(ws, row, pairs)

    fs = snap.get("funnel_rates_summary") or {}
    if fs:
        row += 1
        row = _section(ws, row, "Response rates across every asset")
        row = _kv_block(ws, row, [
            ("Info-seek % (view weighted)", fs.get("view_weighted_issue_research_pct")),
            ("Action page % (view weighted)", fs.get("view_weighted_action_page_pct")),
            ("Info-seek % (average per asset)", fs.get("mean_issue_research_pct")),
            ("Action page % (average per asset)", fs.get("mean_action_page_pct")),
        ])

    months = ((bc.get("reach_basis") or {}).get("months")) or []
    if months:
        row += 1
        row = _section(ws, row, "Monthly reach")
        row = _kv_block(ws, row, [("Views per person per month", (bc.get("reach_basis") or {}).get("views_per_person_month"))])
        headers = ["Month", "Views", "People reached", "New to campaign %"]
        rows = [[m.get("month"), m.get("views"), m.get("monthly_people"),
                 _pct_from_fraction(m.get("new_to_campaign_share"))] for m in months]
        _write_table(ws, headers, rows, start_row=row,
                     number_formats={2: "#,##0", 3: "#,##0", 4: "0.0"}, freeze=False)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70
    for c in "CD":
        ws.column_dimensions[c].width = 18


def _sheet_phases(wb, ws, ctx: dict) -> None:
    ov = ctx.get("overview") or {}
    phases = ov.get("phases") or (ctx.get("snapshot") or {}).get("phases") or []
    phases = sorted(phases, key=lambda p: (p.get("phase_order") is None, p.get("phase_order") or 0))
    demo_phases = {p.get("phase_name"): p for p in (((ctx.get("snapshot") or {}).get("demographics") or {}).get("phases") or [])}
    headers = ["Order", "Phase", "Start", "End", "Assets", "Views", "What happened"]
    rows = []
    for p in phases:
        dp = demo_phases.get(p.get("phase_name")) or {}
        rows.append([p.get("phase_order"), p.get("phase_name"), p.get("start_date"), p.get("end_date"),
                     dp.get("asset_count"), dp.get("view_count"), p.get("description")])
    _write_table(ws, headers, rows, number_formats={5: "#,##0", 6: "#,##0"})
    _autosize(ws)


def _sheet_assets(wb, ws, ctx: dict) -> None:
    cards = ctx.get("assets") or []
    headers = ["Asset ID", "Creator / label", "Channel", "Asset type", "Phase", "Paid or organic",
               "Posted", "Views", "Engagements", "Engagement rate %", "Followers", "Likes", "Comments",
               "Shares", "Saves", "Info-seek %", "Action page %", "Funnel stage", "Talent", "Audience targets",
               "Caption", "URL"]
    rows = []
    for a in cards:
        det = a.get("ext_engagement_detail") or {}
        views = _safe_float(a.get("ext_view_count"))
        eng = _safe_float(a.get("ext_engagement_count"))
        er = round(eng / views * 100.0, 2) if views and eng is not None and views > 0 else None
        rows.append([
            a.get("asset_id"), a.get("action_label"), a.get("channel"), a.get("asset_type"), a.get("phase_name"),
            a.get("paid_or_organic"), a.get("posted_date"), a.get("ext_view_count"), a.get("ext_engagement_count"),
            er, det.get("followers"), det.get("likes"), det.get("comments"), det.get("shares"), det.get("bookmarks"),
            a.get("ext_info_seek_pct"), a.get("ext_website_visit_pct"), a.get("funnel_stage"),
            a.get("talent_tags"), a.get("audience_target_tags"), a.get("note"), a.get("url"),
        ])
    rows.sort(key=lambda r: (str(r[6] or ""), -(_safe_float(r[7]) or 0)))
    fmts = {8: "#,##0", 9: "#,##0", 10: "0.00", 11: "#,##0", 12: "#,##0", 13: "#,##0", 14: "#,##0",
            15: "#,##0", 16: "0.00", 17: "0.00"}
    _write_table(ws, headers, rows, number_formats=fmts)
    _autosize(ws)
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["U"].width = 60
    ws.column_dimensions["V"].width = 50


def _sheet_content_types(wb, ws, ctx: dict) -> None:
    q1 = ctx.get("q1") or {}
    rows_in = q1.get("rows") or []
    headers = ["Asset type", "Paid or organic", "Assets", "Average views (first 7 days)", "Median views (first 7 days)"]
    rows = [[r.get("asset_type"), r.get("paid_or_organic"), r.get("asset_count"), r.get("mean_views_7d"), r.get("median_views_7d")]
            for r in rows_in]
    rows.sort(key=lambda r: -(_safe_float(r[3]) or 0))
    _write_table(ws, headers, rows, number_formats={3: "#,##0", 4: "#,##0", 5: "#,##0"})
    _autosize(ws)


def _sheet_paid_vs_organic(wb, ws, ctx: dict) -> None:
    q2 = ctx.get("q2") or {}
    cum = q2.get("cumulative") or []
    headers = ["Date", "Paid or organic", "Assets posted that day", "Cumulative views to date"]
    rows = [[c.get("date"), c.get("paid_or_organic"), c.get("asset_drops"), c.get("cumulative_views_to_date")] for c in cum]
    _write_table(ws, headers, rows, number_formats={3: "#,##0", 4: "#,##0"})
    _autosize(ws)


def _asset_rows_block(items: List[dict]) -> Tuple[List[str], List[list]]:
    headers = ["Rank", "Asset ID", "Creator / label", "Channel", "Asset type", "Phase", "Paid or organic",
               "Posted", "Views", "Engagements", "Info-seek %", "Action page %", "URL"]
    rows = []
    for i, a in enumerate(items, start=1):
        rows.append([i, a.get("asset_id"), a.get("action_label"), a.get("channel"), a.get("asset_type"),
                     a.get("phase_name"), a.get("paid_or_organic"), a.get("posted_date"),
                     a.get("views_total") if a.get("views_total") is not None else a.get("ext_view_count"),
                     a.get("engagement_total") if a.get("engagement_total") is not None else a.get("ext_engagement_count"),
                     a.get("ext_info_seek_pct"), a.get("ext_website_visit_pct"), a.get("url")])
    return headers, rows


def _sheet_top_assets(wb, ws, ctx: dict) -> None:
    inf = ctx.get("in_flight") or {}
    q2 = ctx.get("q2") or {}
    row = 1
    label = inf.get("window_label") or "Campaign-to-date"
    row = _section(ws, row, f"Top assets by views, {label}")
    headers, rows = _asset_rows_block(inf.get("all_candidates") or q2.get("best_in_flight_ytd") or [])
    row = _write_table(ws, headers, rows, start_row=row, number_formats={9: "#,##0", 10: "#,##0", 11: "0.00", 12: "0.00"}, freeze=False)
    best_paid = inf.get("best_paid")
    best_org = inf.get("best_organic")
    if best_paid or best_org:
        row += 1
        row = _section(ws, row, "Best paid and best organic asset")
        headers, rows = _asset_rows_block([x for x in (best_paid, best_org) if x])
        for r in rows:
            r[0] = "Best paid" if r[6] == "paid" else ("Best organic" if r[6] == "organic" else r[0])
        headers[0] = "Which"
        _write_table(ws, headers, rows, start_row=row, number_formats={9: "#,##0", 10: "#,##0", 11: "0.00", 12: "0.00"}, freeze=False)
    _autosize(ws)
    ws.column_dimensions["C"].width = 60


def _mta_audience_index(ctx: dict) -> Dict[str, dict]:
    mta = ctx.get("mta") or {}
    out = {}
    for k, v in (mta.get("audiences") or {}).items():
        if isinstance(v, dict):
            out[k] = v
    return out


def _sheet_audiences(wb, ws, ctx: dict) -> None:
    cards = ctx.get("audiences") or []
    mta_aud = _mta_audience_index(ctx)
    noun = (ctx.get("mta") or {}).get("conversion_noun") or "action page"
    headers = ["Audience", "Group", "Share of campaign exposure %", "Share of US adults %", "Index vs US adults",
               "Accounts", f"Reached the {noun} %"]
    rows = []
    for a in cards:
        ov_bp = _safe_float(a.get("overlap_bp"))
        gp = _safe_float(a.get("gen_pop_share"))
        idx = round(ov_bp / gp * 100.0) if ov_bp is not None and gp else None
        m = mta_aud.get(a.get("subject_key")) or {}
        rows.append([a.get("display"), a.get("category"), ov_bp, gp, idx,
                     m.get("sample_size"), _pct_from_fraction(m.get("conversion_rate"))])
    rows.sort(key=lambda r: -(_safe_float(r[2]) or 0))
    _write_table(ws, headers, rows, number_formats={3: "0.00", 4: "0.00", 5: "0", 6: "#,##0", 7: "0.00"})
    _autosize(ws)


def _sheet_districts(wb, ws, ctx: dict) -> None:
    d = (ctx.get("snapshot") or {}).get("districts") or {}
    nat = d.get("national") or {}
    row = 1
    if d.get("as_of"):
        ws.cell(row=row, column=1, value=f"As of {_fmt_date(d.get('as_of'))}").font = _MUTED_FONT
        row += 1
    if nat:
        row = _section(ws, row, "National")
        row = _kv_block(ws, row, [
            ("US adults reached", nat.get("reached")),
            ("US adults", nat.get("us_adults")),
            ("Penetration %", nat.get("penetration_pct")),
            ("Researched the issue", nat.get("researched")),
            ("Researched %", nat.get("researched_pct")),
            ("Reached an action page", nat.get("acted")),
            ("Action page %", nat.get("action_pct")),
        ])
        row += 1
    row = _section(ws, row, "By congressional district")
    headers = ["District", "State", "Candidate", "Opponent", "Seat", "Adults 18+", "Reached", "Penetration %",
               "Index vs national", "Researched", "Researched %", "Reached an action page", "Action page %",
               "Named in a post", "Statewide post", "Post mentions", "Views on posts naming the district",
               "Placements naming the district", "First post naming the district"]
    rows = []
    for x in d.get("districts") or []:
        rows.append([x.get("district_label"), x.get("state"), x.get("candidate"), x.get("opponent"), x.get("seat_status"),
                     x.get("adults_18_plus"), x.get("reached"), x.get("penetration_pct"), x.get("index_vs_national"),
                     x.get("researched"), x.get("researched_pct"), x.get("acted"), x.get("action_pct"),
                     x.get("named_in_post"), x.get("statewide_post"), x.get("post_mentions"), x.get("named_post_views"),
                     x.get("named_placements"), x.get("first_named_post")])
    rows.sort(key=lambda r: -(_safe_float(r[8]) or 0))
    fmts = {6: "#,##0", 7: "#,##0", 8: "0.00", 9: "0", 10: "#,##0", 11: "0.00", 12: "#,##0", 13: "0.00",
            16: "#,##0", 17: "#,##0", 18: "#,##0"}
    row = _write_table(ws, headers, rows, start_row=row, number_formats=fmts, freeze=False)
    defs = d.get("definition") or {}
    if defs:
        row += 1
        row = _section(ws, row, "How to read")
        row = _kv_block(ws, row, [(_title_case_key(k), v) for k, v in defs.items()])
    _autosize(ws)
    ws.column_dimensions["B"].width = 40


def _sheet_cohorts(wb, ws, ctx: dict) -> None:
    co = ctx.get("cohorts") or {}
    headers = ["Cohort", "Definition", "Accounts", "Share of US adults %"]
    rows = [[c.get("display_name"), c.get("frequency_band"), c.get("panel_count"), c.get("gen_pop_share")]
            for c in (co.get("cohorts") or [])]
    _write_table(ws, headers, rows, number_formats={3: "#,##0", 4: "0.00"})
    _autosize(ws)


def _sheet_demographics(wb, ws, ctx: dict) -> None:
    demo = (ctx.get("snapshot") or {}).get("demographics") or ctx.get("demographics") or {}
    headers = ["Scope", "Category", "Bucket", "Share %", "Assets", "Views"]
    rows = []
    allc = demo.get("all_campaigns") or {}
    def _emit(scope: str, block: dict):
        dd = block.get("demographics") or {}
        for cat, buckets in dd.items():
            if not isinstance(buckets, dict):
                continue
            for bucket, pct in buckets.items():
                rows.append([scope, _title_case_key(cat), bucket, pct, block.get("asset_count"), block.get("view_count")])
    if allc:
        _emit("All campaigns", allc)
    for p in demo.get("phases") or []:
        _emit(p.get("phase_name") or "Phase", p)
    _write_table(ws, headers, rows, number_formats={4: "0.00", 5: "#,##0", 6: "#,##0"})
    _autosize(ws)


def _tp_rows(block: dict, scope: Optional[str] = None) -> List[list]:
    rows = []
    for t in block.get("touchpoints") or []:
        exp = _safe_float(t.get("exposed_n"))
        conv = _safe_float(t.get("converted_n"))
        rate = round(conv / exp * 100.0, 2) if exp and conv is not None and exp > 0 else None
        r = [t.get("asset_title"), t.get("channel"), t.get("phase"), t.get("paid_or_organic"),
             t.get("odds_ratio"), exp, conv, rate, t.get("significance")]
        if scope is not None:
            r.insert(0, scope)
        rows.append(r)
    return rows


_TP_HEADERS = ["Touchpoint", "Channel", "Phase", "Paid or organic", "Lift in response odds (1.0 = no change)",
               "Exposed accounts", "Responded accounts", "Response rate %", "Signal strength"]


def _sheet_touchpoints(wb, ws, ctx: dict) -> None:
    mta = ctx.get("mta") or {}
    overall = mta.get("overall") or {}
    row = 1
    noun = mta.get("conversion_noun") or "action page"
    row = _kv_block(ws, row, [
        ("Response measured as", f"Reached the {noun} within the response window"),
        ("Accounts in the read", overall.get("sample_size")),
        ("Overall response rate %", _pct_from_fraction(overall.get("conversion_rate"))),
        ("As of", _fmt_date(mta.get("as_of"))),
    ])
    row += 1
    row = _section(ws, row, "Every touchpoint")
    rows = _tp_rows(overall)
    rows.sort(key=lambda r: -(_safe_float(r[4]) or 0))
    _write_table(ws, _TP_HEADERS, rows, start_row=row,
                 number_formats={5: "0.000", 6: "#,##0", 7: "#,##0", 8: "0.00"}, freeze=False)
    _autosize(ws)
    ws.column_dimensions["A"].width = 70


def _journey_rows(block: dict, scope: Optional[str] = None) -> List[list]:
    rows = []
    for j in block.get("journeys") or []:
        tps = j.get("touchpoints") or []
        path = " > ".join(str((t or {}).get("asset_title") or "") for t in tps)
        chans = " > ".join(str((t or {}).get("channel") or "") for t in tps)
        r = [path, chans, j.get("path_length"), j.get("exposed_n"), j.get("converted_n"),
             _pct_from_fraction(j.get("conversion_rate")), j.get("lift_vs_baseline"), _pct_from_fraction(j.get("share_of_exposed"))]
        if scope is not None:
            r.insert(0, scope)
        rows.append(r)
    return rows


_JOURNEY_HEADERS = ["Journey (touchpoints in order)", "Channels", "Touchpoints", "Exposed accounts",
                    "Responded accounts", "Response rate %", "Lift vs campaign average", "Share of exposed %"]


def _sheet_journeys(wb, ws, ctx: dict) -> None:
    overall = (ctx.get("mta") or {}).get("overall") or {}
    rows = _journey_rows(overall)
    rows.sort(key=lambda r: -(_safe_float(r[6]) or 0))
    _write_table(ws, _JOURNEY_HEADERS, rows, number_formats={4: "#,##0", 5: "#,##0", 6: "0.00", 7: "0.000", 8: "0.00"})
    _autosize(ws)
    ws.column_dimensions["A"].width = 80


def _sheet_funnel(wb, ws, ctx: dict) -> None:
    mta = ctx.get("mta") or {}
    paths = ((mta.get("overall") or {}).get("paths")) or {}
    row = 1
    row = _kv_block(ws, row, [
        ("US adults", paths.get("us_gen_pop")),
        ("Accounts in the read", paths.get("panel_sample")),
    ])
    row += 1
    row = _section(ws, row, "Funnel")
    nest = paths.get("nest") or []
    rows = [[n.get("label"), n.get("us_accounts"), n.get("share_of_us_gen_pop_pct"), n.get("drop_from_prior")] for n in nest]
    row = _write_table(ws, ["Stage", "US accounts", "Share of US adults %", "Lost from prior stage"], rows,
                       start_row=row, number_formats={2: "#,##0", 3: "0.0000", 4: "#,##0"}, freeze=False)
    stage_label = {n.get("stage"): n.get("label") for n in nest}

    forks = paths.get("forks") or []
    if forks:
        row += 1
        row = _section(ws, row, "Splits at each stage")
        rows = [[stage_label.get(f.get("of_stage"), f.get("of_stage")), f.get("question"), f.get("yes"), f.get("no")] for f in forks]
        row = _write_table(ws, ["Stage", "Split", "Yes (US accounts)", "No (US accounts)"], rows,
                           start_row=row, number_formats={3: "#,##0", 4: "#,##0"}, freeze=False)

    where = paths.get("where") or {}
    for key, label in (("infoseek_overlap", "Where they researched"), ("ticketer_partition", "Where the response landed")):
        items = where.get(key) or []
        if not items:
            continue
        row += 1
        row = _section(ws, row, label)
        rows = [[w.get("surface"), w.get("us_accounts"), w.get("pct")] for w in items]
        row = _write_table(ws, ["Surface", "US accounts", "Share %"], rows, start_row=row,
                           number_formats={2: "#,##0", 3: "0.0"}, freeze=False)

    attr = paths.get("attribution") or {}
    for key, label in (("first_touch", "First touch"), ("last_touch", "Last touch"), ("assists", "Assists")):
        items = attr.get(key) or []
        if not items:
            continue
        row += 1
        row = _section(ws, row, label)
        rows = [[a.get("touchpoint"), a.get("us_accounts"), a.get("pct")] for a in items]
        row = _write_table(ws, ["Touchpoint", "US accounts", "Share %"], rows, start_row=row,
                           number_formats={2: "#,##0", 3: "0.0"}, freeze=False)

    ttc = paths.get("time_to_conversion") or []
    if ttc:
        row += 1
        row = _section(ws, row, "Time from exposure to response")
        rows = [[t.get("bucket"), t.get("us_accounts"), t.get("pct")] for t in ttc]
        row = _write_table(ws, ["When", "US accounts", "Share %"], rows, start_row=row,
                           number_formats={2: "#,##0", 3: "0.0"}, freeze=False)

    arche = paths.get("path_archetypes") or []
    if arche:
        row += 1
        row = _section(ws, row, "Path types")
        rows = [[a.get("archetype"), a.get("description"), a.get("us_accounts"), a.get("pct")] for a in arche]
        row = _write_table(ws, ["Path type", "What it means", "US accounts", "Share %"], rows, start_row=row,
                           number_formats={3: "#,##0", 4: "0.0"}, freeze=False)

    leaks = paths.get("leaks") or []
    if leaks:
        row += 1
        row = _section(ws, row, "Where people dropped out")
        rows = [[l.get("leak"), stage_label.get(l.get("of_base"), l.get("of_base")), l.get("us_accounts"), l.get("note")] for l in leaks]
        row = _write_table(ws, ["Drop-out", "From stage", "US accounts", "Detail"], rows, start_row=row,
                           number_formats={3: "#,##0"}, freeze=False)
    _autosize(ws)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 60


def _sheet_coexposure(wb, ws, ctx: dict) -> None:
    co = ((ctx.get("mta") or {}).get("overall") or {}).get("co_exposure") or {}
    tps = co.get("touchpoints") or []
    matrix = co.get("matrix") or []
    row = 1
    ws.cell(row=row, column=1, value="Share of accounts exposed to the row touchpoint who were also exposed to the column touchpoint.").font = _MUTED_FONT
    row += 2
    labels = [t.get("asset_title") or t.get("touchpoint_id") for t in tps]
    short = [f"T{i + 1}" for i in range(len(labels))]
    headers = ["Touchpoint", "Channel", "Exposure rate %"] + short
    rows = []
    for i, t in enumerate(tps):
        vals = matrix[i] if i < len(matrix) and isinstance(matrix[i], list) else []
        rows.append([f"{short[i]}: {labels[i]}", t.get("channel"), _pct_from_fraction(t.get("marginal_exposure_rate"))]
                    + [_pct_from_fraction(v) for v in vals])
    fmts = {3: "0.0"}
    for c in range(4, 4 + len(short)):
        fmts[c] = "0.0"
    _write_table(ws, headers, rows, start_row=row, number_formats=fmts, freeze=False)
    ws.freeze_panes = ws.cell(row=row + 1, column=2)
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 16
    for c in range(4, 4 + len(short)):
        ws.column_dimensions[get_column_letter(c)].width = 7


def _sheet_audience_touchpoints(wb, ws, ctx: dict) -> None:
    mta_aud = _mta_audience_index(ctx)
    rows = []
    for key, block in mta_aud.items():
        label = ((block.get("cohort_meta") or {}).get("audience_label")) or key
        rows.extend(_tp_rows(block, scope=label))
    rows.sort(key=lambda r: (str(r[0]), -(_safe_float(r[5]) or 0)))
    _write_table(ws, ["Audience"] + _TP_HEADERS, rows,
                 number_formats={6: "0.000", 7: "#,##0", 8: "#,##0", 9: "0.00"})
    _autosize(ws)
    ws.column_dimensions["B"].width = 70


def _sheet_audience_journeys(wb, ws, ctx: dict) -> None:
    mta_aud = _mta_audience_index(ctx)
    rows = []
    for key, block in mta_aud.items():
        label = ((block.get("cohort_meta") or {}).get("audience_label")) or key
        rows.extend(_journey_rows(block, scope=label))
    rows.sort(key=lambda r: (str(r[0]), -(_safe_float(r[7]) or 0)))
    _write_table(ws, ["Audience"] + _JOURNEY_HEADERS, rows,
                 number_formats={5: "#,##0", 6: "#,##0", 7: "0.00", 8: "0.000", 9: "0.00"})
    _autosize(ws)
    ws.column_dimensions["B"].width = 80


def _sheet_audience_funnels(wb, ws, ctx: dict) -> None:
    mta_aud = _mta_audience_index(ctx)
    headers = ["Audience", "Group", "Accounts", "Response rate %", "Stage", "US accounts", "Share of US adults %"]
    rows = []
    for key, block in mta_aud.items():
        meta = block.get("cohort_meta") or {}
        label = meta.get("audience_label") or key
        nest = ((block.get("paths") or {}).get("nest")) or []
        for n in nest:
            rows.append([label, meta.get("category"), block.get("sample_size"), _pct_from_fraction(block.get("conversion_rate")),
                         n.get("label"), n.get("us_accounts"), n.get("share_of_us_gen_pop_pct")])
    _write_table(ws, headers, rows, number_formats={3: "#,##0", 4: "0.00", 6: "#,##0", 7: "0.0000"})
    _autosize(ws)


# ----------------------------------------------------------------------
# Orchestration
# ----------------------------------------------------------------------

_SHEETS: List[Tuple[str, Callable, str]] = [
    ("Campaign", _sheet_campaign, "Campaign facts, response rates, monthly reach."),
    ("Phases", _sheet_phases, "Each campaign phase with dates, assets and views."),
    ("Assets", _sheet_assets, "Every tracked asset with views, engagement, info-seek % and action page %."),
    ("Content types", _sheet_content_types, "Views by asset type, paid vs organic."),
    ("Paid vs organic", _sheet_paid_vs_organic, "Cumulative views by day, paid vs organic."),
    ("Top assets", _sheet_top_assets, "Highest-view assets campaign-to-date, plus best paid and best organic."),
    ("Audiences", _sheet_audiences, "Every audience with share of exposure, index and response rate."),
    ("Districts", _sheet_districts, "Congressional district reach, research and action rates."),
    ("Cohorts", _sheet_cohorts, "Behavioral cohorts and their size."),
    ("Demographics", _sheet_demographics, "Who was reached, all campaigns and by phase."),
    ("Touchpoints", _sheet_touchpoints, "Every touchpoint's lift in response odds, exposure and response."),
    ("Journeys", _sheet_journeys, "The most common paths to response and how each performed."),
    ("Funnel", _sheet_funnel, "Funnel stages, splits, first and last touch, time to response, drop-outs."),
    ("Co-exposure", _sheet_coexposure, "Which touchpoints the same people saw together."),
    ("Audience funnels", _sheet_audience_funnels, "Funnel stages for each audience."),
    ("Audience touchpoints", _sheet_audience_touchpoints, "Touchpoint lift within each audience."),
    ("Audience journeys", _sheet_audience_journeys, "Paths to response within each audience."),
]


def _try(label: str, fn: Callable, *args, **kwargs):
    try:
        return fn(*args, **kwargs)
    except Exception as exc:  # pragma: no cover - defensive by design
        import logging
        logging.getLogger(__name__).warning("attribution export: %s failed: %s", label, exc)
        return None


def gather(slug: str, intent_iq, mta_iq=None, as_of: Optional[str] = None) -> dict:
    """Pull every dataset the workbook needs. Each pull is independent."""
    ctx: Dict[str, Any] = {"slug": slug, "as_of": as_of}
    ctx["overview"] = _try("overview", intent_iq.get_overview, slug) or {}
    snap = None
    loader = getattr(intent_iq, "_load_normalized_snapshot", None)
    if callable(loader):
        snap = _try("snapshot", loader, slug)
    ctx["snapshot"] = snap or {}
    assets = _try("assets", intent_iq.get_assets, slug, window="all") or {}
    cards = assets.get("cards") or []
    if not cards and ctx["snapshot"].get("assets"):
        cards = ctx["snapshot"]["assets"]
    ctx["assets"] = cards
    ctx["audiences"] = ((_try("audiences", intent_iq.get_audiences, slug) or {}).get("cards")) or ctx["snapshot"].get("audiences") or []
    ctx["cohorts"] = _try("cohorts", intent_iq.get_cohorts, slug) or {"cohorts": ctx["snapshot"].get("cohorts") or []}
    ctx["demographics"] = _try("demographics", intent_iq.get_demographics, slug) or {}
    ctx["q1"] = _try("q1", intent_iq.answer_question, slug, "q1") or {}
    ctx["q2"] = _try("q2", intent_iq.answer_question, slug, "q2") or {}
    ctx["in_flight"] = _try("in_flight", intent_iq.get_in_flight, slug, as_of=as_of, window="all") or {}
    tabs = (ctx["overview"].get("enabled_tabs") or {})
    ctx["mta"] = {}
    if mta_iq is not None and tabs.get("mta"):
        ctx["mta"] = _try("mta", mta_iq.compute_mta_coefficients, slug, as_of=as_of) or {}
    now = _dt.datetime.utcnow()
    ctx["exported_label"] = now.strftime("%B %-d, %Y")
    ctx["exported_stamp"] = now.strftime("%Y_%m_%d")
    return ctx


def build_workbook(slug: str, intent_iq, mta_iq=None, as_of: Optional[str] = None) -> Tuple[bytes, str]:
    """Return ``(xlsx_bytes, filename)`` for the campaign."""
    ctx = gather(slug, intent_iq, mta_iq=mta_iq, as_of=as_of)
    wb = Workbook()
    taken: set = set()

    # Read me first.
    ws0 = wb.active
    ws0.title = _sheet_title("Read me", taken)
    name = (ctx.get("overview") or {}).get("display_name") or ((ctx.get("snapshot") or {}).get("title") or {}).get("display_name") or slug
    ws0.cell(row=1, column=1, value=f"{name}: Attribution IQ export").font = _TITLE_FONT
    ws0.cell(row=2, column=1, value=f"Exported {ctx['exported_label']}. Every sheet is the full campaign-to-date read.").font = _MUTED_FONT
    r = 4
    r = _write_table(ws0, ["Sheet", "What is on it"], [[s[0], s[2]] for s in _SHEETS], start_row=r, freeze=False)
    r += 1
    r = _section(ws0, r, "Definitions")
    noun = (ctx.get("mta") or {}).get("conversion_noun") or "action page"
    win = ((ctx.get("overview") or {}).get("brand_config") or {}).get("attribution_window_days")
    win_txt = f"{win} days" if win else "the response window"
    r = _kv_block(ws0, r, [
        ("Info-seek %", "Of the people who saw the asset, the share who searched or read about the issue within " + win_txt + "."),
        ("Action page %", f"Of the people who saw the asset, the share who reached the {noun} within {win_txt}."),
        ("Index", "100 = the US adult average. 200 means twice as common as the US average."),
        ("Lift in response odds", "1.0 means seeing the touchpoint did not change the odds of responding. 1.25 means the odds went up by a quarter."),
        ("Accounts", "Every count is an individual account, not a household."),
    ])
    ws0.column_dimensions["A"].width = 28
    ws0.column_dimensions["B"].width = 110

    for title, fn, _desc in _SHEETS:
        ws = wb.create_sheet(_sheet_title(title, taken))
        try:
            fn(wb, ws, ctx)
        except Exception as exc:  # pragma: no cover - defensive by design
            import logging
            logging.getLogger(__name__).warning("attribution export: sheet %s failed: %s", title, exc)
            for row in ws.iter_rows():
                for cell in row:
                    cell.value = None
            ws.cell(row=1, column=1, value="This sheet could not be filled for this campaign.").font = _MUTED_FONT

    buf = io.BytesIO()
    wb.save(buf)
    safe = re.sub(r"[^A-Za-z0-9]+", "_", str(name)).strip("_") or "Campaign"
    filename = f"{safe}_Attribution_IQ_{ctx['exported_stamp']}.xlsx"
    return buf.getvalue(), filename


__all__ = ["build_workbook", "gather"]
