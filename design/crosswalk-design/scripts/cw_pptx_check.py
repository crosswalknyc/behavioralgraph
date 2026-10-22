#!/usr/bin/env python3
"""Crosswalk PPTX font-safety gate. Usage: python3 cw_pptx_check.py deck.pptx [--fonts DIR_WITH_CROSSWALK_INTER_TTFS] [--no-render]
Exit 1 on anything that can make text fall back to Calibri, re-wrap unpredictably, or draw twice.
From the crosswalk-pptx-font-safety skill. Changed here: a doubled word at a line break only fails when the
render shows that word more times than the slide's own text contains it, and the duplicate-layer test needs
real overlap and several shared content words (common words like "the" no longer trip it)."""
import sys, os, re, zipfile, subprocess, tempfile, shutil, itertools
from pptx import Presentation
from pptx.util import Emu

HOUSE = "Crosswalk Inter"
ALLOWED_OTHER = {"Geist Mono", "Arial"}   # Geist Mono: raw clickstream displays only. Arial: bullet glyphs only.
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
problems, warnings = [], []
bad, warn = problems.append, warnings.append

path = sys.argv[1]
fonts_dir = sys.argv[sys.argv.index("--fonts") + 1] if "--fonts" in sys.argv else None
z = zipfile.ZipFile(path)

# 1. Theme fonts: anything that inherits must land on the house face, never Calibri
for n in z.namelist():
    if n.startswith("ppt/theme/") and n.endswith(".xml"):
        x = z.read(n).decode("utf8", "ignore")
        for slot in ("majorFont", "minorFont"):
            m = re.search(rf'<a:{slot}>\s*<a:latin typeface="([^"]*)"', x)
            if m and not m.group(1).startswith(HOUSE):
                bad(f"{n}: theme {slot} is '{m.group(1)}', must be '{HOUSE}'")

# 2. Embedding: flags set, and every Crosswalk Inter face the deck uses is embedded in full
p = z.read("ppt/presentation.xml").decode("utf8", "ignore")
if 'embedTrueTypeFonts="1"' not in p: bad('presentation.xml: embedTrueTypeFonts="1" missing')
if 'saveSubsetFonts="1"' in p: bad("presentation.xml: saveSubsetFonts is on; fonts must embed all characters")
slide_xml = "".join(z.read(n).decode("utf8", "ignore") for n in z.namelist()
                    if re.match(r"ppt/(slides|slideLayouts|slideMasters|notesSlides)/[^/]+\.xml$", n))
used = set()
for m in re.finditer(r"<a:(?:rPr|defRPr|endParaRPr)\b([^>]*)>(.*?)</a:(?:rPr|defRPr|endParaRPr)>", slide_xml, re.S):
    lat = re.search(r'<a:latin typeface="([^"]*)"', m.group(2))
    if lat and lat.group(1).startswith(HOUSE):
        b, i = ' b="1"' in m.group(1), ' i="1"' in m.group(1)
        used.add((lat.group(1), "boldItalic" if b and i else "bold" if b else "italic" if i else "regular"))
emb = set()
for fm in re.finditer(r'<p:embeddedFont>\s*<p:font typeface="([^"]*)"[^>]*/>(.*?)</p:embeddedFont>', p, re.S):
    for slot in re.findall(r"<p:(regular|bold|italic|boldItalic) ", fm.group(2)):
        emb.add((fm.group(1), slot))
if not emb:
    (bad if fonts_dir else warn)("No fonts embedded. Run cw_pptx_embed.py with the Crosswalk Inter font folder.")
else:
    for u in sorted(used - emb):
        bad(f"face used but not embedded: {u[0]} {u[1]}")

prs = Presentation(path)
W, H = prs.slide_width, prs.slide_height
words = lambda t: re.findall(r"\S+", t.lower())
norm = lambda t: re.sub(r"[^\w%$]+", "", t.lower())
STOP = set("the and for with that this from what your their they them have more than into over each when were been only also most much many some about which while where".split())
# source word counts per slide, from every text run on the slide (groups and tables included)
src_counts = []
for slide in prs.slides:
    c = {}
    for t in slide._element.iter(f"{A}t"):
        for w_ in (t.text or "").split():
            k = norm(w_)
            if k:
                c[k] = c.get(k, 0) + 1
    src_counts.append(c)

for si, slide in enumerate(prs.slides, 1):
    tshapes = []
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        tf, txt = sh.text_frame, sh.text_frame.text
        tshapes.append(sh)
        short = re.sub(r"\s+", " ", txt[:40]).strip()   # computed outside the f-string: Python 3.11 (Render) forbids a backslash inside one
        tag = f"slide {si} '{short}'"
        # 3. every run names its font explicitly
        for para in tf.paragraphs:
            for r in para.runs:
                if not r.text.strip():
                    continue
                lat = r._r.find(f".//{A}latin")
                face = lat.get("typeface") if lat is not None else None
                if face is None:
                    bad(f"{tag}: run '{r.text[:20]}' has no explicit typeface (inherits the theme)")
                elif not (face.startswith(HOUSE) or face in ALLOWED_OTHER):
                    bad(f"{tag}: run uses '{face}'")
        # 4. doubled words inside one box
        w = words(txt)
        for a, b in zip(w, w[1:]):
            if a == b and re.search(r"[a-z]", a):
                bad(f"{tag}: repeated word '{a} {b}'")
        # 5. pre-broken lines
        xml = tf._txBody.xml
        paras = [pp.text.strip() for pp in tf.paragraphs]
        if "<a:br/>" in xml or "\v" in txt:
            warn(f"{tag}: manual line break inside a text box. Allowed only as a deliberate headline break.")
        for a, b in zip(paras, paras[1:]):
            if a and b and not re.search(r"[.!?:;)”\"]$", a) and b[0].islower():
                bad(f"{tag}: one sentence split across paragraphs ('{a[-15:]}' / '{b[:15]}'). Use one paragraph and let it wrap.")
        # 6. shrink-on-overflow is recomputed differently on every machine
        if "<a:normAutofit" in xml:
            warn(f"{tag}: shrink-on-overflow autofit. Prefer a fixed size with height headroom.")
        if sh.left + sh.width > W + Emu(9525) or sh.top + sh.height > H + Emu(9525):
            bad(f"{tag}: text box extends off the slide")
    # 7. overlapping text boxes sharing words = a duplicate text layer
    for a, b in itertools.combinations(tshapes, 2):
        ix = min(a.left + a.width, b.left + b.width) - max(a.left, b.left)
        iy = min(a.top + a.height, b.top + b.height) - max(a.top, b.top)
        if ix > 0 and iy > 0 and ix * iy > 0.5 * min(a.width * a.height, b.width * b.height):
            common = {c for c in set(map(norm, words(a.text_frame.text))) & set(map(norm, words(b.text_frame.text)))
                      if len(c) > 3 and c not in STOP}
            if len(common) >= 3:
                bad(f"slide {si}: overlapping text boxes share words {sorted(common)[:4]} (a duplicate layer; it shows twice when the font changes)")

# 8. Render the deck as a machine WITHOUT Crosswalk Inter sees it, then look for doubled words
def render(alias, label):
    tmp = tempfile.mkdtemp()
    conf = os.path.join(tmp, "fonts.conf")
    extra = f"<dir>{os.path.abspath(fonts_dir)}</dir>" if (fonts_dir and alias is None) else ""
    rule = "" if alias is None else (
        f'<match target="pattern"><test name="family" compare="contains"><string>{HOUSE}</string></test>'
        f'<edit name="family" mode="assign" binding="strong"><string>{alias}</string></edit></match>')
    open(conf, "w").write('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>'
                          f'<include ignore_missing="yes">/etc/fonts/fonts.conf</include>{extra}{rule}</fontconfig>')
    env = dict(os.environ, FONTCONFIG_FILE=conf, HOME=tmp)
    shutil.copy(path, os.path.join(tmp, "d.pptx"))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, os.path.join(tmp, "d.pptx")],
                   env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    pdf = os.path.join(tmp, "d.pdf")
    if not os.path.exists(pdf):
        warn(f"{label}: render failed, check skipped")
        return
    import pdfplumber
    with pdfplumber.open(pdf) as doc:
        for pi, page in enumerate(doc.pages, 1):
            ws = page.extract_words()
            rendered = {}
            for wd in ws:
                k = norm(wd["text"])
                rendered[k] = rendered.get(k, 0) + 1
            source = src_counts[pi - 1] if pi - 1 < len(src_counts) else {}
            rows = {}
            for wd in ws:
                rows.setdefault(round(wd["top"]), []).append(wd)
            keys = sorted(rows)
            for r1, r2 in zip(keys, keys[1:]):
                if not (0 < r2 - r1 < 40):
                    continue
                for w2 in rows[r2]:
                    above = [w for w in rows[r1] if w2["x0"] - 20 <= w["x0"] <= w2["x0"] + 400]
                    if not above or w2["x0"] > min(w["x0"] for w in above) + 2:
                        continue
                    tail = max(above, key=lambda w: w["x1"])
                    k = norm(w2["text"])
                    if (norm(tail["text"]) == k and len(k) > 2 and re.search(r"[a-z]", k)
                            and rendered.get(k, 0) > source.get(k, 0)):
                        bad(f"{label}, page {pi}: '{w2['text']}' ends one line and starts the next, and appears more often than the slide's text contains it")
            stacked, where = 0, ""
            spots = {}
            for c in page.chars:
                if not c["text"].strip():
                    continue
                k = (c["text"], round(c["x0"]), round(c["top"]))
                if k in spots:
                    stacked += 1
                    where = where or page.crop((max(0, c["x0"] - 60), max(0, c["top"] - 2), min(page.width, c["x0"] + 120), min(page.height, c["bottom"] + 2))).extract_text()
                spots[k] = 1
            if stacked > 3:
                bad(f"{label}, page {pi}: text drawn twice in the same spot, {stacked} characters (stacked text layers) near '{(where or '').strip()[:40]}'")
            for wd in ws:
                if wd["x1"] > page.width + 1 or wd["bottom"] > page.height + 1:
                    bad(f"{label}, page {pi}: text runs off the slide ('{wd['text']}')")

if "--no-render" not in sys.argv:   # --no-render: structural checks only (fast); run the full check before delivery
    render("Carlito", "Fallback render, no Crosswalk Inter (Calibri metrics)")
    render("Liberation Sans", "Fallback render, no Crosswalk Inter (Arial metrics)")
    if fonts_dir:
        render(None, "House-font render")

for w in warnings: print("WARN ", w)
for b in problems: print("FAIL ", b)
print(f"\n{len(problems)} failure(s), {len(warnings)} warning(s)")
sys.exit(1 if problems else 0)
