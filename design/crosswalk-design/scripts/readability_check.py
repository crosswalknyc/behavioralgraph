#!/usr/bin/env python3
"""Plain English check for Crosswalk decks and copy.

Usage:
    python3 readability_check.py deck.pptx
    python3 readability_check.py draft.md
    python3 readability_check.py page.html
    python3 readability_check.py report.docx
    python3 readability_check.py --text "Header line here"

Reports:
    - headers over 12 words (pptx: slide title placeholders; text:
      markdown headings or the first line of each block)
    - clever-header patterns (fragment cadences, colon reveals,
      questions, common idioms)
    - sentences over 30 words
    - Flesch-Kincaid grade for body copy (target 10 or lower)
    - jargon words with a plain substitute
    - words that are fine once but not on repeat ("actually")
    - template tells (no-formulaic-deck-prose.mdc, 2026-10-19):
      three or more cards on one page that open the same way, headers
      that repeat an opener, and phrases of five or more words that
      recur across cards or pages
    - door / doors metaphor (no-doors-in-decks.mdc, 2026-10-28)

Exit code 0 when there are no flags, 1 otherwise. No dependencies
beyond the standard library; python-pptx is used only when a .pptx
is passed, python-docx only for a .docx.

2026-10-21: one script for both Crosswalk skills (crosswalk-design and
crosswalk-plain-english). It carries the design pack's .html / .docx
readers and overuse list together with the formulaic-prose checks.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

_HERE = Path(__file__).resolve().parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))
try:
    from cw_door_copy import door_hits
except Exception:  # pragma: no cover
    def door_hits(text: str) -> list[str]:
        return [m.group(0) for m in re.finditer(r"\bdoors?\b", text or "", re.I)]

HEADER_WORD_CEILING = 12
SENTENCE_WORD_CEILING = 30
GRADE_CEILING = 10.0

JARGON = {
    "leverage": "use",
    "leverages": "uses",
    "leveraging": "using",
    "utilize": "use",
    "utilizes": "uses",
    "utilization": "use",
    "optimize": "improve",
    "optimizing": "improving",
    "robust": "strong",
    "holistic": "whole",
    "synergy": "fit",
    "synergies": "fits",
    "ecosystem": "market",
    "paradigm": "model",
    "granular": "detailed",
    "actionable": "useful",
    "learnings": "lessons",
    "ideate": "come up with",
    "operationalize": "put in place",
    "incentivize": "encourage",
    "attrited": "left",
    "attrition": "people leaving",
    "monetize": "make money from",
    "seamless": "smooth",
    "cutting-edge": "new",
    "best-in-class": "best",
    "delve": "look",
    "landscape": "market",
    "tapestry": "mix",
    "testament": "sign",
    "arguably": "",
    "essentially": "",
    "significant proportion": "half / a third / most",
    "a number of": "several",
    "in order to": "to",
    "it is worth noting": "",
    "at the end of the day": "",
}

IDIOMS = [
    "leaky bucket", "long tail", "low-hanging fruit", "move the needle",
    "tip of the iceberg", "double-edged sword", "elephant in the room",
    "perfect storm", "tale of two", "the good, the bad", "game changer",
    "north star", "secret sauce", "boil the ocean", "at the end of the day",
    "wags", "hides", "whisper", "roars", "quietly", "loudly",
]

STOP_TITLES = {"agenda", "appendix", "thank you", "contents", "sources"}

# Fine once in a piece; flagged when used more often than this.
OVERUSE = {
    "actually": 1,
}

# Voice checks (no-formulaic-deck-prose.mdc). A "card" is a written unit
# of at least this many words that starts like a sentence and ends with
# terminal punctuation; stat captions, chart labels, and chrome are not
# cards and are never compared.
CARD_MIN_WORDS = 8
SAME_SHAPE_MIN = 3        # cards on one page opening the same way
OPENER_REPEAT_MIN = 3     # headers sharing a two-word opener, or consecutive first words
PHRASE_N = 5              # phrase length compared across units
PHRASE_FLAG_MIN = 3       # distinct units carrying the phrase -> FLAG
PHRASE_INFO_MIN = 2       # distinct units carrying the phrase -> info line
INFO_LINES_SHOWN = 10

STOPWORDS = {
    "the", "a", "an", "of", "in", "on", "at", "to", "for", "and", "or", "but",
    "with", "from", "by", "as", "is", "are", "was", "were", "be", "been", "it",
    "its", "that", "this", "these", "those", "who", "which", "what", "her",
    "his", "their", "she", "he", "they", "them", "we", "our", "you", "your",
    "not", "no", "than", "then", "so", "if", "into", "over", "under", "about",
    "after", "before", "up", "down", "out", "off", "all", "any", "each",
    "every", "more", "most", "much", "many", "some", "such", "same", "other",
    "one", "two", "has", "have", "had", "do", "does", "did", "can", "will",
    "would", "should", "could", "there", "here", "when", "where", "while",
    "how", "also", "just", "only", "very", "per", "vs", "within", "between",
}

EYEBROW_RE = re.compile(r"^\s*\d{2}\s*[\u2014\u2013-]\s*[A-Z0-9 ,'&/:]+$")


def syllables(word: str) -> int:
    w = re.sub(r"[^a-z]", "", word.lower())
    if not w:
        return 0
    if len(w) <= 3:
        return 1
    w = re.sub(r"(?:[^laeiouy]es|ed|[^laeiouy]e)$", "", w)
    w = re.sub(r"^y", "", w)
    groups = re.findall(r"[aeiouy]{1,2}", w)
    return max(1, len(groups))


def sentences_of(text: str) -> list[str]:
    parts = re.split(r"(?<=[.!?])\s+", text.strip())
    return [p.strip() for p in parts if p.strip()]


def words_of(text: str) -> list[str]:
    return [w for w in re.findall(r"[A-Za-z0-9$%'.,-]+", text) if re.search(r"[A-Za-z0-9]", w)]


FOOTER_RE = re.compile(r"CROSSWALK\s*/\s*BEHAVIORAL INTELLIGENCE\s*\d*|CONFIDENTIAL\s*\d*", re.I)


def is_prose(sentence: str) -> bool:
    """True when a sentence reads as prose rather than a table, label run, or clickstream row."""
    if not re.search(r"[.!?]$", sentence.strip()):
        return False
    if "·" in sentence or re.search(r"\b[a-z0-9-]+\.(com|tv|net|org|io)\b", sentence, re.I):
        return False
    ws = words_of(sentence)
    if len(ws) < 5:
        return False
    numeric = sum(1 for w in ws if re.search(r"\d", w))
    if numeric / len(ws) > 0.25:
        return False
    lower_words = sum(1 for w in ws if re.match(r"[a-z]", w))
    return lower_words / len(ws) >= 0.6


def fk_grade(text: str) -> float | None:
    sents = sentences_of(text)
    words = words_of(text)
    if not sents or len(words) < 10:
        return None
    syl = sum(syllables(w) for w in words)
    return round(0.39 * (len(words) / len(sents)) + 11.8 * (syl / len(words)) - 15.59, 1)


def clever_flags(header: str) -> list[str]:
    flags = []
    h = header.strip()
    low = h.lower()
    if h.endswith("?"):
        flags.append("rhetorical question")
    frags = [f for f in re.split(r"(?<=[.!?])\s+", h) if f.strip()]
    frags = [f for f in frags if not re.fullmatch(r"[\d$.,%]+", f.strip())]
    if len(frags) >= 2:
        if all(len(f.split()) <= 3 for f in frags):
            flags.append("fragment cadence")
        else:
            flags.append(f"{len(frags)} sentences; a header is one sentence")
    if ":" in h and not re.search(r"\d:\d", h):
        flags.append("colon reveal")
    for idiom in IDIOMS:
        if idiom in low:
            flags.append(f"idiom or metaphor: '{idiom}'")
    if re.search(r"\bisn't\b|\bain't\b", low):
        flags.append("wordplay contraction")
    return flags


Page = tuple[str, list[str]]  # (label, written units on that page, title excluded)


def load_pptx(path: Path) -> tuple[list[str], list[str], list[Page]]:
    from pptx import Presentation  # type: ignore

    prs = Presentation(str(path))
    headers, bodies, pages = [], [], []
    for idx, slide in enumerate(prs.slides, start=1):
        title_text = None
        body_parts = []

        def biggest(shape):
            try:
                return max([r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size] or [0])
            except Exception:
                return 0
        # Fallback title: the largest type on the page (20pt or more, 25 words or fewer). Works for
        # 16:9 decks (34pt titles) and portrait documents (23pt headlines) alike.
        cands = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()
                 and len(sh.text_frame.text.split()) <= 25 and biggest(sh) >= 20]
        fallback = max(cands, key=biggest) if cands else None
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            txt = re.sub(r"\s+", " ", shape.text_frame.text).strip()  # line breaks (\v) and paragraphs read as spaces
            if not txt:
                continue
            is_title = False
            try:
                ph = shape.placeholder_format
                is_title = ph.type in (1, 3)  # TITLE, CENTER_TITLE
            except Exception:
                is_title = False
            if not is_title and title_text is None and shape is fallback:
                is_title = True
            if is_title and title_text is None:
                title_text = txt
            else:
                body_parts.append(txt)
        if title_text:
            headers.append(f"slide {idx}: {title_text}")
        if body_parts:
            # Each text box is its own unit. Closing it with a period keeps
            # stacked labels and stat tiles from fusing into one fake sentence.
            closed = [p if p.endswith((".", "!", "?")) else p + "." for p in body_parts]
            bodies.append(" ".join(closed))
            pages.append((f"slide {idx}", body_parts))
    return headers, bodies, pages


def load_html(text: str) -> tuple[list[str], list[str], list[Page]]:
    """Headers from h1 to h3, body units from p / li / td / th / captions;
    each h1 to h3 opens a new page for the voice checks."""
    from html.parser import HTMLParser

    class P(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.headers, self.pages, self.buf, self.kind, self.skip = [], [], [], None, 0
            self.label, self.units = "top", []

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style", "svg", "noscript"):
                self.skip += 1
            elif tag in ("h1", "h2", "h3"):
                self.flush(); self.kind = "h"
            elif tag in ("p", "li", "td", "th", "figcaption", "blockquote", "h4", "h5", "h6"):
                self.flush(); self.kind = "b"
            elif tag == "br":
                self.buf.append(" ")

        def handle_endtag(self, tag):
            if tag in ("script", "style", "svg", "noscript"):
                self.skip = max(0, self.skip - 1)
            elif tag in ("h1", "h2", "h3", "p", "li", "td", "th", "figcaption", "blockquote", "h4", "h5", "h6"):
                self.flush()

        def handle_data(self, data):
            if not self.skip and self.kind:
                self.buf.append(data)

        def flush(self):
            t = re.sub(r"\s+", " ", "".join(self.buf)).strip()
            if t:
                if self.kind == "h":
                    self.headers.append(t)
                    if self.units:
                        self.pages.append((self.label, self.units))
                    self.label, self.units = f"section '{t[:40]}'", []
                else:
                    self.units.append(t if t.endswith((".", "!", "?")) else t + ".")
            self.buf, self.kind = [], None

        def close_all(self):
            self.flush()
            if self.units:
                self.pages.append((self.label, self.units))

    p = P()
    p.feed(text)
    p.close_all()
    bodies = [" ".join(u for _l, us in p.pages for u in us)] if p.pages else []
    return p.headers, bodies, p.pages


def load_docx(path: Path) -> tuple[list[str], list[str], list[Page]]:
    """Heading / Title styles are headers; every other paragraph is a body
    unit under the most recent heading."""
    from docx import Document  # type: ignore

    headers, pages = [], []
    label, units = "top", []
    for para in Document(str(path)).paragraphs:
        t = re.sub(r"\s+", " ", para.text).strip()
        if not t:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith(("heading", "title")):
            headers.append(t)
            if units:
                pages.append((label, units))
            label, units = f"section '{t[:40]}'", []
        else:
            units.append(t if t.endswith((".", "!", "?", ":")) else t + ".")
    if units:
        pages.append((label, units))
    bodies = [" ".join(u for _l, us in pages for u in us)] if pages else []
    return headers, bodies, pages


def load_text(text: str) -> tuple[list[str], list[str], list[Page]]:
    headers, bodies, pages = [], [], []
    label, units = "top", []
    for block in re.split(r"\n\s*\n", text):
        lines = [l.rstrip() for l in block.strip().splitlines() if l.strip()]
        if not lines:
            continue
        first = lines[0]
        header = None
        if first.startswith("#"):
            header = first.lstrip("# ").strip()
            rest = lines[1:]
        elif len(lines) > 1 and len(first.split()) <= 14 and not first.endswith(","):
            header = first
            rest = lines[1:]
        else:
            rest = lines
        if header is not None:
            headers.append(header)
            if units:
                pages.append((label, units))
            label, units = f"section '{header[:40]}'", []
        if rest:
            bodies.append(" ".join(rest))
            # A bullet or numbered line is its own unit (a card); plain
            # lines in one block form one paragraph unit.
            para = []
            for line in rest:
                if re.match(r"^\s*(?:[-*\u2022]|\d+[.)])\s+", line):
                    if para:
                        units.append(" ".join(para))
                        para = []
                    units.append(re.sub(r"^\s*(?:[-*\u2022]|\d+[.)])\s+", "", line))
                else:
                    para.append(line)
            if para:
                units.append(" ".join(para))
    if units:
        pages.append((label, units))
    return headers, bodies, pages


def is_chrome(unit: str) -> bool:
    """Source lines, footer lockups, eyebrows: never compared for voice."""
    u = unit.strip()
    return bool(u.lower().startswith("source") or FOOTER_RE.search(u) or EYEBROW_RE.match(u))


def is_card(unit: str) -> bool:
    """A written unit worth comparing for shape: sentence-like, long enough, not chrome."""
    u = unit.strip()
    if len(words_of(u)) < CARD_MIN_WORDS:
        return False
    if not re.search(r"[.!?]$", u):
        return False
    if not re.match(r"[A-Z0-9$]", u):
        return False
    if is_chrome(u):
        return False
    letters = re.sub(r"[^A-Za-z]", "", u)
    if letters and sum(1 for c in letters if c.isupper()) / len(letters) > 0.6:
        return False  # an all-caps label run
    return True


GENERIC_OPENERS = {"the", "a", "an", "this", "that", "these", "those", "it", "there"}


def opener_class(text: str) -> str:
    """How a unit opens: 'a number' or its first word. An article or
    demonstrative is too common to count alone, so it takes the next
    word with it ('the field'); a pronoun or name stands on its own."""
    toks = words_of(text)
    if not toks:
        return ""
    first = toks[0]
    if re.search(r"\d", first):
        return "a number"
    w = re.sub(r"[^a-z']", "", first.lower())
    if w in GENERIC_OPENERS and len(toks) > 1:
        second = re.sub(r"[^a-z0-9']", "", toks[1].lower())
        return f"{w} {second}" if second else w
    return w


def two_word_opener(text: str) -> str:
    toks = [re.sub(r"[^a-z0-9']", "", t.lower()) for t in words_of(text)[:2]]
    toks = [t for t in toks if t]
    return " ".join(toks) if len(toks) == 2 else ""


def phrase_tokens(text: str) -> list[str]:
    toks = [re.sub(r"[^a-z0-9']", "", t.lower()) for t in words_of(text)]
    return [t for t in toks if t]


def voice_flags(headers: list[str], pages: list[Page]) -> tuple[list[str], list[str]]:
    """Template tells: same-shape cards per page, repeated header openers,
    phrases recurring across units. Returns (flags, info)."""
    flags: list[str] = []
    info: list[str] = []

    # Same shape: three or more cards on one page opening the same way.
    for label, units in pages:
        cards = [u for u in units if is_card(u)]
        if len(cards) < SAME_SHAPE_MIN:
            continue
        by_opener: dict[str, list[str]] = {}
        for c in cards:
            by_opener.setdefault(opener_class(c), []).append(c)
        for opener, group in by_opener.items():
            if not opener or len(group) < SAME_SHAPE_MIN:
                continue
            counts = {len(sentences_of(c)) for c in group}
            how = f"{len(group)} cards open on {opener if opener == 'a number' else repr(opener)}"
            if len(counts) == 1:
                n = counts.pop()
                how += f", {n} sentence{'s' if n != 1 else ''} each"
            flags.append(f"{label}: {how}. Same shape card to card; rewrite so no three share a skeleton.")

    # Repeated opener across headers.
    titles = []
    for h in headers:
        _, _, text = h.partition(": ") if h.startswith("slide ") else ("", "", h)
        text = (text or h).strip()
        if text.lower().rstrip(".") in STOP_TITLES:
            continue
        titles.append(text)
    two_words: dict[str, list[str]] = {}
    for t in titles:
        key = two_word_opener(t)
        if key:
            two_words.setdefault(key, []).append(t)
    flagged_sets: list[set[str]] = []
    for key, group in two_words.items():
        if len(group) >= OPENER_REPEAT_MIN:
            flags.append(f"{len(group)} headers open '{key} ...': " + " | ".join(g[:50] for g in group))
            flagged_sets.append(set(group))
    run, prev = [], None
    for t in titles + [""]:
        first = opener_class(t) if t else ""
        if first and first == prev:
            run.append(t)
        else:
            if len(run) >= OPENER_REPEAT_MIN and not any(set(run) <= s for s in flagged_sets):
                flags.append(f"{len(run)} consecutive headers open on {repr(prev)}: " + " | ".join(r[:50] for r in run))
            run = [t] if t else []
        prev = first

    # Repeated phrases across units. Headers, cards, and long stat captions
    # count; chrome (source lines, footer, eyebrows), short table or chart
    # labels, and running lines repeated near-verbatim on two or more
    # pages (a base or window line under every chart) do not.
    running: dict[str, set[str]] = {}
    for label, page_units in pages:
        for u in page_units:
            key = re.sub(r"\d+", "#", " ".join(phrase_tokens(u)))
            running.setdefault(key, set()).add(label)
    units: list[str] = list(titles)
    for _, page_units in pages:
        for u in page_units:
            n = len(words_of(u))
            if n < CARD_MIN_WORDS or is_chrome(u):
                continue
            if len(running.get(re.sub(r"\d+", "#", " ".join(phrase_tokens(u))), ())) >= 2:
                continue
            if is_card(u) or n >= CARD_MIN_WORDS + 2:
                units.append(u)
    grams: dict[tuple[str, ...], set[int]] = {}
    for idx, u in enumerate(units):
        toks = phrase_tokens(u)
        for i in range(len(toks) - PHRASE_N + 1):
            g = tuple(toks[i:i + PHRASE_N])
            if any(re.search(r"\d", t) for t in g):
                continue
            if sum(1 for t in g if t not in STOPWORDS) < 2:
                continue
            grams.setdefault(g, set()).add(idx)
    repeated = [(g, s) for g, s in grams.items() if len(s) >= PHRASE_INFO_MIN]
    for phrase, count, merged in merged_phrases(repeated):
        line = f"'{phrase}'{' (or most of it)' if merged else ''} appears in {count} places"
        if count >= PHRASE_FLAG_MIN:
            flags.append(line + ". A reused phrase reads as a template; say it once.")
        else:
            info.append(line)
    return flags, info


def merged_phrases(repeated: list[tuple[tuple[str, ...], set[int]]]) -> list[tuple[str, int, bool]]:
    """Join overlapping repeated n-grams that recur in the same places into
    one phrase each, so 'the top rate in the' and 'top rate in the cast'
    report once as 'the top rate in the cast'. Returns (phrase, count,
    merged_from_several), strongest repeats first."""
    n = len(repeated)
    parent = list(range(n))

    def find(x: int) -> int:
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    for i in range(n):
        gi, si = repeated[i]
        for j in range(i + 1, n):
            gj, sj = repeated[j]
            if len(si & sj) >= PHRASE_INFO_MIN and (gi[1:] == gj[:-1] or gj[1:] == gi[:-1]):
                parent[find(i)] = find(j)
    clusters: dict[int, list[int]] = {}
    for i in range(n):
        clusters.setdefault(find(i), []).append(i)

    out: list[tuple[str, int, bool]] = []
    for members in clusters.values():
        gs = {repeated[i][0] for i in members}
        starts = sorted(g for g in gs if not any(h != g and h[1:] == g[:-1] for h in gs))
        cur = starts[0] if starts else sorted(gs)[0]
        phrase, used = list(cur), {cur}
        while True:
            nxt = sorted(h for h in gs if h not in used and h[:-1] == cur[1:])
            if not nxt:
                break
            cur = nxt[0]
            used.add(cur)
            phrase.append(cur[-1])
        count = max(len(repeated[i][1]) for i in members)
        out.append((" ".join(phrase), count, len(members) > 1))
    out.sort(key=lambda pc: (-pc[1], pc[0]))
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", nargs="?", help=".pptx, .docx, .html, .md, or .txt")
    ap.add_argument("--text", help="check a single string (treated as a header)")
    args = ap.parse_args()

    if args.text:
        headers, bodies, pages = [args.text], [], []
    elif args.path:
        p = Path(args.path)
        if p.suffix.lower() == ".pptx":
            headers, bodies, pages = load_pptx(p)
        elif p.suffix.lower() == ".docx":
            headers, bodies, pages = load_docx(p)
        elif p.suffix.lower() in (".html", ".htm"):
            headers, bodies, pages = load_html(p.read_text(encoding="utf-8", errors="replace"))
        else:
            headers, bodies, pages = load_text(p.read_text(encoding="utf-8"))
    else:
        ap.print_help()
        return 2

    flags = 0

    print("HEADERS")
    for h in headers:
        label, _, text = h.partition(": ") if h.startswith("slide ") else ("", "", h)
        text = text or h
        if text.strip().lower().rstrip(".") in STOP_TITLES:
            continue
        n = len(text.split())
        issues = []
        if n > HEADER_WORD_CEILING:
            issues.append(f"{n} words (ceiling {HEADER_WORD_CEILING})")
        issues.extend(clever_flags(text))
        for snippet in door_hits(text):
            issues.append(f"door metaphor: '{snippet}'")
        for j in JARGON:
            if re.search(rf"\b{re.escape(j)}\b", text.lower()):
                issues.append(f"jargon '{j}'" + (f" -> '{JARGON[j]}'" if JARGON[j] else " -> cut"))
        mark = "FLAG" if issues else " ok "
        flags += len(issues)
        prefix = f"{label}: " if label else ""
        print(f"  [{mark}] {prefix}{text}" + (f"    <- {'; '.join(issues)}" if issues else ""))
    if not headers:
        print("  (none found)")

    print("\nBODY")
    all_body = FOOTER_RE.sub(" ", " ".join(bodies))
    prose_sents = [s for s in sentences_of(all_body) if is_prose(s)]
    long_sents = [s for s in prose_sents if len(words_of(s)) > SENTENCE_WORD_CEILING]
    for s in long_sents:
        flags += 1
        print(f"  [FLAG] {len(words_of(s))} words: {s[:110]}{'...' if len(s) > 110 else ''}")
    prose_text = " ".join(prose_sents)
    grade = fk_grade(prose_text)
    if grade is not None:
        ok = grade <= GRADE_CEILING
        if not ok:
            flags += 1
        print(f"  [{' ok ' if ok else 'FLAG'}] Flesch-Kincaid grade {grade} on prose (target <= {GRADE_CEILING})")
    if prose_sents:
        avg = sum(len(words_of(s)) for s in prose_sents) / len(prose_sents)
        print(f"  [info] {len(prose_sents)} prose sentences, average {avg:.1f} words (tables and labels skipped)")
    jargon_hits = {}
    low_body = all_body.lower()
    for j, sub in JARGON.items():
        c = len(re.findall(rf"\b{re.escape(j)}\b", low_body))
        if c:
            jargon_hits[j] = (c, sub)
    for j, (c, sub) in sorted(jargon_hits.items(), key=lambda kv: -kv[1][0]):
        flags += 1
        print(f"  [FLAG] '{j}' x{c}" + (f" -> '{sub}'" if sub else " -> cut"))
    all_copy = (" ".join(headers) + " " + all_body).lower()
    for w, limit in OVERUSE.items():
        c = len(re.findall(rf"\b{re.escape(w)}\b", all_copy))
        if c > limit:
            flags += 1
            print(f"  [FLAG] '{w}' x{c} -> {limit} at most; cut the rest")
    if not bodies:
        print("  (none found)")

    print("\nDOORS")
    d_hits = door_hits(" ".join(headers) + " " + all_body)
    for snippet in d_hits:
        flags += 1
        print(f"  [FLAG] door metaphor: '{snippet}'")
    if not d_hits:
        print("  [ ok ] no door / doors in copy")

    print("\nVOICE (one analyst wrote this)")
    v_flags, v_info = voice_flags(headers, pages)
    for line in v_flags:
        flags += 1
        print(f"  [FLAG] {line}")
    for line in v_info[:INFO_LINES_SHOWN]:
        print(f"  [info] {line}")
    if len(v_info) > INFO_LINES_SHOWN:
        print(f"  [info] ... and {len(v_info) - INFO_LINES_SHOWN} more phrases used twice; read them before shipping")
    if not v_flags and not v_info:
        print("  [ ok ] no same-shape cards, repeated openers, or repeated phrases" if pages or len(headers) > 1
              else "  (nothing to compare)")

    print(f"\n{flags} flag(s). {'Ready.' if flags == 0 else 'Rewrite before shipping.'}")
    return 0 if flags == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
