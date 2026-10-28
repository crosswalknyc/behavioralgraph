#!/usr/bin/env python3
"""Decks never use the word doors (Jenna 2026-10-28).

The defect is the metaphor: the door to new viewers, opens the door,
a door for, walk through the door, close the door, doorway to. Compound
brand tokens (DoorDash, Glassdoor, NextDoor) do not match. The only
allowlist is a proper name that already contains Door (The Door, The
Door NYC) when it is not followed by to/for.

``scrub_door_metaphor`` rewrites common frames and strips any leftover
door token. Builders, the brand lint, and ``design_skill.finish`` all
call it. A finish never holds the file: the rewrite lands in place.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

DOOR_TOKEN = re.compile(r"\bdoors?\b", re.I)
# Title-case venue name only. "the door to" stays a metaphor.
ALLOWED_PROPER = re.compile(r"\bThe Door(?:\s+NYC)?\b")

_FRAMES = [
    (re.compile(r"\bopens?\s+the\s+doors?\s+to\s+", re.I), "brings "),
    (re.compile(r"\bopen(?:ed|ing)\s+the\s+doors?\s+to\s+", re.I), "brings "),
    (re.compile(r"\bopens?\s+doors?\s+to\s+", re.I), "brings "),
    (re.compile(r"\bopen(?:ed|ing)\s+doors?\s+to\s+", re.I), "brings "),
    (re.compile(r"\bopens?\s+the\s+doors?\b", re.I), "brings"),
    (re.compile(r"\bopen(?:ed|ing)\s+the\s+doors?\b", re.I), "brings"),
    (re.compile(r"\bopens?\s+doors?\b", re.I), "brings"),
    (re.compile(r"\bopen(?:ed|ing)\s+doors?\b", re.I), "brings"),
    (re.compile(r"\ba\s+doors?\s+for\s+", re.I), ""),
    (re.compile(r"\bdoors?\s+for\s+", re.I), ""),
    (re.compile(r"\bwalk(?:s|ed|ing)?\s+through\s+(?:the\s+)?doors?\b", re.I), "arrives"),
    (re.compile(r"\bthrough\s+the\s+doors?\b", re.I), ""),
    (re.compile(r"\bclos(?:e|es|ed|ing)\s+the\s+doors?\b", re.I), ""),
    (re.compile(r"\b(?:the\s+)?doorways?\s+to\s+", re.I), ""),
    (re.compile(r"\b(?:the\s+)?doorways?\b", re.I), ""),
    (re.compile(r"\b(?:the\s+)?doors?\s+to\s+", re.I), ""),
]


def _is_allowed_span(text: str, start: int, end: int) -> bool:
    for m in ALLOWED_PROPER.finditer(text or ""):
        if m.start() <= start and end <= m.end():
            if re.match(r"\s+(to|for)\b", (text or "")[m.end():], re.I):
                return False
            return True
    return False


def door_hits(text: str) -> list[str]:
    """Short snippets for each disallowed door / doors token."""
    text = text or ""
    hits = []
    for m in DOOR_TOKEN.finditer(text):
        if _is_allowed_span(text, m.start(), m.end()):
            continue
        a = max(0, m.start() - 16)
        b = min(len(text), m.end() + 20)
        hits.append(re.sub(r"\s+", " ", text[a:b]).strip())
    return hits


def has_door_metaphor(text: str) -> bool:
    return bool(door_hits(text))


def _protect(text: str) -> tuple[str, list[str]]:
    held: list[str] = []

    def hold(m: re.Match) -> str:
        if re.match(r"\s+(to|for)\b", text[m.end():], re.I):
            return m.group(0)
        held.append(m.group(0))
        return f"\x00DOOR{len(held) - 1}\x00"

    return ALLOWED_PROPER.sub(hold, text or ""), held


def _restore(text: str, held: list[str]) -> str:
    for i, name in enumerate(held):
        text = text.replace(f"\x00DOOR{i}\x00", name)
    return text


def _tidy(text: str, orig: str) -> str:
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r" +\n", "\n", text)
    text = re.sub(r"\n +", "\n", text)
    text = re.sub(r" +([,.;:!?])", r"\1", text)
    text = re.sub(r"\(\s+", "(", text)
    text = re.sub(r"\s+\)", ")", text)
    text = re.sub(r"\s{2,}", " ", text)
    lead = re.match(r"^\s*", text).group(0)
    body = text[len(lead):]
    if body:
        body = re.sub(
            r"^((?:\d+[.)]?\s+)*)([a-z])",
            lambda m: m.group(1) + m.group(2).upper(),
            body,
            count=1,
        )
        o = (orig or "").lstrip()
        if o[:1].isupper() and body[:1].islower() and not o[:1].isdigit():
            body = body[:1].upper() + body[1:]
    return (lead + body).strip() if (orig or "").strip() == (orig or "") else lead + body


def scrub_door_metaphor(text: str) -> str:
    """Rewrite door metaphors. The returned string has no door / doors token
    unless it is an allowed proper name."""
    if text is None:
        return ""
    if not isinstance(text, str):
        text = str(text)
    if not text:
        return text
    protected, held = _protect(text)
    out = protected
    for pat, repl in _FRAMES:
        out = pat.sub(repl, out)
    def drop(m: re.Match) -> str:
        if _is_allowed_span(out, m.start(), m.end()):
            return m.group(0)
        return ""
    out = DOOR_TOKEN.sub(drop, out)
    out = _restore(out, held)
    return _tidy(out, text)


def _scrub_paragraph_runs(paragraph) -> int:
    runs = list(getattr(paragraph, "runs", []) or [])
    raw = "".join((r.text or "") for r in runs)
    if not raw and getattr(paragraph, "text", None):
        raw = paragraph.text or ""
    new = scrub_door_metaphor(raw)
    if new == raw:
        return 0
    if runs:
        runs[0].text = new
        for r in runs[1:]:
            r.text = ""
    else:
        try:
            paragraph.text = new
        except Exception:
            return 0
    return 1


def _scrub_text_frame(tf) -> int:
    n = 0
    try:
        for p in tf.paragraphs:
            n += _scrub_paragraph_runs(p)
    except Exception:
        pass
    return n


def _iter_shapes(shapes):
    for sh in shapes:
        yield sh
        try:
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    for cell in row.cells:
                        if cell.text_frame is not None:
                            yield cell
        except Exception:
            pass
        try:
            if sh.shape_type is not None and int(sh.shape_type) == 6:
                yield from _iter_shapes(sh.shapes)
        except Exception:
            pass


def scrub_pptx(path) -> int:
    from pptx import Presentation  # type: ignore

    prs = Presentation(str(path))
    n = 0
    for slide in prs.slides:
        for sh in _iter_shapes(slide.shapes):
            if getattr(sh, "has_text_frame", False) and sh.text_frame is not None:
                n += _scrub_text_frame(sh.text_frame)
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide
                if notes.notes_text_frame is not None:
                    n += _scrub_text_frame(notes.notes_text_frame)
                for sh in _iter_shapes(notes.shapes):
                    if getattr(sh, "has_text_frame", False) and sh.text_frame is not None:
                        n += _scrub_text_frame(sh.text_frame)
        except Exception:
            pass
    if n:
        prs.save(str(path))
    return n


def scrub_docx(path) -> int:
    from docx import Document  # type: ignore

    doc = Document(str(path))
    n = 0
    for p in doc.paragraphs:
        n += _scrub_paragraph_runs(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    n += _scrub_paragraph_runs(p)
    if n:
        doc.save(str(path))
    return n


def scrub_text_file(path) -> int:
    p = Path(path)
    raw = p.read_text(encoding="utf-8", errors="replace")
    new = scrub_door_metaphor(raw)
    if new == raw:
        return 0
    p.write_text(new, encoding="utf-8")
    return 1


def scrub_file(path) -> int:
    """Rewrite door metaphors in a deck or document in place. Returns
    the number of rewritten units. Never raises; 0 on any failure."""
    try:
        ext = Path(str(path)).suffix.lower()
        if ext == ".pptx":
            return scrub_pptx(path)
        if ext == ".docx":
            return scrub_docx(path)
        if ext in (".html", ".htm", ".md", ".txt"):
            return scrub_text_file(path)
    except Exception:
        return 0
    return 0


def main(argv: list[str] | None = None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    if not args:
        print("usage: cw_door_copy.py --text STRING | FILE")
        return 2
    if args[0] == "--text":
        src = " ".join(args[1:])
        print(scrub_door_metaphor(src))
        return 0 if not has_door_metaphor(scrub_door_metaphor(src)) else 1
    n = scrub_file(args[0])
    print(f"rewrote {n} unit(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
