#!/usr/bin/env python3
"""Work from the bundled Crosswalk PowerPoint templates instead of rebuilding them.

  python3 cw_pptx_tools.py list FILE.pptx [--slide N]
      Every text box on every slide (or one slide): shape id, name, position in inches, and each
      paragraph's runs with their font, size and color. Use it to see what to replace.

  python3 cw_pptx_tools.py pick TEMPLATE.pptx OUT.pptx 1,7,13,13,27
      New deck from template slides, in that order. A slide number may repeat: the second use is a
      true duplicate (images and background carried over). Unpicked slides are dropped.

  python3 cw_pptx_tools.py replace FILE.pptx EDITS.json [OUT.pptx]
      Replace text while keeping each run's font, weight, size and color. EDITS.json:
        [ {"slide": 1, "find": "[Client]", "replace": "Live Nation"},
          {"slide": 2, "shape": 7, "paragraph": 0, "runs": ["Fans who stayed ", "streamed twice as long."],
           "weights": ["bold", "light"]},
          {"slide": "*", "find": "[MONTH YEAR]", "replace": "October 2026"} ]
      "find"/"replace" works inside a run, and across runs and line breaks of one paragraph (write the
      break as a space; the result keeps the first run's formatting and drops the break). "runs" sets
      the runs of one paragraph in order; each new run copies the formatting of the run it replaces.
      "weights" (optional, one per run) sets each run's face: thin, extralight, light, regular, medium,
      semibold, bold, extrabold, black. That is how an all-bold template title becomes a weight pair.
      Every edit must match at least once or the command fails.

  python3 cw_pptx_tools.py delete FILE.pptx 3,5 [OUT.pptx]
      Remove slides (page-number fields renumber themselves).

  python3 cw_pptx_tools.py merge BASE.pptx EXTRA.pptx OUT.pptx [3,7]
      Append slides from EXTRA (all, or the listed ones in that order) to the end of BASE: for example
      template pages added to the person's own deck. Shapes, text, images and backgrounds come across;
      charts and embedded objects do not (the command names them; rebuild those). Reorder afterwards
      with pick (OUT.pptx OUT.pptx 1,2,9,3,...). Slide sizes must match.

  python3 cw_pptx_tools.py leftovers FILE.pptx
      Every template placeholder and annotation still in the file ([Client], [FEE], "TEMPLATE 07 — ...",
      "IMAGE · 13.333 x 7.500") and every shape parked off the page, with slide and shape. Exit status 1
      if any remain. For a deck it also lists text still word for word as in the deck templates (sample
      copy you may have meant to replace); that list is a prompt to check, not a failure.

  python3 cw_pptx_tools.py sow-products SOW.pptx "Profile IQ,Trends" [OUT.pptx] [--keep-unchecked]
      In a file made from Crosswalk_SOW_Template.pptx: marks the listed products as included (the
      template's filled box), deletes the other product rows and closes the gap (or, with
      --keep-unchecked, leaves them as unchecked menu rows), resizes the table panel and removes the
      helper shapes parked off the page.

pick and merge also rebuild the template's gradient backgrounds, graphite wash and scrim at full
resolution (cw_backgrounds.py), so a deck made from the templates is 4K and free of banding.

Finish every deck with: python3 cw_embed_fonts.py OUT.pptx
"""
import copy, io, json, re, sys
from pptx import Presentation
from pptx.util import Emu

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from walk(sh.shapes)
        else:
            yield sh


def cmd_list(path, only=None):
    prs = Presentation(path)
    for i, sl in enumerate(prs.slides, 1):
        if only and i != only:
            continue
        print(f"=== slide {i}  (layout: {sl.slide_layout.name})")
        for sh in walk(sl.shapes):
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                if getattr(sh, "has_table", False) and sh.has_table:
                    print(f"  [{sh.shape_id}] TABLE {sh.name}: " + " / ".join(c.text for c in sh.table.rows[0].cells))
                continue
            pos = f"{Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f}"
            print(f"  [{sh.shape_id}] {sh.name}  @ {pos}")
            for pi, p in enumerate(sh.text_frame.paragraphs):
                if not p.text.strip():
                    continue
                runs = []
                for r in p.runs:
                    lat = r._r.find(f".//{A}latin")
                    face = lat.get("typeface") if lat is not None else "theme"
                    col = r._r.find(f".//{A}solidFill/{A}srgbClr")
                    sz = r.font.size.pt if r.font.size else "-"
                    runs.append(f"«{r.text}» {face}{' bold' if r.font.bold else ''} {sz}pt {col.get('val') if col is not None else ''}".rstrip())
                print(f"     p{pi}: " + " + ".join(runs))


def duplicate(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rel in list(src.part.rels.values()):
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        if rel.is_external:
            rid_map[rel.rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rel.rId] = new.part.relate_to(rel.target_part, rel.reltype)
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        tree.append(copy.deepcopy(el))
    bg = src._element.cSld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    for el in new._element.iter():
        for attr in (R + "embed", R + "id", R + "link"):
            v = el.get(attr)
            if v in rid_map:
                el.set(attr, rid_map[v])
    return new


def upgrade_backgrounds(path):
    """Template gradients, wash and scrim -> full-resolution, dithered versions (cw_backgrounds.py)."""
    import os
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    try:
        import cw_backgrounds
        n = len(cw_backgrounds.upgrade(path))
        if n:
            print(f"  {n} background image(s) rebuilt at full resolution")
    except Exception as e:  # never block a pick on this; lint will flag low-res images
        print("  note: could not rebuild backgrounds at full resolution:", e)


def cmd_pick(tpl, out, spec):
    order = [int(x) for x in spec.split(",") if x.strip()]
    prs = Presentation(tpl)
    slides = list(prs.slides)
    if any(n < 1 or n > len(slides) for n in order):
        sys.exit(f"slide numbers must be 1 to {len(slides)}")
    lst = prs.slides._sldIdLst
    ids = list(lst)
    used, chosen = set(), []
    for n in order:
        if n not in used:
            used.add(n)
            chosen.append(ids[n - 1])
        else:
            duplicate(prs, slides[n - 1])
            chosen.append(list(lst)[-1])
    for el in list(lst):
        lst.remove(el)
    for el in ids:
        if el not in chosen:
            prs.part.drop_rel(el.get(R + "id"))
    for el in chosen:
        lst.append(el)
    prs.save(out)
    print(f"{out}: {len(chosen)} slides from {tpl} ({spec})")
    upgrade_backgrounds(out)


FACES = {"thin": ("Crosswalk Inter Thin", False), "extralight": ("Crosswalk Inter ExtraLight", False),
         "light": ("Crosswalk Inter Light", False), "regular": ("Crosswalk Inter", False),
         "medium": ("Crosswalk Inter Medium", False), "semibold": ("Crosswalk Inter SemiBold", False),
         "bold": ("Crosswalk Inter", True), "extrabold": ("Crosswalk Inter ExtraBold", False),
         "black": ("Crosswalk Inter Black", False)}


def set_face(r, weight):
    face, bold = FACES[weight.lower()]
    rPr = r._r.get_or_add_rPr()
    if bold:
        rPr.set("b", "1")
    elif "b" in rPr.attrib:
        del rPr.attrib["b"]
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(f"{A}{tag}")
        if el is None:
            el = rPr.makeelement(f"{A}{tag}", {})
            # schema order: ln, fills, effects, highlight, uLnTx.., latin, ea, cs, sym, hlinkClick...
            after = [c for c in rPr if c.tag.split("}")[1] in ("sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst")]
            if after:
                after[0].addprevious(el)
            else:
                rPr.append(el)
        el.set("typeface", face)
        for k in ("panose", "pitchFamily", "charset"):
            el.attrib.pop(k, None)


def set_runs(p, texts):
    runs = p.runs
    if not runs:
        raise ValueError("paragraph has no runs")
    while len(runs) < len(texts):  # clone the last run's formatting
        new = copy.deepcopy(runs[-1]._r)
        runs[-1]._r.addnext(new)
        runs = p.runs
    for r, t in zip(runs, texts):
        r.text = t
    for r in runs[len(texts):]:
        r._r.getparent().remove(r._r)


def replace_in_paragraph(p, find, rep):
    hit = False
    for r in p.runs:
        if find in r.text:
            r.text = r.text.replace(find, rep)
            hit = True
    if hit:
        return True
    text = p.text                       # line breaks read as \v
    if find in text:
        full = text.replace(find, rep)
    elif find in text.replace("\v", " "):
        full = text.replace("\v", " ").replace(find, rep)
    else:
        return False
    runs = p.runs
    first = runs[0]._r
    for el in list(p._p):
        if el is not first and el.tag in (A + "r", A + "br"):
            p._p.remove(el)
    parts = full.split("\v")
    runs[0].text = parts[0]
    prev = first
    for part in parts[1:]:              # a break the replacement kept stays a real break
        br = first.makeelement(A + "br", {})
        rpr = first.find(A + "rPr")
        if rpr is not None:
            br.append(copy.deepcopy(rpr))
        prev.addnext(br)
        r = copy.deepcopy(first)
        r.find(A + "t").text = part
        br.addnext(r)
        prev = r
    return True


def cmd_replace(path, edits_path, out=None):
    prs = Presentation(path)
    edits = json.load(open(edits_path))
    slides = list(prs.slides)
    for e in edits:
        targets = slides if e.get("slide") == "*" else [slides[int(e["slide"]) - 1]]
        hits = 0
        for sl in targets:
            for sh in walk(sl.shapes):
                if "shape" in e and sh.shape_id != int(e["shape"]):
                    continue
                frames = []
                if sh.has_text_frame:
                    frames.append(sh.text_frame)
                if getattr(sh, "has_table", False) and sh.has_table:
                    frames += [c.text_frame for row in sh.table.rows for c in row.cells]
                for tf in frames:
                    for pi, p in enumerate(tf.paragraphs):
                        if "runs" in e:
                            if pi == int(e.get("paragraph", 0)) and "shape" in e:
                                set_runs(p, e["runs"]); hits += 1
                                for r, w in zip(p.runs, e.get("weights", [])):
                                    set_face(r, w)
                        elif replace_in_paragraph(p, e["find"], e["replace"]):
                            hits += 1
        if not hits:
            sys.exit(f"edit matched nothing: {json.dumps(e)}")
    prs.save(out or path)
    print(f"{out or path}: {len(edits)} edit(s) applied")


def cmd_delete(path, spec, out=None):
    prs = Presentation(path)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for n in sorted({int(x) for x in spec.split(",")}, reverse=True):
        el = ids[n - 1]
        prs.part.drop_rel(el.get(R + "id"))
        lst.remove(el)
    prs.save(out or path)
    print(f"{out or path}: removed slide(s) {spec}")


PLACEHOLDER = re.compile(r"\[[A-Z][^\]]{0,40}\](?!\()|^TEMPLATE \d{1,2}\b.*|\b(IMAGE|HALF BLEED|FULL BLEED)\s+·.*|\bXX+\b|\bLorem ipsum\b", re.M)


def cmd_leftovers(path):
    prs = Presentation(path)
    found = 0
    for i, sl in enumerate(prs.slides, 1):
        for sh in walk(sl.shapes):
            texts = []
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                texts += [c.text for row in sh.table.rows for c in row.cells]
            for t in texts:
                for m in PLACEHOLDER.finditer(t):
                    print(f"slide {i} [{sh.shape_id}]: {m.group(0)[:70]}")
                    found += 1
        for sh in sl.shapes:
            if sh.left is None or sh.width is None:
                continue
            if sh.left >= prs.slide_width or sh.left + sh.width <= 0 or sh.top >= prs.slide_height or sh.top + sh.height <= 0:
                label = sh.text_frame.text.strip().replace("\v", " ").replace("\n", " ")[:50] if sh.has_text_frame else sh.name
                print(f"slide {i} [{sh.shape_id}]: shape parked off the page «{label}»")
                found += 1
        if sl.has_notes_slide:
            for m in PLACEHOLDER.finditer(sl.notes_slide.notes_text_frame.text if sl.notes_slide.notes_text_frame else ""):
                print(f"slide {i} notes: {m.group(0)[:70]}")
                found += 1
    print(f"{found} placeholder(s) left" if found else "no placeholders left")
    if prs.slide_width > prs.slide_height:
        import os
        tpl = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "templates", "Crosswalk_Deck_Templates.pptx")
        if os.path.exists(tpl) and os.path.abspath(tpl) != os.path.abspath(path):
            sample = set()
            for sl in Presentation(tpl).slides:
                for sh in walk(sl.shapes):
                    if sh.has_text_frame:
                        t = re.sub(r"\s+", " ", sh.text_frame.text).strip()
                        if len(t.split()) >= 3 and not t.startswith("CROSSWALK /"):
                            sample.add(t)
            same = []
            for i, sl in enumerate(prs.slides, 1):
                for sh in walk(sl.shapes):
                    if sh.has_text_frame:
                        t = re.sub(r"\s+", " ", sh.text_frame.text).strip()
                        if t in sample:
                            same.append(f"slide {i} [{sh.shape_id}]: «{t[:70]}»")
            if same:
                print("\nStill word for word as in the deck templates (check each one is meant to stay):")
                print("\n".join(same))
    return found == 0


def cmd_merge(base, extra, out, spec=None):
    prs, src = Presentation(base), Presentation(extra)
    if (prs.slide_width, prs.slide_height) != (src.slide_width, src.slide_height):
        sys.exit("slide sizes differ; merge only decks of the same size")
    layout = min(prs.slide_layouts, key=lambda l: len(l.placeholders))
    order = [int(x) for x in spec.split(",")] if spec else range(1, len(src.slides) + 1)
    sslides = list(src.slides)
    skipped = []
    for n in order:
        s = sslides[n - 1]
        new = prs.slides.add_slide(layout)
        for shp in list(new.shapes):
            shp._element.getparent().remove(shp._element)
        rid_map = {}
        for rel in list(s.part.rels.values()):
            if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
                continue
            if rel.is_external:
                rid_map[rel.rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            elif rel.reltype.endswith("/image"):
                try:
                    _, rid = new.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                    rid_map[rel.rId] = rid
                except Exception:
                    skipped.append(f"slide {n}: image {rel.target_part.partname} (unsupported format)")
            else:
                skipped.append(f"slide {n}: {rel.reltype.rsplit('/', 1)[-1]} {rel.target_part.partname}")
        tree = new.shapes._spTree
        for el in s.shapes._spTree.iterchildren():
            if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
                continue
            tree.append(copy.deepcopy(el))
        bg = s._element.cSld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
        if bg is not None:
            new._element.cSld.insert(0, copy.deepcopy(bg))
        for el in new._element.iter():
            for attr in (R + "embed", R + "id", R + "link"):
                v = el.get(attr)
                if v in rid_map:
                    el.set(attr, rid_map[v])
    prs.save(out)
    print(f"{out}: {len(prs.slides)} slides ({len(list(order))} appended from {extra})")
    upgrade_backgrounds(out)
    for k in skipped:
        print("NOT CARRIED:", k, "- rebuild this element on the merged slide")


def cmd_sow_products(path, names, out=None, keep=False):
    prs = Presentation(path)
    want = [n.strip().lower() for n in names.split(",") if n.strip()]
    E = 914400
    slide = None
    for sl in prs.slides:
        if any(sh.has_text_frame and "PLATFORM AND TOOL ACCESS" in sh.text_frame.text for sh in sl.shapes):
            slide = sl
            break
    if slide is None:
        sys.exit("no PLATFORM AND TOOL ACCESS table found; is this file from Crosswalk_SOW_Template.pptx?")
    shapes = list(slide.shapes)
    W = prs.slide_width
    helpers = [sh for sh in shapes if sh.left >= W]
    green = next((sh for sh in helpers if not sh.has_text_frame or not sh.text_frame.text.strip()), None)
    boxes = [sh for sh in shapes if sh.left < W and sh.width < 0.2 * E and sh.height < 0.2 * E
             and abs(sh.width - sh.height) < 0.01 * E and sh.top > 4.9 * E]
    rows = []
    for b in sorted(boxes, key=lambda b: b.top):
        row = [sh for sh in shapes if sh.left < W and abs(sh.top - b.top) < 0.05 * E]
        name = next((sh for sh in row if sh.has_text_frame and sh.text_frame.text.strip() and sh.left < 2.0 * E), None)
        rows.append({"box": b, "shapes": row, "name": name.text_frame.text.strip() if name else "?", "top": b.top})
    known = [r["name"].lower() for r in rows]
    missing = [n for n in want if n not in known]
    if missing:
        sys.exit(f"not in the table: {missing}. Rows are: {[r['name'] for r in rows]}")
    pitch = rows[1]["top"] - rows[0]["top"]
    lines = sorted([sh for sh in shapes if sh.left < W and sh.height < 0.02 * E and sh.width > 5 * E
                    and rows[0]["top"] < sh.top < rows[-1]["top"]], key=lambda s: s.top)
    panel = max((sh for sh in shapes if sh.left < W and sh.height > 3 * E and sh.top < rows[0]["top"]),
                key=lambda s: s.height)
    tree = slide.shapes._spTree
    next_id = max(sh.shape_id for sh in shapes) + 1
    kept = rows if keep else [r for r in rows if r["name"].lower() in want]
    for r in rows:
        if r not in kept:
            for sh in r["shapes"]:
                tree.remove(sh._element)
    for k, r in enumerate(kept):
        dy = (rows[0]["top"] + k * pitch) - r["top"]
        for sh in r["shapes"]:
            if sh._element.getparent() is not None:
                sh.top = sh.top + dy
        if r["name"].lower() in want and green is not None:
            g = copy.deepcopy(green._element)
            g.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("id", str(next_id))
            g.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("name", f"Included {r['name']}")
            next_id += 1
            off = g.find(f".//{A}off")
            off.set("x", str(r["box"].left)); off.set("y", str(r["box"].top))
            r["box"]._element.addnext(g)
            tree.remove(r["box"]._element)
    for j, ln in enumerate(lines):
        if j < len(kept) - 1:
            ln.top = rows[0]["top"] + (j + 1) * pitch - (rows[1]["top"] - lines[0].top)
        else:
            tree.remove(ln._element)
    gd = panel._element.find(f".//{A}prstGeom/{A}avLst/{A}gd")
    r_abs = int(gd.get("fmla").split()[1]) / 100000 * min(panel.width, panel.height) if gd is not None else None
    bottom_pad = (panel.top + panel.height) - rows[-1]["top"]
    panel.height = int(rows[0]["top"] + (len(kept) - 1) * pitch + bottom_pad - panel.top)
    if r_abs:
        gd.set("fmla", f"val {round(r_abs / min(panel.width, panel.height) * 100000)}")
    for sh in helpers:
        tree.remove(sh._element)
    prs.save(out or path)
    print(f"{out or path}: included {[r['name'] for r in kept if r['name'].lower() in want]}; "
          + ("kept the other rows unchecked" if keep else f"removed {[r['name'] for r in rows if r not in kept]}")
          + "; helper shapes removed")


if __name__ == "__main__":
    a = sys.argv[1:]
    if not a:
        sys.exit(__doc__)
    if a[0] == "list":
        cmd_list(a[1], int(a[a.index("--slide") + 1]) if "--slide" in a else None)
    elif a[0] == "pick" and len(a) == 4:
        cmd_pick(a[1], a[2], a[3])
    elif a[0] == "replace" and len(a) in (3, 4):
        cmd_replace(*a[1:])
    elif a[0] == "delete" and len(a) in (3, 4):
        cmd_delete(*a[1:])
    elif a[0] == "merge" and len(a) in (4, 5):
        cmd_merge(*a[1:])
    elif a[0] == "sow-products" and len([x for x in a if not x.startswith("--")]) in (3, 4):
        b = [x for x in a if not x.startswith("--")]
        cmd_sow_products(b[1], b[2], b[3] if len(b) > 3 else None, keep="--keep-unchecked" in a)
    elif a[0] == "leftovers" and len(a) == 2:
        sys.exit(0 if cmd_leftovers(a[1]) else 1)
    else:
        sys.exit(__doc__)
