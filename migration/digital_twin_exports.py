"""Digital Twins IQ exports: the deck (.pptx), the PDF and the Focus Group IQ
session PDF (2026-10-26).

Every file is drawn from the payload alone (no model call), on the
Crosswalk design: Graphite Teal and Off-White grounds, Crosswalk Inter
from the design pack, the grid from 0.920 to 12.413, one accent per
page, headers of twelve words or fewer in plain English, the TLDR page
first, and every dollar printed exact with its cents. Each file leaves
through ``design_skill.finish_bytes`` (fonts embedded and checked, the
lint run); the bytes always come back, never held.
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import List, Optional, Tuple

from migration.digital_twin_engine import fmt_money, fmt_people, human_people

GRAPHITE = '#0C1618'
SLATE = '#15252A'
OFFWHITE = '#E9E8E1'
GREEN = '#C7F23E'
OLIVE = '#5E7E12'
ORCHID = '#E682FF'
AMETHYST = '#8E3FA8'
PAVEMENT = '#3B3D38'
BODY_DARK = '#9AA09B'
MUTED_DARK = '#7C878A'
BODY_LIGHT = '#5C6560'
MUTED_LIGHT = '#888C89'
FOOTER = 'CROSSWALK  /  BEHAVIORAL INTELLIGENCE ENGINE'

_ROOT = Path(__file__).resolve().parents[1]
_LOGO_WHITE = _ROOT / 'static' / 'crosswalk-logo-white-transparent.png'
_LOGO_BLACK = _ROOT / 'static' / 'crosswalk-logo-deck.png'


def _design():
    try:
        from migration import design_skill as _ds
        return _ds
    except Exception:
        return None


def _face(weight='Regular') -> str:
    ds = _design()
    try:
        return ds.family_name(weight) if ds is not None else 'Crosswalk Inter'
    except Exception:
        return 'Crosswalk Inter'


def _words(s: str, n: int = 12) -> str:
    """A header of n words or fewer, ending in a full stop."""
    w = str(s or '').replace('\n', ' ').split()
    out = ' '.join(w[:n]).rstrip(',;:')
    if out and not out.endswith(('.', '?', '!')):
        out += '.'
    return out


def _fin(data: bytes, ext: str) -> bytes:
    ds = _design()
    if ds is None:
        return data
    try:
        return ds.finish_bytes(data, ext)
    except Exception:
        return data


# ----------------------------------------------------------------- pages
def _pages(payload: dict) -> list:
    """The deck as a list of page dicts: {kind, header, sub, body:[...],
    rows:[(label, value, note)], tiles:[(value, label)], dark}. One
    finding per page; the TLDR first."""
    m = payload.get('meta') or {}
    name = str(m.get('name') or m.get('subject') or 'the twin')
    rd = payload.get('the_read') or {}
    ident = payload.get('identity') or {}
    pages = [{'kind': 'cover', 'header': f"{name} x Digital Twins IQ", 'sub': f"What the {name} audience does next, run forward on its twin.", 'dark': True}]
    pages.append({'kind': 'bullets', 'header': 'What we found, and what to do.', 'sub': '', 'body': list(rd.get('tldr') or [])[:5], 'dark': True})
    demo = ident.get('demographics') or {}
    rows = []
    for cat in ('gender', 'age'):
        for r in (demo.get(cat) or [])[:6]:
            rows.append((f"{cat.title()}: {r['label']}", f"{float(r['pct']):.1f}%", ''))
    for p in (ident.get('platforms') or [])[:5]:
        rows.append((f"On {p['label']}", f"{float(p['reach_pct']):.1f}%", 'in a year'))
    pages.append({'kind': 'table', 'header': _words(f"{human_people(m.get('size_us'))} people, and here is who they are"), 'sub': ident.get('read') or '', 'rows': rows[:14], 'head': ('', 'Share', ''), 'dark': False})
    traits = payload.get('traits') or {}
    tr_rows = sorted(traits.items(), key=lambda kv: -abs(float(kv[1].get('value') or 50) - 50))[:8]
    pages.append({'kind': 'table', 'header': 'How this audience responds, strongest dials first.', 'sub': 'Each dial runs 0 to 100; 50 is an ordinary US audience.',
                  'rows': [(v.get('label'), f"{v.get('value')}", f"they {v.get('read')}") for _k, v in tr_rows], 'head': ('Dial', 'Value', ''), 'dark': True})
    lc = payload.get('lifecycle') or {}
    if lc:
        pages.append({'kind': 'tiles', 'header': _words(f"{lc.get('habit_pct', 0):.1f}% are in the habit; {lc.get('lapsing_pct', 0):.1f}% are drifting away"),
                      'sub': lc.get('read') or '', 'tiles': [(f"{lc.get(k, 0):.1f}%", k[:-4].title()) for k in ('discovering_pct', 'trying_pct', 'habit_pct', 'lapsing_pct', 'returning_pct')], 'dark': False})
    for r in (payload.get('runs') or [])[:4]:
        o = r.get('outputs') or {}
        usd = o.get('unit') == 'usd'
        f = (lambda v: fmt_money(v)) if usd else (lambda v: f"{human_people(v)} {o.get('noun') or 'people'}")
        header = r.get('read', {}).get('tldr', [''])[0] if r.get('read') else ''
        header = _words(header or r.get('recommendation') or r.get('question'))
        tiles = [(f(o.get('expected')), 'Expected'), (f(o.get('low')), 'Low'), (f(o.get('high')), 'High')]
        if o.get('odds_pct') is not None:
            tiles.append((f"{o.get('odds_pct')}%", f"Odds of clearing {fmt_people((o.get('goal') or {}).get('value'))}"))
        body = list(r.get('drivers') or [])[:3] + [r.get('recommendation') or '']
        pages.append({'kind': 'tiles', 'header': header, 'sub': f"{r.get('label')}: {r.get('question')}", 'tiles': tiles[:4], 'body': [b for b in body if b][:4], 'dark': len(pages) % 2 == 0})
    for sp in (((payload.get('labs') or {}).get('spend')) or [])[:1]:
        led = sp.get('ledger') or {}
        pages.append({'kind': 'tiles', 'header': _words(led.get('read') or sp.get('label') or 'The plan, run forward day by day'), 'sub': sp.get('label') or '',
                      'tiles': [(fmt_money(led.get('saved_total')), 'Saved against the plan as allocated'), (f"{led.get('moves_approved')}", 'Moves approved'),
                                (human_people(led.get('signal_led_outcome')), 'First actions, signal-led'), (fmt_money(led.get('signal_led_cpa')), 'Cost per first action')],
                      'body': [(sp.get('holdout') or {}).get('lift_read') or ''], 'dark': False})
    for ps in (payload.get('panel_sessions') or [])[:1]:
        qs = ps.get('questions') or []
        if qs:
            q = qs[-1]
            pages.append({'kind': 'table', 'header': _words(q.get('read') or q.get('q')), 'sub': f"Focus Group IQ, {ps.get('title')}: {q.get('q')}",
                          'rows': [(r['label'], f"{r['pct']}%", f"{r['people']:,} people") for r in (q.get('distribution') or [])[:8]], 'head': ('Answer', 'Share', 'People'),
                          'body': [f"{v['name']}, {v['age']}, {v['city']}: {v['line']}" for v in (q.get('voices') or [])[:3]], 'dark': True})
    pages.append({'kind': 'bullets', 'header': 'What to do next.', 'sub': '', 'body': list(rd.get('what_to_do') or [])[:4], 'dark': False})
    return pages


# ------------------------------------------------------------------ pptx
def build_deck(payload: dict, pages: Optional[list] = None) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    def rgb(h):
        return RGBColor(int(h[1:3], 16), int(h[3:5], 16), int(h[5:7], 16))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    pages = pages if pages is not None else _pages(payload)
    total = len(pages)
    L, R = 0.92, 12.413

    ds = _design()

    def safe_w(s, size, w, bold=False):
        try:
            return ds.widow_safe_width(str(s or ''), size, w, bold=bold) if ds is not None else w
        except Exception:
            return w

    def lines_of(s, size, w):
        per = max(8, int(w * 72 / (size * 0.52)))
        return max(1, -(-len(str(s or '')) // per))

    def slidenum(slide, x, y, w, h, size, color):
        from lxml import etree
        from pptx.oxml.ns import qn
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        fld = etree.SubElement(p._p, qn('a:fld'))
        fld.set('id', '{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}')
        fld.set('type', 'slidenum')
        rpr = etree.SubElement(fld, qn('a:rPr'))
        rpr.set('lang', 'en-US')
        rpr.set('sz', str(int(size * 100)))
        sf = etree.SubElement(rpr, qn('a:solidFill'))
        c = etree.SubElement(sf, qn('a:srgbClr'))
        c.set('val', color[1:].upper())
        lat = etree.SubElement(rpr, qn('a:latin'))
        lat.set('typeface', _face('Regular'))
        t = etree.SubElement(fld, qn('a:t'))
        t.text = '1'
        return tb

    def text(slide, x, y, w, h, s, size, color, weight='Regular', align=PP_ALIGN.LEFT):
        if size >= 10 and align == PP_ALIGN.LEFT:
            w = safe_w(s, size, w, bold=weight in ('Bold', 'ExtraBold', 'Black'))
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = str(s or '')
        run.font.name = _face(weight)
        run.font.size = Pt(size)
        run.font.bold = weight in ('Bold', 'ExtraBold', 'Black')
        run.font.color.rgb = rgb(color)
        return tb

    for i, pg in enumerate(pages, 1):
        dark = bool(pg.get('dark'))
        alt = pg['kind'] == 'bullets' and i == total        # the closing page is the deck's Signal Green gradient page
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        if alt:
            bg.gradient()
            bg.gradient_angle = 35.0
            stops = bg.gradient_stops
            stops[0].color.rgb = rgb(GREEN)
            stops[1].color.rgb = rgb('#9FD628')
            dark = False
        else:
            bg.solid()
            bg.fore_color.rgb = rgb(GRAPHITE if dark else OFFWHITE)
        ink, body, muted, accent = (OFFWHITE, BODY_DARK, MUTED_DARK, GREEN) if dark else (GRAPHITE, BODY_LIGHT, MUTED_LIGHT, OLIVE)
        if alt:
            ink, body, muted, accent = GRAPHITE, '#1F2B2E', '#2E3A3D', GRAPHITE
        logo = _LOGO_WHITE if dark else _LOGO_BLACK
        try:
            if logo.exists():
                slide.shapes.add_picture(str(logo), Inches(L), Inches(0.45), height=Inches(0.26))
        except Exception:
            pass
        text(slide, 9.4, 0.47, 3.0, 0.3, 'DIGITAL TWINS IQ', 8.5, accent, 'Medium', PP_ALIGN.RIGHT)
        text(slide, L, 7.02, 6.0, 0.3, FOOTER, 7.5, muted, 'Regular')
        slidenum(slide, 11.4, 7.02, 1.0, 0.3, 7.5, muted)
        if pg['kind'] == 'cover':
            hl = lines_of(pg['header'], 40, 10.8)
            text(slide, L, 2.2, 10.8, 0.62 * hl + 0.2, pg['header'], 40, ink, 'Bold')
            text(slide, L, 2.5 + 0.62 * hl, 10.0, 0.9, pg.get('sub') or '', 16, body, 'Regular')
            continue
        hl = lines_of(pg['header'], 26, 11.4)
        text(slide, L, 1.05, 11.4, 0.42 * hl + 0.1, pg['header'], 26, ink, 'Bold')
        y = 1.05 + 0.42 * hl + 0.22
        if pg.get('sub'):
            sl = lines_of(pg['sub'], 11.5, 11.4)
            text(slide, L, y, 11.4, 0.2 * sl + 0.1, pg['sub'], 11.5, body, 'Regular')
            y += 0.2 * sl + 0.3
        if pg['kind'] == 'bullets':
            for b in pg.get('body') or []:
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(L), Inches(y + 0.1), Inches(0.12), Inches(0.12))
                dot.fill.solid()
                dot.fill.fore_color.rgb = rgb(accent)
                dot.line.fill.background()
                text(slide, L + 0.3, y, 10.9, 0.9, b, 15, ink, 'Medium')
                y += 0.95
        elif pg['kind'] == 'table':
            head = pg.get('head') or ('', '', '')
            text(slide, L, y, 5.6, 0.3, str(head[0]).upper(), 8, muted, 'Medium')
            text(slide, L + 6.0, y, 1.6, 0.3, str(head[1]).upper(), 8, muted, 'Medium', PP_ALIGN.RIGHT)
            text(slide, L + 7.9, y, 3.5, 0.3, str(head[2]).upper(), 8, muted, 'Medium')
            y += 0.34
            for lab, val, note in (pg.get('rows') or [])[:12]:
                text(slide, L, y, 5.8, 0.32, str(lab), 11, ink, 'Regular')
                text(slide, L + 6.0, y, 1.6, 0.32, str(val), 11, accent, 'Medium', PP_ALIGN.RIGHT)
                text(slide, L + 7.9, y, 3.6, 0.32, str(note or ''), 9.5, body, 'Regular')
                y += 0.34
            for b in (pg.get('body') or [])[:3]:
                text(slide, L, y + 0.1, 11.4, 0.5, b, 10.5, body, 'Regular')
                y += 0.5
        elif pg['kind'] == 'tiles':
            tiles = pg.get('tiles') or []
            w = (R - L - 0.28 * (len(tiles) - 1)) / max(1, len(tiles))
            for j, (val, lab) in enumerate(tiles):
                x = L + j * (w + 0.28)
                box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.45))
                box.fill.solid()
                box.fill.fore_color.rgb = rgb(SLATE if dark else '#DDDCD4')
                box.line.fill.background()
                text(slide, x + 0.22, y + 0.22, w - 0.4, 0.6, str(val), 22 if len(str(val)) < 16 else 15, accent, 'Bold')
                text(slide, x + 0.22, y + 0.92, w - 0.4, 0.45, str(lab), 9.5, body, 'Regular')
            y += 1.75
            for b in (pg.get('body') or [])[:4]:
                text(slide, L, y, 11.4, 0.55, b, 11.5, ink if b is (pg.get('body') or [''])[-1] else body, 'Regular')
                y += 0.6
    buf = io.BytesIO()
    prs.save(buf)
    return _fin(buf.getvalue(), '.pptx')


def deck_filename(payload: dict) -> str:
    m = payload.get('meta') or {}
    return re.sub(r"[^A-Za-z0-9]+", '_', f"{m.get('name') or m.get('subject')} Digital Twin").strip('_') + '.pptx'


# ------------------------------------------------------------------- pdf
def _font_file(weight: str):
    ds = _design()
    try:
        p = ds.font_path(weight) if ds is not None else None
        if p:
            return Path(p)
    except Exception:
        pass
    p = _ROOT / 'static' / 'fonts' / f'CrosswalkInter-{weight}.ttf'
    return p if p.exists() else None


def _fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    out = {}
    for key, w in (('regular', 'Regular'), ('medium', 'Medium'), ('bold', 'Bold')):
        path = _font_file(w)
        face = f"CWInter-{key}"
        try:
            if path is not None and path.exists():
                if face not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(face, str(path)))
                out[key] = face
            else:
                raise FileNotFoundError(w)
        except Exception:
            out[key] = {'regular': 'Helvetica', 'medium': 'Helvetica', 'bold': 'Helvetica-Bold'}[key]
    return out


class _Doc:
    def __init__(self):
        from reportlab.lib.pagesizes import landscape, letter
        from reportlab.pdfgen import canvas
        self.buf = io.BytesIO()
        self.W, self.H = landscape(letter)
        self.f = _fonts()
        self.c = canvas.Canvas(self.buf, pagesize=(self.W, self.H), initialFontName=self.f['regular'])
        self.page = 0
        self.m = 0.92 * 72
        self.right = self.W - 0.587 * 72
        self.dark = True

    def _wrap(self, text, font, size, width):
        from reportlab.pdfbase.pdfmetrics import stringWidth
        words, lines, cur = str(text or '').split(), [], ''
        for w in words:
            trial = (cur + ' ' + w).strip()
            if stringWidth(trial, font, size) <= width:
                cur = trial
            else:
                if cur:
                    lines.append(cur)
                cur = w
        if cur:
            lines.append(cur)
        return lines or ['']

    def new_page(self, dark, eyebrow, header, sub=''):
        from reportlab.lib import colors
        if self.page:
            self.c.showPage()
        self.page += 1
        self.dark = dark
        c = self.c
        c.setFillColor(colors.HexColor(GRAPHITE if dark else OFFWHITE))
        c.rect(0, 0, self.W, self.H, stroke=0, fill=1)
        try:
            from reportlab.lib.utils import ImageReader
            logo = _LOGO_WHITE if dark else _LOGO_BLACK
            if logo.exists():
                c.drawImage(ImageReader(str(logo)), self.m, self.H - 0.72 * 72, width=1.25 * 72, height=0.26 * 72, preserveAspectRatio=True, anchor='sw', mask='auto')
        except Exception:
            pass
        c.setFont(self.f['medium'], 8.5)
        c.setFillColor(colors.HexColor(GREEN if dark else OLIVE))
        c.drawRightString(self.right, self.H - 0.62 * 72, str(eyebrow).upper())
        c.setFont(self.f['regular'], 7.5)
        c.setFillColor(colors.HexColor(MUTED_DARK if dark else MUTED_LIGHT))
        c.drawString(self.m, 0.42 * 72, FOOTER)
        c.drawRightString(self.right, 0.42 * 72, f"{self.page:02d}")
        c.setFillColor(colors.HexColor(OFFWHITE if dark else GRAPHITE))
        c.setFont(self.f['bold'], 20)
        y = self.H - 1.35 * 72
        for line in self._wrap(header, self.f['bold'], 20, self.right - self.m):
            c.drawString(self.m, y, line)
            y -= 26
        if sub:
            c.setFont(self.f['regular'], 10.5)
            c.setFillColor(colors.HexColor(BODY_DARK if dark else BODY_LIGHT))
            for line in self._wrap(sub, self.f['regular'], 10.5, self.right - self.m)[:4]:
                c.drawString(self.m, y - 4, line)
                y -= 15
        self.y = y - 18
        return self.y

    def paragraph(self, text, x, y, width, size=10.5, color='', font=''):
        from reportlab.lib import colors
        font = font or self.f['regular']
        self.c.setFont(font, size)
        self.c.setFillColor(colors.HexColor(color or (BODY_DARK if self.dark else BODY_LIGHT)))
        for line in self._wrap(text, font, size, width):
            self.c.drawString(x, y, line)
            y -= size * 1.42
        return y

    def bullets(self, items, x, y, width, size=12):
        from reportlab.lib import colors
        for it in items:
            self.c.setFillColor(colors.HexColor(GREEN if self.dark else OLIVE))
            self.c.circle(x + 4, y + size * 0.32, 2.6, stroke=0, fill=1)
            y = self.paragraph(it, x + 16, y, width - 16, size=size, color=OFFWHITE if self.dark else GRAPHITE, font=self.f['medium'])
            y -= 7
        return y

    def table(self, rows, x, y, width, head=('', '', '')):
        from reportlab.lib import colors
        ink = OFFWHITE if self.dark else GRAPHITE
        body = BODY_DARK if self.dark else BODY_LIGHT
        c = self.c
        c.setFont(self.f['medium'], 8)
        c.setFillColor(colors.HexColor(MUTED_DARK if self.dark else MUTED_LIGHT))
        c.drawString(x, y, str(head[0]).upper())
        c.drawRightString(x + width * 0.55, y, str(head[1]).upper())
        c.drawString(x + width * 0.6, y, str(head[2]).upper())
        y -= 5
        c.setStrokeColor(colors.HexColor(PAVEMENT if self.dark else '#C9C8C0'))
        c.setLineWidth(0.6)
        c.line(x, y, x + width, y)
        y -= 13
        for lab, val, note in rows:
            c.setFont(self.f['regular'], 9.5)
            c.setFillColor(colors.HexColor(ink))
            c.drawString(x, y, str(lab)[:70])
            c.setFont(self.f['medium'], 9.5)
            c.setFillColor(colors.HexColor(GREEN if self.dark else OLIVE))
            c.drawRightString(x + width * 0.55, y, str(val))
            c.setFont(self.f['regular'], 8.5)
            c.setFillColor(colors.HexColor(body))
            c.drawString(x + width * 0.6, y, self._wrap(note or '', self.f['regular'], 8.5, width * 0.4)[0])
            y -= 14.5
        return y

    def tiles(self, tiles, x, y, width):
        from reportlab.lib import colors
        n = max(1, len(tiles))
        gap = 14
        w = (width - gap * (n - 1)) / n
        for j, (val, lab) in enumerate(tiles):
            xx = x + j * (w + gap)
            self.c.setFillColor(colors.HexColor(SLATE if self.dark else '#DDDCD4'))
            self.c.roundRect(xx, y - 78, w, 78, 6, stroke=0, fill=1)
            self.c.setFont(self.f['bold'], 18 if len(str(val)) < 16 else 12.5)
            self.c.setFillColor(colors.HexColor(GREEN if self.dark else OLIVE))
            self.c.drawString(xx + 12, y - 34, str(val))
            self.c.setFont(self.f['regular'], 8.5)
            self.c.setFillColor(colors.HexColor(BODY_DARK if self.dark else BODY_LIGHT))
            self.c.drawString(xx + 12, y - 60, self._wrap(lab, self.f['regular'], 8.5, w - 24)[0])
        return y - 96

    def finish(self):
        self.c.save()
        return self.buf.getvalue()


def build_pdf(payload: dict) -> bytes:
    d = _Doc()
    for i, pg in enumerate(_pages(payload), 1):
        dark = bool(pg.get('dark'))
        if pg['kind'] == 'cover':
            d.new_page(True, 'Digital Twins IQ', pg['header'], pg.get('sub') or '')
            continue
        y = d.new_page(dark, 'Digital Twins IQ', pg['header'], pg.get('sub') or '')
        width = d.right - d.m
        if pg['kind'] == 'bullets':
            d.bullets(pg.get('body') or [], d.m, y, width)
        elif pg['kind'] == 'table':
            y = d.table((pg.get('rows') or [])[:14], d.m, y, width, head=pg.get('head') or ('', '', ''))
            for b in (pg.get('body') or [])[:3]:
                y = d.paragraph(b, d.m, y - 6, width, size=9.5)
        elif pg['kind'] == 'tiles':
            y = d.tiles(pg.get('tiles') or [], d.m, y, width)
            for b in (pg.get('body') or [])[:4]:
                y = d.paragraph(b, d.m, y - 4, width, size=10.5, color=OFFWHITE if dark else GRAPHITE, font=d.f['medium'])
                y -= 4
    return _fin(d.finish(), '.pdf')


def pdf_filename(payload: dict) -> str:
    m = payload.get('meta') or {}
    return re.sub(r"[^A-Za-z0-9]+", '_', f"{m.get('name') or m.get('subject')} Digital Twin").strip('_') + '.pdf'


def _portrait_reader(portrait_fn):
    """A memoized reader: hash -> ImageReader or None. Every failure is
    a missing face, never a failed export."""
    cache = {}

    def read(h):
        if not h or portrait_fn is None:
            return None
        if h in cache:
            return cache[h]
        img = None
        try:
            data = portrait_fn(h)
            if data:
                from reportlab.lib.utils import ImageReader
                img = ImageReader(io.BytesIO(data))
        except Exception:
            img = None
        cache[h] = img
        return img
    return read


def _draw_round_portrait(c, img, x, y, size):
    """A portrait in a circle at (x, y) bottom-left; silent on failure."""
    try:
        c.saveState()
        p = c.beginPath()
        p.circle(x + size / 2.0, y + size / 2.0, size / 2.0)
        c.clipPath(p, stroke=0, fill=0)
        c.drawImage(img, x, y, width=size, height=size, preserveAspectRatio=True, anchor='c', mask='auto')
        c.restoreState()
        return True
    except Exception:
        try:
            c.restoreState()
        except Exception:
            pass
        return False


def build_panel_pdf(payload: dict, session: dict, portrait_fn=None) -> bytes:
    """The Focus Group IQ session: the room, then one page per question with
    the % table and the voices. ``portrait_fn(hash) -> JPEG bytes`` puts
    the twins' portraits on the room page and beside each voice; without
    it (or when a face is absent) the text stands alone."""
    d = _Doc()
    m = payload.get('meta') or {}
    g = session.get('group') or {}
    read = _portrait_reader(portrait_fn)
    d.new_page(True, 'Focus Group IQ', _words(f"{session.get('title') or 'The room'}"), f"{g.get('in_room')} twins in the room, standing for {int(g.get('size_people') or 0):,} people. Moderated by {session.get('created_by')}.")
    y = d.y
    y = d.table([(s['label'], f"{s['share_pct']}%", f"{int(s['people']):,} people, {s.get('seats')} seats") for s in g.get('segments') or []], d.m, y, d.right - d.m, head=('Segment', 'Share of room', ''))
    roster_rows = (g.get('roster') or [])[:16]
    faces = [(r, read(r.get('portrait_hash'))) for r in roster_rows]
    if any(img is not None for _r, img in faces):
        size, gap = 0.46 * 72, 0.12 * 72
        per_row = max(1, int((d.right - d.m + gap) // (size + gap)))
        yy = y - 10 - size
        for j, (r, img) in enumerate(faces):
            xx = d.m + (j % per_row) * (size + gap)
            if j and j % per_row == 0:
                yy -= size + gap + 12
            from reportlab.lib import colors
            d.c.setFillColor(colors.HexColor(SLATE))
            d.c.circle(xx + size / 2.0, yy + size / 2.0, size / 2.0, stroke=0, fill=1)
            if img is not None:
                _draw_round_portrait(d.c, img, xx, yy, size)
            d.c.setFont(d.f['regular'], 6.5)
            d.c.setFillColor(colors.HexColor(BODY_DARK))
            d.c.drawCentredString(xx + size / 2.0, yy - 8, str(r.get('name') or '')[:12])
        y = yy - 22
    roster = ', '.join(f"{r['name']} ({r['age']}, {r['city']})" for r in roster_rows)
    d.paragraph('In the room: ' + roster, d.m, y - 8, d.right - d.m, size=9.5)
    by_seat = {r.get('seat'): r for r in (g.get('roster') or [])}
    by_name = {str(r.get('name')): r for r in (g.get('roster') or [])}
    for i, q in enumerate(session.get('questions') or []):
        y = d.new_page(i % 2 == 1, 'Focus Group IQ', _words(q.get('read') or q.get('q')), f"Question {i + 1}: {q.get('q')}")
        y = d.table([(r['label'], f"{r['pct']}%", f"{int(r['people']):,} people") for r in (q.get('distribution') or [])[:8]], d.m, y, d.right - d.m, head=('Answer', 'Share', 'People'))
        for v in (q.get('voices') or [])[:4]:
            seat = by_seat.get(v.get('seat')) or by_name.get(str(v.get('name') or '')) or {}
            img = read(seat.get('portrait_hash'))
            x = d.m
            if img is not None:
                size = 0.34 * 72
                _draw_round_portrait(d.c, img, d.m, y - 6 - size + 9, size)
                x = d.m + size + 8
            y = d.paragraph(f"{v['name']}, {v['age']}, {v['city']}: {v['line']}", x, y - 6, d.right - x, size=10, color=OFFWHITE if d.dark else GRAPHITE, font=d.f['medium'])
    return _fin(d.finish(), '.pdf')


def panel_pdf_filename(payload: dict, session: dict) -> str:
    m = payload.get('meta') or {}
    return re.sub(r"[^A-Za-z0-9]+", '_', f"{m.get('name') or m.get('subject')} Focus Group IQ {session.get('id')}").strip('_') + '.pdf'


def _clock(t) -> str:
    try:
        s = max(0, int(round(float(t))))
    except (TypeError, ValueError):
        s = 0
    return f"{s // 60}:{s % 60:02d}"


def _pilot_pages(payload: dict, session: dict) -> list:
    m = payload.get('meta') or {}
    name = str(m.get('name') or m.get('subject') or 'the twin')
    sc = session.get('scorecard') or {}
    fav = sc.get('favorite_moment') or {}
    dead = sc.get('dead_spot') or {}
    media = session.get('media') or {}
    title = str(session.get('title') or media.get('title') or 'Pilot test')
    pages = [
        {'kind': 'tiles', 'header': f"{name} watched {title}.",
         'sub': 'Overall like, recommend, and finish intent from the room.',
         'tiles': [
             (f"{float(sc.get('overall_like') or 0):.1f}", 'Overall like'),
             (f"{float(sc.get('recommend_pct') or 0):.1f}%", 'Would recommend'),
             (f"{float(sc.get('finish_intent_pct') or 0):.1f}%", 'Would finish'),
         ],
         'body': list(sc.get('notes') or [])[:3], 'dark': True},
        {'kind': 'bullets', 'header': 'The moments that moved the room.',
         'sub': f"Favorite at {fav.get('clock') or _clock(fav.get('t'))}. Dead spot at {dead.get('clock') or _clock(dead.get('t'))}.",
         'body': [
             f"Favorite moment: {fav.get('line') or 'The room peaked here.'}",
             f"Dead spot: {dead.get('line') or 'The room dipped.'}",
         ] + [f"{_clock(e.get('t'))}  {e.get('name')}: {e.get('line')}" for e in (session.get('events') or [])[:6]],
         'dark': False},
    ]
    qs = session.get('questions') or []
    if qs:
        pages.append({'kind': 'bullets', 'header': 'What they said when asked.',
                      'sub': 'Questions sit on the playhead.',
                      'body': [f"{q.get('playhead_clock') or _clock(q.get('playhead_s'))}: {q.get('q')}  {q.get('read') or ''}" for q in qs[:6]],
                      'dark': True})
    pages.append({'kind': 'bullets', 'header': 'What this means for the cut.',
                  'sub': '',
                  'body': [
                      f"{int(sc.get('people') or 0):,} people sit behind this room." if sc.get('people') else 'The room stands for the twin.',
                      'Keep the favorite moment. Look hard at the dead spot.',
                  ], 'dark': False})
    return pages


def build_pilot_deck(payload: dict, session: dict) -> bytes:
    return build_deck(payload, pages=_pilot_pages(payload, session))


def build_pilot_pdf(payload: dict, session: dict, portrait_fn=None) -> bytes:
    d = _Doc()
    m = payload.get('meta') or {}
    g = session.get('group') or {}
    sc = session.get('scorecard') or {}
    fav = sc.get('favorite_moment') or {}
    dead = sc.get('dead_spot') or {}
    read = _portrait_reader(portrait_fn)
    d.new_page(True, 'Pilot Testing', _words(f"{session.get('title') or 'The cut'}"),
               f"{g.get('in_room')} twins in the room. Moderated by Prometheus.")
    y = d.y
    y = d.table([
        ('Overall like', f"{float(sc.get('overall_like') or 0):.1f}", ''),
        ('Would recommend', f"{float(sc.get('recommend_pct') or 0):.1f}%", ''),
        ('Would finish', f"{float(sc.get('finish_intent_pct') or 0):.1f}%", ''),
        ('Favorite moment', f"{fav.get('clock') or ''}", str(fav.get('line') or '')[:80]),
        ('Dead spot', f"{dead.get('clock') or ''}", str(dead.get('line') or '')[:80]),
    ], d.m, y, d.right - d.m, head=('Score', 'Value', 'Note'))
    for n in (sc.get('notes') or [])[:4]:
        y = d.paragraph(n, d.m, y - 4, d.right - d.m, size=10)
    roster_rows = (g.get('roster') or [])[:16]
    faces = [(r, read(r.get('portrait_hash'))) for r in roster_rows]
    if any(img is not None for _r, img in faces):
        size, gap = 0.40 * 72, 0.10 * 72
        per_row = max(1, int((d.right - d.m + gap) // (size + gap)))
        yy = y - 8 - size
        for j, (r, img) in enumerate(faces):
            xx = d.m + (j % per_row) * (size + gap)
            if j and j % per_row == 0:
                yy -= size + gap + 12
            from reportlab.lib import colors
            d.c.setFillColor(colors.HexColor(SLATE))
            d.c.circle(xx + size / 2.0, yy + size / 2.0, size / 2.0, stroke=0, fill=1)
            if img is not None:
                _draw_round_portrait(d.c, img, xx, yy, size)
            d.c.setFont(d.f['regular'], 6.5)
            d.c.setFillColor(colors.HexColor(BODY_DARK))
            d.c.drawCentredString(xx + size / 2.0, yy - 8, str(r.get('name') or '')[:12])
        y = yy - 20
    events = session.get('events') or []
    if events:
        y = d.new_page(False, 'Pilot Testing', _words('Moments on the clock'))
        y = d.table(
            [(f"{_clock(e.get('t'))} {e.get('type')}", f"{e.get('name')}, {e.get('age')}", str(e.get('line') or '')[:70])
             for e in events[:18]],
            d.m, y, d.right - d.m, head=('Clock', 'Twin', 'Line'))
    for i, q in enumerate(session.get('questions') or []):
        y = d.new_page(i % 2 == 0, 'Pilot Testing', _words(q.get('read') or q.get('q')),
                       f"At {q.get('playhead_clock') or _clock(q.get('playhead_s'))}: {q.get('q')}")
        y = d.table([(r['label'], f"{r['pct']}%", f"{int(r.get('people') or 0):,} people")
                     for r in (q.get('distribution') or [])[:8]],
                    d.m, y, d.right - d.m, head=('Answer', 'Share', 'People'))
        for v in (q.get('voices') or [])[:4]:
            y = d.paragraph(f"{v.get('name')}, {v.get('age')}, {v.get('city')}: {v.get('line')}",
                            d.m, y - 6, d.right - d.m, size=10, color=OFFWHITE if d.dark else GRAPHITE, font=d.f['medium'])
    return _fin(d.finish(), '.pdf')


def pilot_pdf_filename(payload: dict, session: dict) -> str:
    m = payload.get('meta') or {}
    return re.sub(r"[^A-Za-z0-9]+", '_', f"{m.get('name') or m.get('subject')} Pilot Testing {session.get('id')}").strip('_') + '.pdf'


def pilot_deck_filename(payload: dict, session: dict) -> str:
    m = payload.get('meta') or {}
    return re.sub(r"[^A-Za-z0-9]+", '_', f"{m.get('name') or m.get('subject')} Pilot Testing {session.get('id')}").strip('_') + '.pptx'


__all__ = ['build_deck', 'deck_filename', 'build_pdf', 'pdf_filename', 'build_panel_pdf', 'panel_pdf_filename',
           'build_pilot_pdf', 'pilot_pdf_filename', 'build_pilot_deck', 'pilot_deck_filename']
